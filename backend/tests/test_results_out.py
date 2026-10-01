"""Results that leave the app and feed the next plan (improvement plan, phase 3).

A small cover model on places with shapes -- open depots so every site has one
within reach, as few as possible -- solved for real, then: its answer on a map,
exported as Excel / CSV / GeoJSON, asked "what if depot D2 is closed?" as a data
what-if, and its decision kept as data for the next problem.
"""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.api.answer_map import answer_map
from app.solve import whatif
from app.solve.service import claim_next, enqueue_run, execute_run
from tests.test_api_problems import auth_headers, client  # noqa: F401
from tests.test_quadratic import empty_queue  # noqa: F401
from tests.test_v1_problem_run import (  # noqa: F401
    db, make_attribute_def, make_domain, make_entity, make_entity_type, make_model_version, make_parameter_def,
    make_parameter_value, make_problem,
)


def _pt(lon, lat):
    return {"type": "Point", "coordinates": [lon, lat]}


IR = {
    "version": 1,
    "sets": ["depot", "site"],
    "parameters": {"reach": {"index": ["depot", "site"]}},
    "variables": {"open": {"index": ["depot"], "domain": "binary"},
                  "serve": {"index": ["depot", "site"], "domain": "binary"}},
    "constraints": [
        {"id": "c_served", "note": "every site is served", "forall": [{"index": "s", "set": "site"}],
         "left": {"sum": {"mul": [{"par": "reach", "index": ["d", "s"]}, {"var": "serve", "index": ["d", "s"]}]},
                  "over": [{"index": "d", "set": "depot"}]},
         "relation": ">=", "right": {"const": 1}, "severity": "soft", "weight": 100},
        {"id": "c_open", "note": "only an open depot serves", "forall": [{"index": "d", "set": "depot"}, {"index": "s", "set": "site"}],
         "left": {"var": "serve", "index": ["d", "s"]}, "relation": "<=", "right": {"var": "open", "index": ["d"]},
         "severity": "hard"},
    ],
    "objective": {"sense": "minimize", "terms": [{"id": "o_open", "weight": 1, "expression": {
        "sum": {"var": "open", "index": ["d"]}, "over": [{"index": "d", "set": "depot"}]}}]},
}


@pytest.fixture
def placed(db):  # noqa: F811
    domain = make_domain(db, "towns")
    depot = make_entity_type(db, domain, "depot", "location")
    site = make_entity_type(db, domain, "site", "location")
    for t in (depot, site):
        make_attribute_def(db, t, "shape", "geometry")
    d = {k: make_entity(db, depot, k, attrs={"shape": _pt(lon, 0)}, label=f"Depot {k}") for k, lon in (("D1", 0.0), ("D2", 1.0))}
    s = {k: make_entity(db, site, k, attrs={"shape": _pt(lon, 0.01)}) for k, lon in (("S1", 0.0), ("S2", 1.0), ("S3", 5.0))}
    reach = make_parameter_def(db, domain, "reach", [depot, site], default_value=0)
    make_parameter_value(db, reach, [d["D1"], s["S1"]], 1)
    make_parameter_value(db, reach, [d["D2"], s["S2"]], 1)
    problem = make_problem(db, domain)
    version = make_model_version(db, problem, IR)
    db.commit()
    return {"domain": domain, "problem": problem, "version": version, "depot": depot, "site": site}


def _scenario(db, placed, patch=None):  # noqa: F811
    scenario = db.execute(text("INSERT INTO scenario (problem_id, model_version_id, name, patch)"
                               " VALUES (:p, :v, :n, CAST(:patch AS jsonb)) RETURNING id"),
                          {"p": placed["problem"], "v": placed["version"], "n": f"s{json.dumps(patch)}",
                           "patch": json.dumps(patch or {})}).scalar_one()
    db.commit()
    return scenario


def _solve(db, scenario):  # noqa: F811
    run = enqueue_run(db, scenario, time_limit=10.0, reuse=False)
    assert claim_next(db) == run
    return run, execute_run(db, run)


def test_any_answer_is_drawn_where_it_happens(db, placed, empty_queue):  # noqa: F811
    run, outcome = _solve(db, _scenario(db, placed))
    assert outcome.status == "optimal"
    row = db.execute(text("SELECT mv.ir, d.data, so.assignments, so.amounts FROM run r JOIN scenario s ON s.id = r.scenario_id"
                          " JOIN model_version mv ON mv.id = s.model_version_id JOIN dataset d ON d.id = r.dataset_id"
                          " JOIN solution so ON so.run_id = r.id WHERE r.id = :r"), {"r": run}).one()
    results = [dict(r) for r in db.execute(text("SELECT constraint_id, satisfied, violations FROM constraint_result WHERE run_id = :r"),
                                           {"r": run}).mappings()]
    mapped = answer_map(row[0], row[1], row[2], row[3], results)
    layers = {l["id"]: l for l in mapped["layers"]}
    assert layers["open"]["kind"] == "places" and layers["open"]["title"] == "open: 2 of 2 depot"
    assert layers["serve"]["kind"] == "links" and layers["serve"]["title"] == "serve: 2 links"
    lines = [f for f in mapped["features"] if f["properties"]["layer"] == "serve"]
    assert {f["properties"]["key"] for f in lines} == {"D1|S1", "D2|S2"}
    assert lines[0]["geometry"]["type"] == "LineString"
    # S3 has no depot within reach: the bent rule is marked at its place.
    unmet = [f for f in mapped["features"] if f["properties"]["layer"] == "unmet"]
    assert [f["properties"]["key"] for f in unmet] == ["S3"] and "c_served short by 1" in unmet[0]["properties"]["title"]
    # Who serves whom, read from the 0/1 reach data and the open depots (user test: which yard covers which hotspot).
    assert layers["reach_served"]["title"] == "site served from a chosen depot (reach): 2 of 3"
    served = {f["properties"]["key"]: f for f in mapped["features"] if f["properties"]["layer"] == "reach_served"}
    assert served["D1|S1"]["geometry"]["type"] == "LineString" and "served from Depot D1" in served["D1|S1"]["properties"]["title"]
    assert served["S3"]["properties"]["status"] == "short" and "no chosen depot within reach" in served["S3"]["properties"]["title"]


def test_an_answer_exports_as_excel_csv_and_geojson(db, placed, empty_queue, client, auth_headers):  # noqa: F811
    run, _ = _solve(db, _scenario(db, placed))
    xlsx = client.get(f"/api/v1/runs/{run}/export", params={"format": "xlsx"}, headers=auth_headers)
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert {"Summary", "open", "serve", "serve grid", "Rules"} <= set(wb.sheetnames)
    rows = list(wb["open"].iter_rows(values_only=True))
    assert rows[0] == ("depot", "depot name", "value") and ("D1", "Depot D1", 1) in rows
    rules = {r[0]: r for r in wb["Rules"].iter_rows(values_only=True)}
    assert rules["c_served"][2] is False and "S3 by 1" in rules["c_served"][6]
    csv_text = client.get(f"/api/v1/runs/{run}/export", params={"format": "csv"}, headers=auth_headers).text
    assert "serve,D1,S1,1" in csv_text.replace("\r", "")
    geo = client.get(f"/api/v1/runs/{run}/export", params={"format": "geojson"}, headers=auth_headers).json()
    assert geo["type"] == "FeatureCollection" and any(f["properties"]["layer"] == "serve" for f in geo["features"])
    mapped = client.get(f"/api/v1/runs/{run}/answer-map", headers=auth_headers).json()
    assert {l["id"] for l in mapped["layers"]} >= {"open", "serve", "unmet"}


def test_an_answer_prints_as_a_report_with_its_map(db, placed, empty_queue, client, auth_headers):  # noqa: F811
    """Improvement plan 3.2: a printable report (saved as PDF by the browser) with the map drawn."""
    run, _ = _solve(db, _scenario(db, placed))
    page = client.get(f"/api/v1/runs/{run}/export", params={"format": "html", "print": "true"}, headers=auth_headers)
    assert page.status_code == 200 and page.headers["content-type"].startswith("text/html")
    html = page.text
    assert "<svg" in html and "<circle" in html, "the answer map is drawn"
    assert "Rules" in html and "window.print" in html


def test_what_if_a_depot_is_closed_is_a_scenario_on_the_same_data(db, placed, empty_queue):  # noqa: F811
    base, _ = _solve(db, _scenario(db, placed))
    closed, outcome = _solve(db, _scenario(db, placed, {"remove": {"depot": ["D2"]}}))
    assert outcome.status == "optimal"
    same_data = db.execute(text("SELECT count(DISTINCT dataset_id) FROM run WHERE id IN (:a, :b)"), {"a": base, "b": closed}).scalar_one()
    assert same_data in (1, 2)  # the stored data is never edited; the what-if is applied to a copy
    opened = db.execute(text("SELECT assignments FROM solution WHERE run_id = :r"), {"r": closed}).scalar_one()["open"]
    assert opened == [["D1"]]
    shortfalls = db.execute(text("SELECT violations FROM constraint_result WHERE run_id = :r AND constraint_id = 'c_served'"),
                            {"r": closed}).scalar_one()
    assert sorted(v["index"][0] for v in shortfalls) == ["S2", "S3"]


def test_a_what_if_scales_and_sets_values_on_a_copy():
    data = {"sets": {"depot": [{"id": "D1"}, {"id": "D2"}]},
            "parameters": {"reach": [{"depot": "D1", "site": "S1", "value": 1}, {"depot": "D2", "site": "S2", "value": 1}]},
            "relationships": {}}
    out = whatif.apply(data, IR, {"scale_param": {"reach": 3}, "set_param": [{"param": "reach", "index": ["D1", "S2"], "value": 1}]})
    assert [r["value"] for r in out["parameters"]["reach"]] == [3, 3, 1]
    assert data["parameters"]["reach"][0]["value"] == 1  # the frozen data is untouched
    assert whatif.describe({"remove": {"depot": ["D2"]}, "scale_param": {"rain": 1.3}}) == ["without depot D2", "rain × 1.3"]


def test_the_api_refuses_a_what_if_on_a_set_the_model_does_not_have(db, placed, client, auth_headers):  # noqa: F811
    got = client.post("/api/v1/scenarios", json={"problem_id": placed["problem"], "model_version_id": placed["version"],
                                                  "name": "bad", "patch": {"remove": {"truck": ["T1"]}}}, headers=auth_headers)
    assert got.status_code == 422 and "whatif_unknown" in got.text


def test_an_answer_is_kept_as_data_for_the_next_problem(db, placed, empty_queue, client, auth_headers):  # noqa: F811
    run, _ = _solve(db, _scenario(db, placed))
    kept = client.post(f"/api/v1/runs/{run}/promote", json={"decision": "serve", "name": "served_by", "as": "relationship"},
                       headers=auth_headers)
    assert kept.status_code == 201, kept.text
    assert kept.json()["links"] == 2
    as_param = client.post(f"/api/v1/runs/{run}/promote", json={"decision": "open", "name": "opened", "as": "parameter"},
                           headers=auth_headers)
    assert as_param.status_code == 201 and as_param.json()["cells"] == 2
    refused = client.post(f"/api/v1/runs/{run}/promote", json={"decision": "open", "name": "x", "as": "relationship"},
                          headers=auth_headers)
    assert refused.status_code == 422


def test_data_checks_name_unfilled_and_stale_map_data(db, placed):  # noqa: F811
    from app.api.preflight import data_findings

    empty = make_parameter_def(db, placed["domain"], "capacity", [placed["depot"]], default_value=0)
    db.execute(text("UPDATE parameter_def SET source = CAST(:s AS jsonb) WHERE domain_id = :d AND name = 'reach'"),
               {"s": json.dumps({"kind": "within", "missing": ["S9"], "computed_at": "2020-01-01T00:00:00+00:00"}), "d": placed["domain"]})
    db.commit()
    ir = {**IR, "parameters": {**IR["parameters"], "capacity": {"index": ["depot"]}}}
    codes = {f["code"]: f for f in data_findings(db, placed["domain"], ir)}
    assert set(codes) == {"parameter_unfilled", "computed_without_places", "computed_stale"}
    assert codes["parameter_unfilled"]["parameter"] == "capacity" and empty
    assert "S9" in codes["computed_without_places"]["says"]



def test_a_new_field_on_the_places_is_not_a_move(db, placed):  # noqa: F811
    """Writing a height or a count on the same records leaves map data current; moving one does not."""
    from app.api.preflight import data_findings
    from app.spatial.ops import shapes_fingerprint

    types = [placed["depot"], placed["site"]]
    db.execute(text("UPDATE parameter_def SET source = CAST(:s AS jsonb) WHERE domain_id = :d AND name = 'reach'"),
               {"s": json.dumps({"kind": "within", "computed_at": "2020-01-01T00:00:00+00:00",
                                 "shapes": shapes_fingerprint(db, types)}), "d": placed["domain"]})
    make_attribute_def(db, placed["site"], "ground_m", "number")
    db.execute(text("UPDATE entity SET attrs = attrs || '{\"ground_m\": 4}'::jsonb WHERE entity_type_id = :t"), {"t": placed["site"]})
    db.commit()
    assert "computed_stale" not in {f["code"] for f in data_findings(db, placed["domain"], IR)}
    db.execute(text("UPDATE entity SET attrs = jsonb_set(attrs, '{shape}', CAST(:g AS jsonb)) WHERE entity_type_id = :t AND key = 'S1'"),
               {"t": placed["site"], "g": json.dumps(_pt(9.0, 0))})
    db.commit()
    assert "computed_stale" in {f["code"] for f in data_findings(db, placed["domain"], IR)}

def test_two_answers_compare_on_the_map():
    """Improvement plan 3.4: only what changed is drawn -- opened, closed, newly short, fixed."""
    from app.api.answer_map import compare_maps

    def f(layer, key, status):
        return {"type": "Feature", "geometry": {"type": "Point", "coordinates": [0, 0]},
                "properties": {"layer": layer, "key": key, "label": key, "status": status, "title": key}}

    before = {"features": [f("open", "D1", "chosen"), f("open", "D2", "chosen"), f("unmet", "S3", "short")]}
    now = {"features": [f("open", "D2", "chosen"), f("open", "D3", "chosen"), f("open", "D1", "not_chosen")]}
    out = compare_maps(now, before)
    status = {(x["properties"]["layer"], x["properties"]["key"]): x["properties"]["status"] for x in out["features"]}
    assert status == {("open", "D1"): "removed", ("open", "D2"): "same", ("open", "D3"): "added", ("unmet", "S3"): "fixed"}
    assert out["changed"] == 3
    assert {l["id"]: l["title"] for l in out["layers"]}["open"] == "open: +1, −1"


def test_the_approved_plan_is_fetched_by_the_systems_that_act_on_it(db, placed, empty_queue, client, auth_headers):  # noqa: F811
    """Improvement plan 3.6: an approved plan as JSON (or any export), never the newest run."""
    problem = placed["problem"]
    assert client.get(f"/api/v1/problems/{problem}/approved-plan", headers=auth_headers).status_code == 404
    run, _ = _solve(db, _scenario(db, placed))
    approved = client.post(f"/api/v1/runs/{run}/approve", json={"reason": "for dispatch"}, headers=auth_headers)
    assert approved.status_code == 201, approved.text
    _solve(db, _scenario(db, placed, {"name": "newer"}))  # a newer run does not replace the approved one
    plan = client.get(f"/api/v1/problems/{problem}/approved-plan", headers=auth_headers).json()
    assert plan["run"]["id"] == run and plan["approval"]["reason"] == "for dispatch"
    serve = plan["decisions"]["serve"]
    assert serve["index"] == ["depot", "site"] and {tuple(r["keys"]) for r in serve["rows"]} == {("D1", "S1"), ("D2", "S2")}
    xlsx = client.get(f"/api/v1/problems/{problem}/approved-plan", params={"format": "xlsx"}, headers=auth_headers)
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    db.execute(text("DELETE FROM approved_plan WHERE problem_id = :p"), {"p": problem})
    db.commit()


def test_data_that_follows_a_problem_is_rewritten_when_a_new_plan_is_approved(db, placed, empty_queue, client, auth_headers):  # noqa: F811
    """Improvement plan 5.1: placement then roster -- the next problem reads the approved placement."""
    problem = placed["problem"]
    first, _ = _solve(db, _scenario(db, placed))
    kept = client.post(f"/api/v1/runs/{first}/promote", json={"decision": "open", "name": "open_plan", "as": "parameter",
                                                              "follow": True}, headers=auth_headers)
    assert kept.status_code == 201, kept.text
    assert kept.json()["source"]["follow"] == "approved"
    second, _ = _solve(db, _scenario(db, placed, {"name": "second"}))
    assert client.post(f"/api/v1/runs/{second}/approve", json={"reason": "go"}, headers=auth_headers).status_code == 201
    source = db.execute(text("SELECT source FROM parameter_def WHERE name = 'open_plan'")).scalar_one()
    assert source["run_id"] == second and source["follow"] == "approved"
    db.execute(text("DELETE FROM approved_plan WHERE problem_id = :p"), {"p": problem})
    db.commit()


def test_the_printed_map_has_the_chosen_base_map_under_it(monkeypatch):
    """User test (Alexandria): the printed report had no background, so the lines floated on white."""
    import io as _io
    import urllib.request

    from PIL import Image

    from app.api import run_export

    fetched = []

    class Tile:
        def __init__(self, url):
            fetched.append(url)
            buffer = _io.BytesIO()
            Image.new("RGB", (256, 256), (30, 90, 40)).save(buffer, "PNG")
            self.body = buffer.getvalue()

        def read(self):
            return self.body

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(urllib.request, "urlopen", lambda request, timeout=0: Tile(request.full_url))
    features = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [29.9, 31.2]},
                 "properties": {"layer": "open", "status": "chosen", "label": "Y1", "title": "Y1"}},
                {"type": "Feature", "geometry": {"type": "Point", "coordinates": [30.05, 31.25]},
                 "properties": {"layer": "hotspot", "status": "place", "title": "H01"}}]
    plain = run_export._svg_map(features)
    assert "<image" not in plain and ">Y1</text>" in plain and ">H01</text>" not in plain  # chosen places named
    under = run_export._svg_map(features, basemap=run_export.basemap_of(None, None, "builtin-streets"))
    assert under.count("<image") == 1 and "data:image/jpeg;base64," in under
    assert ">Y1</text>" in under and "OpenStreetMap" in under
    assert fetched and all(u.startswith("https://tile.openstreetmap.org/") for u in fetched)
    # Only base maps the app offers are fetched: an address in the request is not one.
    assert run_export.basemap_of(None, None, "https://evil.example/{z}/{x}/{y}.png") is None
