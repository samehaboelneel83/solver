"""One workbook, one sheet per kind: kinds written before the kinds that refer to them, references
to the same kind (or round a circle) written in a second pass, and all or nothing."""
from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.workbook import plan
from app.main import app
from tests.test_api_relationships import auth_headers, domain_id, ensure_admin_seeded  # noqa: F401

client = TestClient(app)


def test_the_plan_orders_kinds_and_defers_what_cannot_wait():
    # truck -> depot -> region -> region (parent); a <-> b round a circle.
    refers = {1: {"parent": {1}}, 2: {"region": {1}}, 3: {"depot": {2}}, 4: {"b": {5}}, 5: {"a": {4}}}
    order, deferred = plan([3, 2, 1, 4, 5], refers)
    assert order.index(1) < order.index(2) < order.index(3)
    assert deferred[1] == {"parent"} and deferred[2] == set() and deferred[3] == set()
    assert (deferred[4], deferred[5]) in (({"b"}, set()), (set(), {"a"}))


def _kind(headers, domain, name, fields=()):
    made = client.post("/api/v1/entity-types", json={"domain_id": domain, "name": name, "role": "location"}, headers=headers)
    assert made.status_code == 201, made.text
    kind = made.json()["id"]
    for field in fields:
        got = client.post(f"/api/v1/entity-types/{kind}/attributes", json=field, headers=headers)
        assert got.status_code == 201, got.text
    return kind


def _book(sheets: dict[str, list[list]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for r in rows:
            sheet.append(r)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _upload(headers, domain, data: bytes, **params):
    return client.post(f"/api/v1/domains/{domain}/workbook", params=params, headers=headers,
                       files={"file": ("records.xlsx", data, "application/octet-stream")})


def _keys(headers, kind):
    got = client.get("/api/v1/entities", params={"entity_type_id": kind, "limit": 500}, headers=headers).json()
    return {e["key"]: e["attrs"] for e in got["items"]}


def _setup(headers, domain):
    region = _kind(headers, domain, "region")
    client.post(f"/api/v1/entity-types/{region}/attributes",
                json={"name": "parent", "data_type": "reference", "target_type_id": region}, headers=headers)
    depot = _kind(headers, domain, "depot", [{"name": "region", "data_type": "reference", "target_type_id": region, "required": True},
                                             {"name": "bays", "data_type": "integer"}])
    return region, depot


# Sheets in the wrong order, and a child region listed before its parent.
SHEETS = {
    "depot": [["key", "region", "bays"], ["D1", "cairo", 4], ["D2", "giza", 2]],
    "region": [["key", "label", "parent"], ["cairo", "Cairo", "egypt"], ["giza", "Giza", "egypt"], ["egypt", "Egypt", None]],
    "notes": [["anything"], ["ignored"]],
}


def test_a_workbook_is_written_in_order_with_parents_after_children(auth_headers, domain_id):  # noqa: F811
    region, depot = _setup(auth_headers, domain_id)

    dry = _upload(auth_headers, domain_id, _book(SHEETS), dry_run=True)
    assert dry.status_code == 200, dry.text
    report = dry.json()
    assert report["ok"] is True and report["kept"] is False and report["order"] == ["region", "depot"]
    assert report["ignored"] == ["notes"]
    assert {s["kind"]: s["second_pass"] for s in report["sheets"]} == {"region": ["parent"], "depot": []}
    assert _keys(auth_headers, region) == {}  # a dry run keeps nothing

    done = _upload(auth_headers, domain_id, _book(SHEETS)).json()
    assert done["kept"] is True and [s["written"] for s in done["sheets"]] == [3, 2]
    assert _keys(auth_headers, region)["cairo"] == {"parent": "egypt"}
    assert _keys(auth_headers, depot) == {"D1": {"region": "cairo", "bays": 4}, "D2": {"region": "giza", "bays": 2}}


def test_one_bad_cell_anywhere_keeps_nothing_and_names_its_sheet_row_and_column(auth_headers, domain_id):  # noqa: F811
    region, depot = _setup(auth_headers, domain_id)
    sheets = {**SHEETS, "depot": [["key", "region", "bays"], ["D1", "cairo", "four"], ["D2", "nowhere", 2]]}
    report = _upload(auth_headers, domain_id, _book(sheets)).json()
    assert report["ok"] is False and report["kept"] is False
    faults = {s["kind"]: s["faults"] for s in report["sheets"]}
    assert faults["region"] == []
    assert [(f["row"], f["column"]) for f in faults["depot"]][0] == (2, "bays")
    assert any(f["row"] == 3 for f in faults["depot"])  # a region that is in no sheet and not stored
    assert _keys(auth_headers, region) == {} and _keys(auth_headers, depot) == {}


def test_the_template_has_a_sheet_per_kind_in_writing_order_and_notes(auth_headers, domain_id):  # noqa: F811
    _setup(auth_headers, domain_id)
    got = client.get(f"/api/v1/domains/{domain_id}/workbook", headers=auth_headers)
    assert got.status_code == 200
    book = load_workbook(io.BytesIO(got.content))
    assert book.sheetnames == ["region", "depot", "about"]
    assert [c.value for c in book["depot"][1]] == ["key", "label", "sort_order", "active", "region", "bays"]
    notes = {(r[0], r[1]): r for r in book["about"].iter_rows(min_row=2, values_only=True)}
    assert notes[("depot", "region")][5] == "region"
    assert notes[("region", "parent")][6] == "second pass"
