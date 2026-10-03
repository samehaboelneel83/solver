"""Download templates and bulk uploads (queue R21): the right columns, every row checked
first, each fault by row and column, nothing written until the file is clean."""

from __future__ import annotations

import csv
import json
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.test_tenancy import tenants  # noqa: F401
from tests.test_v1_problem_run import db  # noqa: F401


@pytest.fixture
def shop(tenants, db):  # noqa: F811
    client, headers, domain = TestClient(app), tenants["a"], tenants["domain_a"]

    def call(method, path, body=None, ok=(200, 201, 204), **kw):
        response = client.request(method, path, json=body, headers=headers, **kw)
        if ok:
            assert response.status_code in ok, response.text
        return response

    def post(path, body):
        return call("POST", path, body).json()

    site = post("/api/v1/entity-types", {"domain_id": domain, "name": "site", "role": "location"})
    post(f"/api/v1/entity-types/{site['id']}/attributes", {"name": "capacity", "data_type": "integer", "required": True})
    post(f"/api/v1/entity-types/{site['id']}/attributes",
         {"name": "tier", "data_type": "enum", "enum_values": ["a", "b"], "default_value": "b"})
    post(f"/api/v1/entity-types/{site['id']}/attributes", {"name": "opens", "data_type": "time"})
    day = post("/api/v1/entity-types", {"domain_id": domain, "name": "day", "role": "time"})
    for k, key in enumerate(["mon", "tue"]):
        post("/api/v1/entities", {"entity_type_id": day["id"], "key": key, "sort_order": k})
    return {"call": call, "post": post, "domain": domain, "site": site, "day": day}


def _csv(rows):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return {"file": ("up.csv", out.getvalue().encode(), "text/csv")}


def _upload(shop, path, rows, **params):
    query = "&".join(f"{k}={str(v).lower()}" for k, v in params.items())
    return shop["call"]("POST", f"{path}?{query}", files=_csv(rows)).json()


def test_an_entity_template_names_every_column_and_says_their_types(shop):
    got = shop["call"]("GET", f"/api/v1/entity-types/{shop['site']['id']}/template?format=csv")
    assert got.headers["content-type"].startswith("text/csv")
    assert next(csv.reader(io.StringIO(got.text.lstrip("﻿")))) == ["key", "label", "sort_order", "active", "capacity", "tier", "opens"]
    book = load_workbook(io.BytesIO(shop["call"]("GET", f"/api/v1/entity-types/{shop['site']['id']}/template?format=xlsx").content))
    about = {row[0]: row for row in book["about"].iter_rows(min_row=2, values_only=True)}
    assert about["capacity"][1:3] == ("integer", "yes") and about["tier"][3] == "a, b"


def test_a_file_with_one_bad_row_writes_nothing_and_names_row_and_column(shop):
    path = f"/api/v1/entity-types/{shop['site']['id']}/upload"
    report = _upload(shop, path, [["key", "capacity", "tier", "opens"],
                                  ["north", "10", "a", "08:00"],
                                  ["south", "ten", "c", "8am"],
                                  ["", "3", "", ""]])
    assert not report["ok"] and report["written"] == 0
    assert {(f["row"], f["column"]) for f in report["faults"]} == {(3, "capacity"), (3, "tier"), (3, "opens"), (4, "key")}
    listed = shop["call"]("GET", f"/api/v1/entities?entity_type_id={shop['site']['id']}").json()
    assert listed["total"] == 0


def test_clean_only_writes_the_clean_rows_and_a_second_upload_updates_by_key(shop):
    path = f"/api/v1/entity-types/{shop['site']['id']}/upload"
    report = _upload(shop, path, [["key", "label", "capacity"], ["north", "North", "10"], ["south", "", "x"]], clean_only=True)
    assert report["written"] == 1 and report["skipped"] == 1
    # The database's own rule: capacity is required -- a trigger fault, still by row and column.
    report = _upload(shop, path, [["key", "tier"], ["east", "a"]])
    assert not report["ok"] and report["faults"][0]["row"] == 2 and report["faults"][0]["column"] == "capacity"
    report = _upload(shop, path, [["key", "capacity"], ["north", "25"]])
    assert report["ok"] and report["written"] == 1
    north = shop["call"]("GET", f"/api/v1/entities?entity_type_id={shop['site']['id']}").json()["items"][0]
    assert north["label"] == "North" and north["attrs"] == {"capacity": 25, "tier": "b"}


def test_a_dry_run_checks_and_writes_nothing(shop):
    path = f"/api/v1/entity-types/{shop['site']['id']}/upload"
    report = _upload(shop, path, [["key", "capacity"], ["north", "10"]], dry_run=True)
    assert report["ok"] and report["dry_run"] and report["written"] == 0
    assert shop["call"]("GET", f"/api/v1/entities?entity_type_id={shop['site']['id']}").json()["total"] == 0


def test_an_unknown_column_is_named_on_row_one(shop):
    report = _upload(shop, f"/api/v1/entity-types/{shop['site']['id']}/upload", [["key", "capacity", "colour"], ["n", "1", "red"]])
    assert report["faults"] == [{"row": 1, "column": "colour", "message": "is not a column of this template"}]



def test_add_fields_turns_unknown_columns_into_typed_fields(shop):
    site = shop["site"]["id"]
    path = f"/api/v1/entity-types/{site}/upload"
    rows = [["key", "capacity", "floors", "shift_start", "Bad Name"], ["n", "1", "3", "06:00", "x"], ["s", "2", "5", "18:00", "y"]]
    refused = _upload(shop, path, rows, add_fields=True, dry_run=True)
    assert [f["column"] for f in refused["faults"]] == ["Bad Name"]  # not a name a field may have
    attrs = {a["name"] for a in shop["call"]("GET", f"/api/v1/entity-types/{site}").json()["attributes"]}
    assert "floors" not in attrs  # a dry run leaves the type alone
    report = _upload(shop, path, [r[:4] for r in rows], add_fields=True)
    assert report["ok"] and report["written"] == 2, report
    kinds = {a["name"]: a["data_type"] for a in shop["call"]("GET", f"/api/v1/entity-types/{site}").json()["attributes"]}
    assert kinds["floors"] == "integer" and kinds["shift_start"] == "time"

def test_xlsx_round_trips_the_stored_rows(shop):
    path = f"/api/v1/entity-types/{shop['site']['id']}/upload"
    _upload(shop, path, [["key", "capacity", "opens"], ["north", "10", "08:00"], ["south", "4", ""]])
    content = shop["call"]("GET", f"/api/v1/entity-types/{shop['site']['id']}/template?format=xlsx&rows=true").content
    book = load_workbook(io.BytesIO(content))
    sheet = book["data"]
    sheet["E2"] = 11  # north's capacity, edited in a spreadsheet
    out = io.BytesIO()
    book.save(out)
    report = shop["call"]("POST", f"{path}", files={"file": ("sites.xlsx", out.getvalue(), "application/octet-stream")}).json()
    assert report["ok"] and report["written"] == 2
    items = shop["call"]("GET", f"/api/v1/entities?entity_type_id={shop['site']['id']}").json()["items"]
    assert {e["key"]: e["attrs"]["capacity"] for e in items} == {"north": 11, "south": 4}


def test_relationships_upload_by_keys(shop):
    post = shop["post"]
    for key in ("north", "south"):
        post("/api/v1/entities", {"entity_type_id": shop["site"]["id"], "key": key, "attrs": {"capacity": 1}})
    rel = post("/api/v1/relationship-types", {"domain_id": shop["domain"], "name": "open_on",
                                              "from_type_id": shop["site"]["id"], "to_type_id": shop["day"]["id"]})
    path = f"/api/v1/relationship-types/{rel['id']}/upload"
    bad = _upload(shop, path, [["from", "to"], ["north", "mon"], ["west", "tue"], ["north", "mon"]])
    assert {(f["row"], f["column"]) for f in bad["faults"]} == {(3, "from"), (4, "to")} and bad["written"] == 0
    good = _upload(shop, path, [["from", "to", "valid_from"], ["north", "mon", "2026-01-01"], ["south", "tue", ""]])
    assert good["ok"] and good["written"] == 2
    csv_text = shop["call"]("GET", f"/api/v1/relationship-types/{rel['id']}/template?rows=true").text.lstrip("﻿")
    assert list(csv.reader(io.StringIO(csv_text)))[1:] == [["north", "mon", "2026-01-01", ""], ["south", "tue", "", ""]]


def test_a_parameter_template_lists_every_cell_and_an_upload_fills_them(shop):
    post = shop["post"]
    for key in ("north", "south"):
        post("/api/v1/entities", {"entity_type_id": shop["site"]["id"], "key": key, "attrs": {"capacity": 1}})
    par = post("/api/v1/parameters", {"domain_id": shop["domain"], "name": "demand", "default_value": 1,
                                      "index_type_ids": [shop["site"]["id"], shop["day"]["id"]]})
    listed = shop["call"]("GET", f"/api/v1/parameters/{par['id']}/template").text.lstrip("﻿")
    assert list(csv.reader(io.StringIO(listed))) == [["site", "day", "value"], ["north", "mon", ""], ["north", "tue", ""],
                                                      ["south", "mon", ""], ["south", "tue", ""]]
    path = f"/api/v1/parameters/{par['id']}/upload"
    bad = _upload(shop, path, [["site", "day", "value"], ["north", "mon", "3"], ["north", "sun", "2"], ["south", "tue", "x"]])
    assert {(f["row"], f["column"]) for f in bad["faults"]} == {(3, "day"), (4, "value")}
    good = _upload(shop, path, [["site", "day", "value"], ["north", "mon", "3"], ["south", "tue", "1"], ["south", "mon", "2.5"]])
    assert good["ok"] and good["written"] == 3
    cells = shop["call"]("GET", f"/api/v1/parameters/{par['id']}/values").json()["cells"]
    # The default (1) is not stored: the grid stays sparse.
    assert sorted(c["value"] for c in cells) == [2.5, 3]


def test_a_place_cell_takes_a_latitude_and_longitude():
    from app.api.bulk import _parse

    assert _parse("geometry", "30.0444, 31.2357", None) == ({"type": "Point", "coordinates": [31.2357, 30.0444]}, None)
    assert _parse("geometry", '{"type": "Point", "coordinates": [31.2, 30.0]}', None)[0] == {"type": "Point", "coordinates": [31.2, 30.0]}
    assert "latitude and longitude" in _parse("geometry", "120, 31", None)[1]
    assert "such as 30.04, 31.23" in _parse("geometry", "near the river", None)[1]


@pytest.fixture
def teams(shop):
    """hospitals with a code, teams that link to one."""
    post = shop["post"]
    hospital = post("/api/v1/entity-types", {"domain_id": shop["domain"], "name": "hospital", "role": "location"})
    post(f"/api/v1/entity-types/{hospital['id']}/attributes", {"name": "hospital_code", "data_type": "text"})
    for key, label, code in (("imbaba", "Imbaba General Hospital", "H1"), ("haram", "Haram Hospital", "H3")):
        post("/api/v1/entities", {"entity_type_id": hospital["id"], "key": key, "label": label, "attrs": {"hospital_code": code}})
    team = post("/api/v1/entity-types", {"domain_id": shop["domain"], "name": "medical_team", "role": "agent"})
    post(f"/api/v1/entity-types/{team['id']}/attributes",
         {"name": "base_hospital", "data_type": "reference", "target_type_id": hospital["id"]})
    return {**shop, "team": team, "hospital": hospital}


def _upload_mapped(shop, type_id, rows, mapping=None, **params):
    query = "&".join(f"{k}={str(v).lower()}" for k, v in params.items())
    data = {"mapping": json.dumps(mapping)} if mapping is not None else None
    return shop["call"]("POST", f"/api/v1/entity-types/{type_id}/upload?{query}", files=_csv(rows), data=data).json()


def test_a_preview_reads_a_people_s_sheet_against_the_kind(teams):
    rows = [["team", "name", "hospital", "doctors"], ["T1", "Team 1", "H1", "2"], ["T2", "Team 2", "Haram Hospital", "1"]]
    got = teams["call"]("POST", f"/api/v1/entity-types/{teams['team']['id']}/upload/preview", files=_csv(rows)).json()
    assert got["rows"] == 2
    guesses = {c["name"]: c["suggestion"] for c in got["columns"]}
    assert guesses == {"team": "key", "name": "label", "hospital": "base_hospital", "doctors": None}
    assert {t["name"]: t["links_to"] for t in got["targets"]}["base_hospital"] == "hospital"
    assert got["columns"][0]["sample"] == ["T1", "T2"] and got["columns"][0]["unique"] is True


def test_a_mapped_upload_links_by_label_or_code_and_says_so(teams):
    rows = [["team", "name", "hospital", "notes"], ["T1", "Team 1", "H1", "x"], ["T2", "Team 2", "haram hospital", "y"]]
    got = _upload_mapped(teams, teams["team"]["id"], rows, {"team": "key", "name": "label", "hospital": "base_hospital", "notes": ""})
    assert got["ok"] and got["written"] == 2, got
    assert got["notes"] == ["base_hospital: 2 value(s) matched their record by its label or a code"]
    stored = teams["call"]("GET", f"/api/v1/entities?entity_type_id={teams['team']['id']}").json()["items"]
    assert {e["key"]: (e["label"], e["attrs"].get("base_hospital")) for e in stored} == {
        "T1": ("Team 1", "imbaba"), "T2": ("Team 2", "haram")}


def test_a_name_two_records_share_is_a_fault_not_a_guess(teams):
    teams["post"]("/api/v1/entities", {"entity_type_id": teams["hospital"]["id"], "key": "haram2", "label": "Haram Hospital"})
    rows = [["key", "base_hospital"], ["T9", "Haram Hospital"]]
    got = _upload_mapped(teams, teams["team"]["id"], rows)
    assert not got["ok"]
    assert got["faults"][0]["column"] == "base_hospital" and "matches 2 records" in got["faults"][0]["message"]


def test_two_columns_read_as_one_target_are_refused(teams):
    rows = [["a", "b"], ["T1", "T2"]]
    got = teams["call"]("POST", f"/api/v1/entity-types/{teams['team']['id']}/upload", files=_csv(rows),
                        data={"mapping": json.dumps({"a": "label", "b": "label"})}, ok=None)
    assert got.status_code == 422 and "'label'" in got.text


def test_several_columns_read_as_the_key_make_one_key(teams):
    """Benchmark, October 2026: a route and a stop were joined into one key in the spreadsheet first."""
    rows = [["route", "stop", "name"], ["R1", "3", "Market"], ["R1", "4", "Bridge"], ["R2", "", "Gap"]]
    got = _upload_mapped(teams, teams["team"]["id"], rows, {"route": "key", "stop": "key", "name": "label"}, clean_only=True)
    assert got["written"] == 2 and got["skipped"] == 1, got  # a missing part is a missing key, not "R2_"
    stored = teams["call"]("GET", f"/api/v1/entities?entity_type_id={teams['team']['id']}").json()["items"]
    assert sorted((e["key"], e["label"]) for e in stored) == [("R1_3", "Market"), ("R1_4", "Bridge")]
    # The key's own columns stay as fields, to link and forecast by (benchmark round 3).
    assert sorted((e["attrs"]["route"], e["attrs"]["stop"]) for e in stored) == [("R1", 3), ("R1", 4)]


def test_a_value_upload_reads_columns_as_mapped(shop):
    for key in ("north", "south"):
        shop["post"]("/api/v1/entities", {"entity_type_id": shop["site"]["id"], "key": key, "attrs": {"capacity": 1}})
    par = shop["post"]("/api/v1/parameters", {"domain_id": shop["domain"], "name": "demand", "default_value": 0,
                                              "index_type_ids": [shop["site"]["id"], shop["day"]["id"]]})
    rows = [["Weekday", "Site name", "score", "comment"], ["mon", "north", "4", "x"], ["tue", "south", "2", "y"]]
    got = shop["call"]("POST", f"/api/v1/parameters/{par['id']}/upload", files=_csv(rows),
                       data={"mapping": json.dumps({"Weekday": "day", "Site name": "site", "score": "value", "comment": ""})}).json()
    assert got["ok"] and got["written"] == 2, got
    cells = shop["call"]("GET", f"/api/v1/parameters/{par['id']}/values").json()["cells"]
    assert sorted(c["value"] for c in cells) == [2, 4]


def test_a_values_preview_does_not_guess_between_number_columns(shop):
    """Benchmark round 3: road_km went into lane_cost because it was the first number column."""
    par = shop["post"]("/api/v1/parameters", {"domain_id": shop["domain"], "name": "lane_cost", "default_value": 0,
                                              "index_type_ids": [shop["site"]["id"], shop["day"]["id"]]})

    def guess(header):
        rows = [header, ["north", "mon", "4", "7", "2"]]
        got = shop["call"]("POST", f"/api/v1/parameters/{par['id']}/upload/preview", files=_csv(rows)).json()
        return {c["name"]: c["suggestion"] for c in got["columns"]}

    # One shares a word with the data's name: that one.
    assert guess(["site", "day", "road_km", "cost_per_pallet", "drive_hours"])["cost_per_pallet"] == "value"
    # None does: the person chooses.
    assert "value" not in guess(["site", "day", "road_km", "fee", "drive_hours"]).values()


def test_a_refused_row_among_many_is_named_and_the_rest_kept_when_asked(shop):
    """Rows are written together, and one at a time only when the database refuses one: that row is
    still named, and with "write the clean rows" the others are kept."""
    site = shop["site"]["id"]
    path = f"/api/v1/entity-types/{site}/upload"
    rows = [["key", "capacity", "tier"]] + [[f"s{i}", str(i), "a"] for i in range(300)]
    rows[151][2] = "z"  # parses as text; the database refuses it as no tier choice
    refused = _upload(shop, path, rows)
    assert not refused["ok"] and refused["written"] == 0
    assert [(f["row"], f["column"]) for f in refused["faults"]] == [(152, "tier")]
    kept = _upload(shop, path, rows, clean_only=True)
    assert kept["written"] == 299 and kept["skipped"] == 1
    items = shop["call"]("GET", f"/api/v1/entities?entity_type_id={site}&limit=500").json()["items"]
    assert len(items) == 299


def test_an_import_finds_stored_records_by_key_whatever_its_case_and_says_what_it_would_make(shop):
    """Benchmark, October 2026: "h001" made a second H001, the name column was taken for the key,
    and "Check only" said every row was clean."""
    site = shop["site"]["id"]
    path = f"/api/v1/entity-types/{site}/upload"
    _upload(shop, path, [["key", "capacity"], ["H001", "5"], ["H002", "6"]])
    files = _csv([["name", "site_code", "capacity"], ["Tanta", "h001 ", "7"], ["Banha", "H002", "8"], ["Zagazig", "H003", "9"]])
    preview = shop["call"]("POST", f"{path}/preview", files=files).json()
    guess = {c["name"]: c for c in preview["columns"]}
    assert preview["existing"] == 2 and guess["site_code"]["suggestion"] == "key" and guess["site_code"]["matches_keys"] == 2
    mapping = {"name": "label", "site_code": "key", "capacity": "capacity"}
    checked = shop["call"]("POST", f"{path}?dry_run=true", files=_csv([["name", "site_code", "capacity"], ["Tanta", "h001 ", "7"],
                           ["Banha", "H002", "8"], ["Zagazig", "H003", "9"]]), data={"mapping": json.dumps(mapping)}).json()
    assert checked["ok"] and (checked["created"], checked["updated"]) == (1, 2)
    done = shop["call"]("POST", path, files=_csv([["name", "site_code", "capacity"], ["Tanta", "h001 ", "7"],
                        ["Banha", "H002", "8"], ["Zagazig", "H003", "9"]]), data={"mapping": json.dumps(mapping)}).json()
    assert done["written"] == 3
    items = {e["key"]: e for e in shop["call"]("GET", f"/api/v1/entities?entity_type_id={site}").json()["items"]}
    assert sorted(items) == ["H001", "H002", "H003"] and items["H001"]["attrs"]["capacity"] == 7 and items["H001"]["label"] == "Tanta"


def test_latitude_and_longitude_columns_put_the_records_on_the_map(shop):
    kind = shop["post"]("/api/v1/entity-types", {"domain_id": shop["domain"], "name": "call", "role": "other"})
    report = _upload(shop, f"/api/v1/entity-types/{kind['id']}/upload",
                     [["key", "lat", "lon"], ["c1", "30.05", "31.24"], ["c2", "", "31.2"]], add_fields=True)
    assert report["ok"], report
    assert any("1 rows placed on the map in location" in n for n in report["notes"])
    items = {e["key"]: e["attrs"] for e in shop["call"]("GET", f"/api/v1/entities?entity_type_id={kind['id']}").json()["items"]}
    assert items["c1"]["location"] == {"type": "Point", "coordinates": [31.24, 30.05]} and "location" not in items["c2"]


def test_a_sheet_never_turns_areas_into_points_and_left_out_columns_place_nothing(shop):
    """Benchmark round 5: a CSV onto the population zones (polygons), its lat/lon "left out", replaced
    every zone's polygon with a point, without a word."""
    kind = shop["post"]("/api/v1/entity-types", {"domain_id": shop["domain"], "name": "zone_area", "role": "location"})
    shop["post"](f"/api/v1/entity-types/{kind['id']}/attributes", {"name": "shape", "data_type": "geometry"})
    square = {"type": "Polygon", "coordinates": [[[31, 30], [31.1, 30], [31.1, 30.1], [31, 30.1], [31, 30]]]}
    shop["post"]("/api/v1/entities", {"entity_type_id": kind["id"], "key": "z1", "attrs": {"shape": square}})
    rows = [["key", "lat", "lon", "people"], ["z1", "30.05", "31.05", "900"]]
    left = _upload_mapped(shop, kind["id"], rows, {"key": "key", "lat": "", "lon": "", "people": "people"}, add_fields=True)
    assert left["ok"], left
    report = _upload_mapped(shop, kind["id"], rows, add_fields=True)
    assert report["ok"], report
    assert any("already have shapes in shape" in n for n in report["notes"])
    items = {e["key"]: e["attrs"] for e in shop["call"]("GET", f"/api/v1/entities?entity_type_id={kind['id']}").json()["items"]}
    assert items["z1"]["shape"] == square and items["z1"]["people"] == 900


def test_several_records_are_deleted_at_once_or_none(shop):
    site = shop["site"]["id"]
    _upload(shop, f"/api/v1/entity-types/{site}/upload", [["key", "capacity"], ["a", "1"], ["b", "2"], ["c", "3"]])
    items = {e["key"]: e["id"] for e in shop["call"]("GET", f"/api/v1/entities?entity_type_id={site}").json()["items"]}
    missing = shop["call"]("POST", "/api/v1/entities/delete", {"ids": [items["a"], 999999999]}, ok=None)
    assert missing.status_code == 404
    assert len(shop["call"]("GET", f"/api/v1/entities?entity_type_id={site}").json()["items"]) == 3  # none of them
    done = shop["call"]("POST", "/api/v1/entities/delete", {"ids": [items["a"], items["b"]]}).json()
    assert done == {"deleted": 2}
    assert [e["key"] for e in shop["call"]("GET", f"/api/v1/entities?entity_type_id={site}").json()["items"]] == ["c"]


def test_a_values_preview_reads_index_columns_by_name_or_by_their_keys(shop):
    for key in ("north", "south"):
        shop["post"]("/api/v1/entities", {"entity_type_id": shop["site"]["id"], "key": key, "attrs": {"capacity": 1}})
    par = shop["post"]("/api/v1/parameters", {"domain_id": shop["domain"], "name": "demand", "default_value": 0,
                                              "index_type_ids": [shop["site"]["id"], shop["day"]["id"]]})
    rows = [["Weekday", "Where", "score", "comment"], ["mon", "North", "4", "x"], ["tue", "south", "2", "y"]]
    got = shop["call"]("POST", f"/api/v1/parameters/{par['id']}/upload/preview", files=_csv(rows)).json()
    assert {c["name"]: c["suggestion"] for c in got["columns"]} == {"Weekday": "day", "Where": "site", "score": "value", "comment": None}
    assert [t["name"] for t in got["targets"]] == ["site", "day", "value"] and got["targets"][0]["links_to"] == "site"
