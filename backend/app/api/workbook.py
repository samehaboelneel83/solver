"""A whole domain's records as one Excel workbook: one sheet per kind of record.

    GET  /api/v1/domains/{id}/workbook          ?rows=true      the template (or the data)
    POST /api/v1/domains/{id}/workbook          ?dry_run=       (multipart `file`)

The one-kind upload (`bulk.py`) needs the kinds a file refers to already filled: regions before
depots, depots before trucks. A workbook brings them together, so this module puts them in order:

1. **Kinds first, then what refers to them.** A sheet is written after every sheet whose kind its
   reference fields name.
2. **References that cannot wait are deferred.** A reference to the same kind (a manager, a
   parent region) or one closing a circle between kinds can name a record that is further down
   the workbook. Those columns are left out of the first pass and written in a second pass,
   once every record exists.
3. **All or nothing.** Every sheet goes through the same checks as a one-kind upload -- each cell
   against its type, each row against the database's triggers in a savepoint -- inside one
   transaction, kept only if the whole workbook is clean; `dry_run` checks and keeps nothing.

4. **Then links and values.** A sheet `links <relationship>` holds that relationship's rows (from,
   to, ...) and `values <parameter>` that parameter's cells, written after every record, so they
   may name records added by this same workbook. A reference field's own relationship is not a
   sheet: it is written through its field, on the kind's sheet.

Sheets named after nothing here are reported as ignored; `about` (the template's notes) is skipped.
"""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app import audit
from app.api.bulk import (
    MAX_BYTES,
    MAX_ROWS,
    _entity_columns,
    _parameter_columns,
    _relationship_columns,
    entity_writer,
    parameter_writer,
    relationship_writer,
    write_rows,
)
from app.api.deps import get_current_user, requires
from app.core.db import get_db
from app.models.iam import UserAccount
from app.models.v1_domain import AttributeDef, Entity, EntityType, ParameterDef, RelationshipType

router = APIRouter(prefix="/api/v1", tags=["bulk"])

#: Excel's limit on a sheet's name.
SHEET_NAME = 31
NOTES = "about"


LINKS = "links "
VALUES = "values "


def sheet_name(prefix: str, name: str) -> str:
    return (prefix + name)[:SHEET_NAME]


def _links(db: Session, domain_id: int) -> list[RelationshipType]:
    """Relationship types written as rows: all but a reference field's mirror."""
    mirrors = select(AttributeDef.references_id).where(AttributeDef.references_id.is_not(None))
    return list(db.execute(select(RelationshipType).where(RelationshipType.domain_id == domain_id,
                                                          RelationshipType.id.not_in(mirrors))
                           .order_by(RelationshipType.name)).scalars())


def _values(db: Session, domain_id: int) -> list[ParameterDef]:
    return list(db.execute(select(ParameterDef).where(ParameterDef.domain_id == domain_id)
                           .order_by(ParameterDef.name)).scalars())


def _kinds(db: Session, domain_id: int) -> list[EntityType]:
    return list(db.execute(select(EntityType).where(EntityType.domain_id == domain_id, EntityType.is_abstract.is_(False))
                           .order_by(EntityType.name)).scalars())


def _refers(db: Session, kinds: list[EntityType]) -> dict[int, dict[str, set[int]]]:
    """Per kind: each reference field (own or inherited) and the concrete kinds it may name."""
    ids = [k.id for k in kinds]
    out: dict[int, dict[str, set[int]]] = {k.id: {} for k in kinds}
    rows = db.execute(text(
        "SELECT k.id AS kind, ad.name, ARRAY(SELECT unnest(entity_type_family(rt.to_type_id))) AS targets"
        "  FROM entity_type k"
        "  JOIN attribute_def ad ON ad.entity_type_id = ANY(entity_type_lineage(k.id)) AND ad.data_type::text = 'reference'"
        "  JOIN relationship_type rt ON rt.id = ad.references_id"
        " WHERE k.id = ANY(:ids)"), {"ids": ids}).mappings()
    for r in rows:
        out[r["kind"]][r["name"]] = set(r["targets"]) & set(ids)
    return out


def plan(kind_ids: list[int], refers: dict[int, dict[str, set[int]]]) -> tuple[list[int], dict[int, set[str]]]:
    """An order to write the kinds in, and per kind the reference fields to write in a second pass.

    Depth-first over "refers to": a kind is placed after the kinds its fields name. A field naming
    its own kind, or a kind still being placed (a circle), is deferred instead of followed."""
    order: list[int] = []
    deferred: dict[int, set[str]] = {k: set() for k in kind_ids}
    state: dict[int, str] = {}

    def visit(k: int) -> None:
        state[k] = "placing"
        for field, targets in sorted(refers.get(k, {}).items()):
            for t in sorted(targets):
                if t == k or state.get(t) == "placing":
                    deferred[k].add(field)
                elif t not in state:
                    visit(t)
        state[k] = "placed"
        order.append(k)

    for k in kind_ids:
        if k not in state:
            visit(k)
    # A field is written in the first pass only if every kind it names is already written by then.
    position = {k: i for i, k in enumerate(order)}
    for k in kind_ids:
        for field, targets in refers.get(k, {}).items():
            if any(position[t] >= position[k] for t in targets):
                deferred[k].add(field)
    return order, deferred


def _sheet_for(kind: EntityType, names: list[str]) -> str | None:
    for n in names:
        if n == kind.name or (len(kind.name) > SHEET_NAME and n == kind.name[:SHEET_NAME]):
            return n
    return None


def _read_book(upload: UploadFile) -> dict[str, tuple[list[str], list[list[Any]]]]:
    data = upload.file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"the file is over {MAX_BYTES // (1024 * 1024)} MB; split it")
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- any unreadable workbook is the same answer
        raise HTTPException(422, "the file is not a readable .xlsx workbook") from exc
    sheets: dict[str, tuple[list[str], list[list[Any]]]] = {}
    total = 0
    for sheet in book.worksheets:
        table = [list(r) for r in sheet.iter_rows(values_only=True)]
        while table and all(v is None or str(v).strip() == "" for v in table[-1]):
            table.pop()
        if not table:
            continue
        total += len(table) - 1
        sheets[sheet.title] = ([str(h).strip() if h is not None else "" for h in table[0]], table[1:])
    if total > MAX_ROWS:
        raise HTTPException(413, f"at most {MAX_ROWS:,} rows per workbook; split it")
    return sheets


def _project(header: list[str], rows: list[list[Any]], keep: list[str]) -> tuple[list[str], list[list[Any]]]:
    at = [header.index(h) for h in keep]
    return keep, [[r[i] if i < len(r) else None for i in at] for r in rows]


def _cell(value: Any) -> Any:
    """What a spreadsheet cell can hold: a shape or a list as its JSON (as the one-kind template
    writes it, and the upload reads back), a date and time without its zone, in UTC. Benchmark,
    October 2026: records made from a map layer made the whole download an HTTP 500."""
    if isinstance(value, (dict, list)):
        return json.dumps(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


@router.get("/domains/{domain_id}/workbook")
def workbook_template(domain_id: int, rows: bool = False, db: Session = Depends(get_db),
                      _: UserAccount = Depends(get_current_user)) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    kinds = _kinds(db, domain_id)
    if not kinds:
        raise HTTPException(404, "this domain has no kinds of record to fill")
    by_id = {k.id: k for k in kinds}
    refers = _refers(db, kinds)
    order, deferred = plan([k.id for k in kinds], refers)
    book = Workbook()
    book.remove(book.active)
    notes: list[list[Any]] = []
    for kid in order:
        kind = by_id[kid]
        columns, attributes = _entity_columns(db, kind)
        sheet = book.create_sheet(kind.name[:SHEET_NAME])
        sheet.append([c.name for c in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        if rows:
            for e in db.execute(select(Entity).where(Entity.entity_type_id == kid).order_by(Entity.sort_order, Entity.key)).scalars():
                sheet.append([_cell(v) for v in [e.key, e.label, e.sort_order, e.active,
                                                 *[(e.attrs or {}).get(a.name) for a in attributes]]])
        for c in columns:
            names = sorted(by_id[t].name for t in refers[kid].get(c.name, set()))
            notes.append([kind.name, c.name, c.kind, "yes" if c.required else "", ", ".join(c.values or []),
                          ", ".join(names), "second pass" if c.name in deferred[kid] else "", c.note])
    for rel in _links(db, domain_id):
        columns, attributes = _relationship_columns(db, rel)
        name = sheet_name(LINKS, rel.name)
        sheet = book.create_sheet(name)
        sheet.append([c.name for c in columns])
        if rows:
            for a, b, valid_from, valid_to, attrs in db.execute(text(
                    "SELECT ef.key, et.key, r.valid_from, r.valid_to, r.attrs FROM relationship r"
                    " JOIN entity ef ON ef.id = r.from_entity_id JOIN entity et ON et.id = r.to_entity_id"
                    " WHERE r.relationship_type_id = :t ORDER BY ef.key, et.key"), {"t": rel.id}).all():
                sheet.append([_cell(v) for v in [a, b, valid_from, valid_to, *[(attrs or {}).get(x.name) for x in attributes]]])
        notes += [[name, c.name, c.kind, "yes" if c.required else "", ", ".join(c.values or []), "", "after the records", c.note]
                  for c in columns]
    for parameter in _values(db, domain_id):
        columns, _heads = _parameter_columns(db, parameter)
        name = sheet_name(VALUES, parameter.name)
        sheet = book.create_sheet(name)
        sheet.append([c.name for c in columns])
        if rows:
            for keys, value, of in db.execute(text(
                    "SELECT ARRAY(SELECT e.key FROM unnest(pv.entity_ids) WITH ORDINALITY u(id, n)"
                    "              JOIN entity e ON e.id = u.id ORDER BY u.n),"
                    "       pv.value, (SELECT key FROM entity WHERE id = pv.value_entity_id)"
                    "  FROM parameter_value pv WHERE pv.parameter_def_id = :p ORDER BY 1"), {"p": parameter.id}).all():
                sheet.append([*keys, of if parameter.value_type_id else (float(value) if value is not None else None)])
        notes += [[name, c.name, c.kind, "yes" if c.required else "", "", "", "after the records", c.note] for c in columns]
    for sheet in book.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    about = book.create_sheet(NOTES)
    about.append(["sheet", "column", "type", "required", "allowed values", "refers to", "written", "note"])
    for cell in about[1]:
        cell.font = Font(bold=True)
    for n in notes:
        about.append(n)
    out = io.BytesIO()
    book.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="records.xlsx"'})


@router.post("/domains/{domain_id}/workbook")
def workbook_upload(domain_id: int, request: Request, file: UploadFile = File(...), dry_run: bool = Query(False),
                    db: Session = Depends(get_db), user: UserAccount = Depends(requires("domain.edit"))) -> dict[str, Any]:
    kinds = _kinds(db, domain_id)
    if not kinds:
        raise HTTPException(404, "this domain has no kinds of record to fill")
    by_id = {k.id: k for k in kinds}
    book = _read_book(file)
    matched = {k.id: s for k in kinds if (s := _sheet_for(k, list(book)))}
    refers = _refers(db, kinds)
    order, deferred = plan([k for k in matched], {k: {f: t & set(matched) for f, t in v.items()} for k, v in refers.items() if k in matched})

    reports: dict[int, dict[str, Any]] = {}
    for kid in order:
        header, rows = book[matched[kid]]
        first = [h for h in header if h not in deferred[kid]]
        h1, r1 = _project(header, rows, first)
        columns, write = entity_writer(db, by_id[kid], h1)
        written, faults, bad = write_rows(db, h1, r1, columns, write)
        reports[kid] = {"sheet": matched[kid], "kind": by_id[kid].name, "what": "records", "rows": len(rows),
                        "written": written, "faults": faults, "bad": bad, "second_pass": sorted(deferred[kid] & set(header))}
    for kid in order:
        later = [h for h in (deferred[kid] & set(book[matched[kid]][0]))]
        if not later:
            continue
        header, rows = book[matched[kid]]
        h2, r2 = _project(header, rows, ["key", *sorted(later)])
        columns, write = entity_writer(db, by_id[kid], h2)
        _, faults, bad = write_rows(db, h2, r2, columns, write)
        report = reports[kid]
        # A row refused in the first pass is refused again here only for want of itself: say it once.
        report["faults"] += [f for f in faults if f.row not in report["bad"]]
        report["bad"] |= bad
    done = [reports[k] for k in order]

    # Links and values last: their writers look records up when made, so every record above is there.
    for what, prefix, items, writer in (
            ("links", LINKS, _links(db, domain_id), lambda r, h: relationship_writer(db, r, h)),
            ("values", VALUES, _values(db, domain_id), lambda p, h: parameter_writer(db, p))):
        for item in items:
            name = sheet_name(prefix, item.name)
            if name not in book:
                continue
            header, rows = book[name]
            columns, write = writer(item, header)
            written, faults, _bad = write_rows(db, header, rows, columns, write)
            done.append({"sheet": name, "kind": item.name, "what": what, "rows": len(rows), "written": written,
                         "faults": faults, "second_pass": []})
    used = {r["sheet"] for r in done}
    ignored = [n for n in book if n not in used and n != NOTES]

    total_faults = sum(len(r["faults"]) for r in done)
    keep = not dry_run and total_faults == 0
    if keep:
        audit.write(db, user, request, action="bulk.workbook", object_type="domain", object_id=domain_id,
                    after={"sheets": [r["sheet"] for r in done], "rows": sum(r["written"] for r in done)})
        db.commit()
    else:
        db.rollback()
    sheets = [{"sheet": r["sheet"], "kind": r["kind"], "what": r["what"], "rows": r["rows"],
               "written": r["written"] if keep else 0, "second_pass": r["second_pass"],
               "faults": [f.model_dump() for f in r["faults"][:500]]} for r in done]
    return {"ok": total_faults == 0, "dry_run": dry_run, "kept": keep, "order": [r["sheet"] for r in done],
            "sheets": sheets, "ignored": ignored}
