"""Download templates and bulk uploads (queue R21), for any entity type, relationship
type and parameter.

    GET  /api/v1/entity-types/{id}/template        ?format=csv|xlsx&rows=
    POST /api/v1/entity-types/{id}/upload          ?clean_only=&dry_run=   (multipart `file`)
    GET  /api/v1/relationship-types/{id}/template
    POST /api/v1/relationship-types/{id}/upload
    GET  /api/v1/parameters/{id}/template
    POST /api/v1/parameters/{id}/upload

A template carries the right columns -- an entity's key, label, sort order,
active and each attribute (its own and inherited); an edge's from and to keys,
its dates and its attributes; a parameter's one column per index plus value --
and, as XLSX, an `about` sheet saying each column's type, whether it is
required and its allowed values. `rows=true` fills it with what is stored
(a parameter's template always lists every cell of its index), so a file
can go out, be edited and come back.

An upload checks every row -- each value against its column's type here,
then each row against the database's own triggers inside a savepoint -- and
reports each fault by row and column. Nothing is written unless the whole
file is clean, or (`clean_only`) the clean rows are; `dry_run` checks and
writes nothing. The triggers stay the one place the rules live: this module
parses cells so that "must be a number" can name its row, and leaves every
other rule to them.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from itertools import product
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import delete, func, select, text
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from app import audit
from app.api.deps import get_current_user, requires
from app.core.db import get_db
from app.crud.db_errors import translate_db_error
from app.models.iam import UserAccount
from app.models.v1_domain import (
    ATTRIBUTE_ORDER,
    AttributeDef,
    Entity,
    EntityType,
    ParameterDef,
    ParameterValue,
    Relationship,
    RelationshipType,
)

router = APIRouter(prefix="/api/v1", tags=["bulk"])

#: The most rows one file may carry, and the largest file read.
MAX_ROWS = 100_000
MAX_BYTES = 20 * 1024 * 1024
#: The most cells a parameter's full template lists.
MAX_TEMPLATE_CELLS = 100_000
_TIME_RE = re.compile(r"^([01]\d|2[0-3]):[0-5]\d(:[0-5]\d)?$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_LAT_LON = re.compile(r"^\s*(-?\d+(?:\.\d+)?)\s*[,; ]\s*(-?\d+(?:\.\d+)?)\s*$")


class Fault(BaseModel):
    row: int
    column: str | None
    message: str


class UploadReport(BaseModel):
    ok: bool
    rows: int
    written: int
    skipped: int
    dry_run: bool
    faults: list[Fault]
    #: What was done that the person did not spell out: links matched by a name or a code.
    notes: list[str] = []
    #: Of the rows written (or, on a check, that would be): records made new, and stored ones updated.
    created: int | None = None
    updated: int | None = None


# --- columns -----------------------------------------------------------------


class Column:
    def __init__(self, name: str, kind: str, *, required: bool = False, values: list[str] | None = None,
                 note: str = "", structural: bool = False) -> None:
        self.name, self.kind, self.required, self.values, self.note = name, kind, required, values, note
        # A structural column (a key, an end, an index) is in every file; an attribute's is
        # optional -- an update may carry only the columns it changes, and the trigger
        # still refuses a new row that lacks a required value.
        self.structural = structural


def _attributes(db: Session, *, entity_type_id: int | None = None, relationship_type_id: int | None = None) -> list[AttributeDef]:
    if relationship_type_id is not None:
        where = AttributeDef.relationship_type_id == relationship_type_id
    else:
        where = AttributeDef.entity_type_id == func.any(func.entity_type_lineage(entity_type_id))
    return list(db.execute(select(AttributeDef).where(where).order_by(*ATTRIBUTE_ORDER)).scalars())


def _attribute_columns(attributes: list[AttributeDef]) -> list[Column]:
    return [
        Column(a.name, a.data_type, required=a.required, values=list(a.enum_values or []) or None,
               note=(f"unit {a.unit}" if a.unit else "") + ("; key of an entity it refers to" if a.data_type == "reference" else ""))
        for a in attributes
    ]


def _parse(kind: str, raw: Any, values: list[str] | None) -> tuple[Any, str | None]:
    """A cell as the value its column's type stores, or why it is not one. `None` is empty."""
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None, None
    if isinstance(raw, str) and "\x00" in raw:
        return None, "holds a NUL character, which no value may"
    if kind in ("integer", "number"):
        if isinstance(raw, bool):
            return None, "must be a number"
        try:
            number = Decimal(str(raw).strip())
        except InvalidOperation:
            return None, f"must be {'a whole number' if kind == 'integer' else 'a number'}, not {raw!r}"
        if not number.is_finite():
            return None, "must be a finite number"
        if kind == "integer":
            if number % 1 != 0:
                return None, f"must be a whole number, not {raw!r}"
            return int(number), None
        return (int(number) if number % 1 == 0 else float(number)), None
    if kind == "boolean":
        if isinstance(raw, bool):
            return raw, None
        word = str(raw).strip().lower()
        if word in ("true", "yes", "y", "1"):
            return True, None
        if word in ("false", "no", "n", "0"):
            return False, None
        return None, f"must be true or false (yes or no), not {raw!r}"
    if kind == "date":
        if isinstance(raw, datetime):
            return raw.date().isoformat(), None
        if isinstance(raw, date):
            return raw.isoformat(), None
        word = str(raw).strip()
        try:
            if not _DATE_RE.match(word):
                raise ValueError(word)
            date.fromisoformat(word)
        except ValueError:
            return None, f"must be a date, YYYY-MM-DD, not {raw!r}"
        return word, None
    if kind == "time":
        if isinstance(raw, time):
            return raw.strftime("%H:%M"), None
        if isinstance(raw, datetime):
            return raw.strftime("%H:%M"), None
        word = str(raw).strip()
        if not _TIME_RE.match(word):
            return None, f"must be a time of day, HH:MM, not {raw!r}"
        return word, None
    if kind == "geometry":
        from app.spatial.geometry import validate_geometry

        # "30.04, 31.23" -- latitude first, as a map app copies it -- is a point.
        place = _LAT_LON.match(raw) if isinstance(raw, str) else None
        if place:
            lat, lon = float(place.group(1)), float(place.group(2))
            if abs(lat) > 90 or abs(lon) > 180:
                return None, f"{raw!r} is not a latitude and longitude"
            return {"type": "Point", "coordinates": [lon, lat]}, None
        try:
            shape = json.loads(raw) if isinstance(raw, str) else raw
        except ValueError:
            return None, ("must be a latitude and longitude, such as 30.04, 31.23, or GeoJSON, such as "
                          "{\"type\": \"Point\", \"coordinates\": [31.2, 30.0]}")
        fault = validate_geometry(shape)
        return (None, fault) if fault else (shape, None)
    # A spreadsheet stores 101 as a number; as a key or a text it is "101", not "101.0".
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    word = raw.isoformat() if isinstance(raw, (date, datetime)) else str(raw)
    word = word.strip() if kind != "text" else word
    if kind == "enum" and word not in (values or []):
        return None, f"must be one of {', '.join(values or [])}, not {word!r}"
    return word, None


# --- files ---------------------------------------------------------------------


def _read(upload: UploadFile) -> tuple[list[str], list[list[Any]]]:
    data = upload.file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"the file is over {MAX_BYTES // (1024 * 1024)} MB; split it")
    name = (upload.filename or "").lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 -- any unreadable workbook is the same answer
            raise HTTPException(422, "the file is not a readable .xlsx workbook") from exc
        sheet = book["data"] if "data" in book.sheetnames else book.worksheets[0]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
    else:
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "the file is neither UTF-8 CSV nor .xlsx") from exc
        table = [list(row) for row in csv.reader(io.StringIO(decoded))]
    while table and all(v is None or str(v).strip() == "" for v in table[-1]):
        table.pop()
    if not table:
        raise HTTPException(422, "the file is empty: the first row names the columns")
    if len(table) - 1 > MAX_ROWS:
        raise HTTPException(413, f"at most {MAX_ROWS:,} rows per file; split it")
    header = [str(h).strip() if h is not None else "" for h in table[0]]
    return header, table[1:]


def _check_header(header: list[str], columns: list[Column]) -> list[Fault]:
    known = {c.name for c in columns}
    faults = [Fault(row=1, column=h or f"column {i + 1}", message="is not a column of this template")
              for i, h in enumerate(header) if h not in known]
    faults += [Fault(row=1, column=h, message="appears more than once") for h in {h for h in header if header.count(h) > 1}]
    faults += [Fault(row=1, column=c.name, message="is required and missing")
               for c in columns if c.structural and c.name not in header]
    return faults


def _file(name: str, columns: list[Column], rows: list[list[Any]], fmt: str) -> Response:
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        book = Workbook()
        sheet = book.active
        sheet.title = "data"
        sheet.append([c.name for c in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([json.dumps(v) if isinstance(v, (dict, list)) else v for v in row])
        about = book.create_sheet("about")
        about.append(["column", "type", "required", "allowed values", "note"])
        for cell in about[1]:
            cell.font = Font(bold=True)
        for c in columns:
            about.append([c.name, c.kind, "yes" if c.required else "", ", ".join(c.values or []), c.note])
        out = io.BytesIO()
        book.save(out)
        return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
    out_text = io.StringIO()
    writer = csv.writer(out_text)
    writer.writerow([c.name for c in columns])
    for row in rows:
        writer.writerow(["" if v is None else json.dumps(v) if isinstance(v, (dict, list)) else
                         ("true" if v is True else "false" if v is False else v) for v in row])
    return Response("﻿" + out_text.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})


def write_rows(db: Session, header: list[str], rows: list[list[Any]], columns: list[Column], write_row,
               *, first_row: int = 2) -> tuple[int, list[Fault], set[int]]:
    """Parse each row and write it; nothing is committed here. Returns the rows written, the faults
    and the numbers of the rows that had any. Several files (a workbook's sheets) can share one
    transaction this way.

    A writer that can start again (`write_row.reset`) writes every row in one savepoint, which is
    fast; only when the database refuses a row is the lot undone and written again a row per
    savepoint, so that row is named and the others kept (benchmark, October 2026: 2,880 rows took
    18 s a row at a time)."""
    faults = _check_header(header, columns)
    if faults:
        return 0, faults, set(range(first_row, first_row + len(rows)))
    kinds = {c.name: c for c in columns}
    bad_rows: set[int] = set()
    parsed: list[tuple[int, dict[str, Any]]] = []
    for offset, raw in enumerate(rows):
        number = offset + first_row  # the header is row 1
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        values: dict[str, Any] = {}
        row_faults = []
        for i, name in enumerate(header):
            column = kinds[name]
            value, fault = _parse(column.kind, raw[i] if i < len(raw) else None, column.values)
            if fault:
                row_faults.append(Fault(row=number, column=name, message=fault))
            elif value is None and column.required:
                row_faults.append(Fault(row=number, column=name, message="is required"))
            else:
                values[name] = value
        if row_faults:
            faults += row_faults
            bad_rows.add(number)
        else:
            parsed.append((number, values))

    if hasattr(write_row, "reset"):
        together = db.begin_nested()
        try:
            written, problems = 0, []
            for number, values in parsed:
                problem = write_row(dict(values))
                if problem:
                    problems.append(Fault(row=number, column=problem[0], message=problem[1]))
                else:
                    written += 1
            db.flush()
            together.commit()
            faults += problems
            bad_rows.update(f.row for f in problems)
            return written, sorted(faults, key=lambda f: f.row or 0), bad_rows
        except DBAPIError:
            together.rollback()
            write_row.reset()

    written = 0
    for number, values in parsed:
        row_faults = []
        savepoint = db.begin_nested()
        try:
            problem = write_row(values)
            if problem:
                row_faults.append(Fault(row=number, column=problem[0], message=problem[1]))
                savepoint.rollback()
            else:
                db.flush()
                savepoint.commit()
        except DBAPIError as exc:
            savepoint.rollback()
            http = translate_db_error(exc, "bulk upload")
            detail = http.detail
            if isinstance(detail, list) and detail:
                loc = detail[0].get("loc") or []
                row_faults.append(Fault(row=number, column=str(loc[-1]) if len(loc) > 1 else None,
                                        message=str(detail[0].get("msg"))))
            else:
                row_faults.append(Fault(row=number, column=None, message=str(detail)))
        if row_faults:
            faults += row_faults
            bad_rows.add(number)
        else:
            written += 1
    return written, sorted(faults, key=lambda f: f.row or 0), bad_rows


def run_rows(db: Session, header: list[str], rows: list[list[Any]], columns: list[Column], write_row,
         clean_only: bool, dry_run: bool, *, user=None, request=None, audit_object: tuple[str, Any] | None = None,
         before_commit=None) -> UploadReport:
    """Parse each row, write it inside a savepoint, and keep or undo the lot."""
    header_faults = _check_header(header, columns)
    if header_faults:
        return UploadReport(ok=False, rows=len(rows), written=0, skipped=len(rows), dry_run=dry_run, faults=header_faults)
    written, faults, bad_rows = write_rows(db, header, rows, columns, write_row)
    keep = not dry_run and (not faults or clean_only)
    if keep:
        if user is not None and audit_object is not None and written:
            ot, oid = audit_object
            audit.write(
                db, user, request,
                action="bulk.upload",
                object_type=ot,
                object_id=oid,
                after={"rows": written, "faults": len(faults), "clean_only": clean_only},
            )
        if before_commit is not None:
            # In the same transaction as the rows: an import's lineage row (Epic UX, U-4).
            before_commit(written)
        db.commit()
    else:
        db.rollback()
    return UploadReport(ok=not faults, rows=len(rows), written=written if keep else 0,
                        skipped=len(bad_rows) if keep else len(rows), dry_run=dry_run, faults=faults[:1000])


# --- entity types ----------------------------------------------------------------


def _entity_columns(db: Session, entity_type: EntityType) -> tuple[list[Column], list[AttributeDef]]:
    attributes = _attributes(db, entity_type_id=entity_type.id)
    return [
        Column("key", "text", required=True, structural=True, note="unique within the type; how a model names the entity"),
        Column("label", "text", note="shown instead of the key"),
        Column("sort_order", "integer", note="lower first"),
        Column("active", "boolean", note="false leaves it out of every run"),
        *_attribute_columns(attributes),
    ], attributes


def _get(db: Session, model, id_: int, what: str):
    row = db.get(model, id_)
    if row is None:
        raise HTTPException(404, f"{what} not found")
    return row


class LinkMatcher:
    """A reference cell names its record by key -- or, as people's sheets do, by its label or a code
    (a text field whose values are unique among that kind, such as hospital_code). A value that is
    no key but matches exactly one record by one of those is written as that record's key."""

    def __init__(self, db: Session, target_type_id: int) -> None:
        self.db, self.target = db, target_type_id
        self.keys: set[str] | None = None
        self.aliases: dict[str, set[str]] = {}
        self.matched = 0

    def _load(self) -> None:
        rows = self.db.execute(text("SELECT key, label, attrs FROM entity WHERE entity_type_id = ANY (entity_type_family(:t))"),
                               {"t": self.target}).all()
        self.keys = {r.key for r in rows}
        by_field: dict[str, dict[str, set[str]]] = {}
        for r in rows:
            if r.label:
                self.aliases.setdefault(_fold(r.label), set()).add(r.key)
            for name, value in (r.attrs or {}).items():
                if isinstance(value, str) and value.strip():
                    by_field.setdefault(name, {}).setdefault(_fold(str(value)), set()).add(r.key)
        for values in by_field.values():
            # A code only counts where it tells the records apart.
            if all(len(k) == 1 for k in values.values()):
                for word, keys in values.items():
                    self.aliases.setdefault(word, set()).update(keys)
        self.aliases.update({_fold(k): {k} for k in self.keys if _fold(k) not in self.aliases})

    def resolve(self, value: str, written: set[str]) -> tuple[str, str | None]:
        """The key to store, or the value unchanged with why it could not be matched."""
        if value in written:
            return value, None
        if self.keys is None:
            self._load()
        if value in (self.keys or set()):
            return value, None
        found = self.aliases.get(_fold(value), set())
        if len(found) == 1:
            self.matched += 1
            return next(iter(found)), None
        if len(found) > 1:
            return value, f"{value!r} matches {len(found)} records by name or code ({', '.join(sorted(found)[:5])}); use the key"
        return value, None


def _fold(word: str) -> str:
    return " ".join(word.split()).casefold()


def entity_writer(db: Session, entity_type: EntityType, header: list[str], notes: list[str] | None = None):
    """The columns of an entity type and the function that writes one parsed row of it:
    shared by a file upload and a database import (Epic UX, U-4), so both meet the same checks.
    `notes`, when given, is told how many links were matched by a name or a code (see `finish`)."""
    columns, attributes = _entity_columns(db, entity_type)
    names = {a.name for a in attributes}
    seen: set[str] = set()
    matchers = {a.name: LinkMatcher(db, a.target_type_id) for a in attributes
                if a.data_type == "reference" and a.target_type_id is not None and a.name in header}
    same_kind = {a.name for a in attributes if a.data_type == "reference" and a.target_type_id == entity_type.id}
    # The kind's records, read once: a query per row also flushed every row before it, one by one.
    existing: dict[str, Entity] = {}
    # The same key in another case or spacing is the same record ("h001" is H001): matching it exactly
    # only made 22 hospitals twice (benchmark, October 2026).
    folded: dict[str, Entity | None] = {}
    counts = {"created": 0, "updated": 0}

    def load() -> None:
        existing.clear()
        folded.clear()
        counts.update(created=0, updated=0)
        existing.update({e.key: e for e in db.execute(select(Entity).where(Entity.entity_type_id == entity_type.id)).scalars()})
        for e in existing.values():
            folded[_fold(e.key)] = None if _fold(e.key) in folded else e  # two that fold alike match neither

    load()

    def write(values: dict[str, Any]):
        key = values["key"]
        if key in seen:
            return "key", f"{key!r} appears earlier in this file"
        seen.add(key)
        for name, matcher in matchers.items():
            if isinstance(values.get(name), str):
                values[name], problem = matcher.resolve(values[name], seen if name in same_kind else set())
                if problem:
                    return name, problem
        found = existing.get(key) or folded.get(_fold(key))
        if found is not None and found.key != key:
            if found.key in seen:
                return "key", f"{key!r} is {found.key!r} again, which appears earlier in this file"
            seen.add(found.key)
        counts["updated" if found is not None else "created"] += 1
        attrs = dict(found.attrs or {}) if found else {}
        for name in names & set(header):
            if values.get(name) is None:
                attrs.pop(name, None)
            else:
                attrs[name] = values[name]
        if found is None:
            found = Entity(entity_type_id=entity_type.id, key=key, attrs=attrs)
            db.add(found)
            existing[key] = found
            folded[_fold(key)] = found
        else:
            found.attrs = attrs
        if "label" in header:
            found.label = values.get("label")
        if values.get("sort_order") is not None:
            found.sort_order = values["sort_order"]
        if values.get("active") is not None:
            found.active = values["active"]
        return None

    def finish() -> None:
        if notes is None:
            return
        for name, matcher in matchers.items():
            if matcher.matched:
                notes.append(f"{name}: {matcher.matched} value(s) matched their record by its label or a code")

    def reset() -> None:
        """Ready to write the same rows again, after what was written was undone."""
        seen.clear()
        for matcher in matchers.values():
            matcher.matched = 0
        load()

    write.finish = finish  # type: ignore[attr-defined]
    write.reset = reset  # type: ignore[attr-defined]
    write.counts = counts  # type: ignore[attr-defined]
    return columns, write


@router.get("/entity-types/{entity_type_id}/template")
def entity_template(entity_type_id: int, format: str = Query("csv", pattern="^(csv|xlsx)$"), rows: bool = False,
                    db: Session = Depends(get_db), _: UserAccount = Depends(get_current_user)) -> Response:
    entity_type = _get(db, EntityType, entity_type_id, "entity type")
    columns, attributes = _entity_columns(db, entity_type)
    body: list[list[Any]] = []
    if rows:
        for e in db.execute(select(Entity).where(Entity.entity_type_id == entity_type.id)
                            .order_by(Entity.sort_order, Entity.key)).scalars():
            body.append([e.key, e.label, e.sort_order, e.active, *[(e.attrs or {}).get(a.name) for a in attributes]])
    return _file(entity_type.name, columns, body, format)


_FIELD_NAME = re.compile(r"^[a-z][a-z0-9_]{0,62}$")


def add_new_fields(db: Session, entity_type: EntityType, header: list[str], rows: list[list[Any]]) -> list[str]:
    """Columns the type has no field for become fields, typed from their own values (the same
    reading a spreadsheet start uses). Only names a field may have; anything else stays a fault.
    Inside the upload's transaction, so a dry run or a refused file leaves the type as it was."""
    from app.api.start import infer

    known = {c.name for c in _entity_columns(db, entity_type)[0]}
    added = []
    for i, name in enumerate(header):
        if name in known or not _FIELD_NAME.match(name or "") or header.count(name) > 1:
            continue
        values = [r[i] for r in rows if i < len(r) and r[i] is not None and str(r[i]).strip() != ""]
        data_type, choices = infer(values)
        db.add(AttributeDef(entity_type_id=entity_type.id, name=name, data_type=data_type, enum_values=choices))
        added.append(name)
    if added:
        db.flush()
    return added


def apply_mapping(header: list[str], rows: list[list[Any]], mapping: dict[str, str]) -> tuple[list[str], list[list[Any]]]:
    """The file's columns renamed as the person chose: `{"team": "key", "notes": ""}` reads the
    column `team` as the key and leaves `notes` out. Columns not named keep their names."""
    for target in mapping.values():
        if not isinstance(target, str):
            raise HTTPException(422, "a mapping names each file column's target as text, or \"\" to leave it out")
    keep = [i for i, h in enumerate(header) if mapping.get(h, h) != ""]
    renamed = [mapping.get(header[i], header[i]) for i in keep]
    # Several columns read as the key make one key, joined by "_": a depot and a day, a route and
    # a stop (benchmark, October 2026: composite keys were built in the spreadsheet first).
    parts = [i for i, h in zip(keep, renamed) if h == "key"] if renamed.count("key") > 1 else []
    twice = sorted({h for h in renamed if renamed.count(h) > 1 and not (parts and h == "key")})
    if twice:
        raise HTTPException(422, f"two columns are read as {', '.join(repr(t) for t in twice)}; choose one")
    if parts:
        keep = [i for i in keep if i not in parts[1:]]
        renamed = [mapping.get(header[i], header[i]) for i in keep]

    def cell(r: list[Any], i: int) -> Any:
        if parts and i == parts[0]:
            values = [str(r[j]).strip() if j < len(r) and r[j] is not None else "" for j in parts]
            return None if any(v == "" for v in values) else "_".join(values)
        return r[i] if i < len(r) else None

    return renamed, [[cell(r, i) for i in keep] for r in rows]


def _mapping(raw: str | None) -> dict[str, str]:
    if not raw:
        return {}
    try:
        found = json.loads(raw)
    except ValueError as exc:
        raise HTTPException(422, "mapping is not JSON") from exc
    if not isinstance(found, dict):
        raise HTTPException(422, "mapping is an object of file column -> target")
    return found


_LAT = ("lat", "latitude", "y_lat")
_LON = ("lon", "lng", "long", "longitude", "x_lon")


def with_places(db: Session, entity_type: EntityType, header: list[str], rows: list[list[Any]],
                notes: list[str]) -> tuple[list[str], list[list[Any]]]:
    """A sheet with a latitude and a longitude column gives each row a place: a point in the kind's
    geometry field (made, as `location`, when it has none). The columns themselves stay, to be kept
    or left out like any other. Benchmark, October 2026: the page said so, and the records got
    two number fields and no place."""
    norm = [re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_") for h in header]
    lat = next((i for i, n in enumerate(norm) if n in _LAT), None)
    lon = next((i for i, n in enumerate(norm) if n in _LON), None)
    if lat is None or lon is None:
        return header, rows
    geometry = next((a.name for a in _attributes(db, entity_type_id=entity_type.id) if a.data_type == "geometry"), None)
    if geometry is None:
        geometry = "location"
        db.add(AttributeDef(entity_type_id=entity_type.id, name=geometry, data_type="geometry"))
        db.flush()
    if geometry in header:
        return header, rows

    def point(row: list[Any]) -> str | None:
        try:
            y, x = float(row[lat]), float(row[lon])
        except (TypeError, ValueError, IndexError):
            return None
        return json.dumps({"type": "Point", "coordinates": [x, y]}) if -90 <= y <= 90 and -180 <= x <= 180 else None

    placed = [point(r) for r in rows]
    notes.append(f"{header[lat]} and {header[lon]}: {sum(p is not None for p in placed)} rows placed on the map in {geometry}")
    return [*header, geometry], [[*r, p] for r, p in zip(rows, placed)]


@router.post("/entity-types/{entity_type_id}/upload")
def entity_upload(entity_type_id: int, request: Request, file: UploadFile = File(...), clean_only: bool = False,
                  dry_run: bool = False, add_fields: bool = False, mapping: str | None = Form(None),
                  db: Session = Depends(get_db),
                  user: UserAccount = Depends(requires("domain.edit"))) -> UploadReport:
    """`add_fields`: a column the type has no field for becomes a new field instead of a fault.
    `mapping`: JSON, file column -> what it is read as (`key`, `label`, a field, or "" to leave it out)."""
    entity_type = _get(db, EntityType, entity_type_id, "entity type")
    if entity_type.is_abstract:
        raise HTTPException(422, f"{entity_type.name!r} is abstract and holds no entities of its own")
    header, rows = _read(file)
    notes: list[str] = []
    header, rows = with_places(db, entity_type, header, rows, notes)
    header, rows = apply_mapping(header, rows, _mapping(mapping))
    if add_fields:
        add_new_fields(db, entity_type, header, rows)
    columns, write = entity_writer(db, entity_type, header, notes)
    report = run_rows(db, header, rows, columns, write, clean_only, dry_run,
                      user=user, request=request, audit_object=("entity_type", entity_type_id))
    write.finish()
    report.notes = notes
    report.created, report.updated = write.counts["created"], write.counts["updated"]
    return report


class ColumnGuess(BaseModel):
    name: str
    #: Its first different values (benchmark re-test, October 2026: a sparse column read "bus lane
    #: planned, bus lane planned, bus lane planned" while most rows were empty).
    sample: list[str]
    #: How many rows have a value in it.
    filled: int | None = None
    #: Every non-empty value differs: it could be the key.
    unique: bool
    #: What it is read as unless the person says otherwise: key, label, a field, or None (new or left out).
    suggestion: str | None
    #: How many of its values are the key of a record already stored (case and spaces aside).
    matches_keys: int = 0


class Target(BaseModel):
    name: str
    kind: str
    required: bool
    #: For a link: the kind it names.
    links_to: str | None = None


class UploadPreview(BaseModel):
    rows: int
    #: Records of the kind already stored: a key column that matches none of them makes them all again.
    existing: int = 0
    columns: list[ColumnGuess]
    targets: list[Target]


_KEYISH = ("key", "code", "id", "ref", "no", "number")
_LABELISH = ("label", "name", "title", "description")


def suggest_mapping(header: list[str], rows: list[list[Any]], entity_type: EntityType, targets: list[Target],
                    existing: set[str] | None = None) -> dict[str, str | None]:
    """A first reading of a file's columns against a kind: the same name; a link named after the kind
    it links to; a key from a code-like, then any all-different column; a label from a name-like one."""
    norm = {h: re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_") for h in header}
    names = {t.name for t in targets}
    guess: dict[str, str | None] = {h: (norm[h] if norm[h] in names else None) for h in header}
    for t in targets:
        if t.links_to and t.name not in guess.values():
            for h in header:
                if guess[h] is None and norm[h] in (t.links_to, f"{t.links_to}_id", f"{t.links_to}_code", f"{t.links_to}_name"):
                    guess[h] = t.name
                    break

    def column(h: str) -> list[str]:
        i = header.index(h)
        return [str(r[i]).strip() for r in rows if i < len(r) and r[i] is not None and str(r[i]).strip() != ""]

    def unique(h: str) -> bool:
        values = column(h)
        return bool(values) and len(values) == len(rows) and len(set(values)) == len(values)

    kind = entity_type.name
    if "key" not in guess.values():
        free = [h for h in header if guess[h] is None and unique(h)]
        keyish = [h for h in free if norm[h] in _KEYISH or norm[h] in (kind, f"{kind}_id", f"{kind}_code")
                  or norm[h].endswith(("_code", "_id", "_key"))]
        named = [h for h in free if norm[h] in _LABELISH]
        # A kind with records already: the column holding their keys is the key, whatever it is called
        # (benchmark, October 2026: the warehouses' names were taken for the key and 12 came in twice).
        matching = sorted(((sum(_fold(v) in existing for v in column(h)), h) for h in free), reverse=True) if existing else []
        stored = [h for n, h in matching if n * 2 >= len(column(h)) and n > 0][:1]
        pick = (stored or keyish or free[:1] or named[:1])
        if pick:
            guess[pick[0]] = "key"
    if "label" not in guess.values():
        for h in header:
            if guess[h] is None and norm[h] in _LABELISH:
                guess[h] = "label"
                break
    return guess


@router.post("/entity-types/{entity_type_id}/upload/preview")
def entity_upload_preview(entity_type_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                          _: UserAccount = Depends(get_current_user)) -> UploadPreview:
    """A file's columns, a few values of each, and what each would be read as -- before anything is written."""
    entity_type = _get(db, EntityType, entity_type_id, "entity type")
    header, rows = _read(file)
    columns, attributes = _entity_columns(db, entity_type)
    kinds = {t.id: t.name for t in db.execute(select(EntityType).where(EntityType.domain_id == entity_type.domain_id)).scalars()}
    links = {a.name: kinds.get(a.target_type_id) for a in attributes if a.data_type == "reference"}
    targets = [Target(name=c.name, kind=c.kind, required=c.structural or c.required, links_to=links.get(c.name)) for c in columns]
    existing = {_fold(k) for k in db.execute(select(Entity.key).where(Entity.entity_type_id == entity_type.id)).scalars()}
    guess = suggest_mapping(header, rows, entity_type, targets, existing)
    out = []
    for i, h in enumerate(header):
        values = [str(r[i]).strip() for r in rows if i < len(r) and r[i] is not None and str(r[i]).strip() != ""]
        out.append(ColumnGuess(name=h, sample=list(dict.fromkeys(values))[:3], filled=len(values),
                               unique=bool(values) and len(values) == len(rows) and len(set(values)) == len(values),
                               suggestion=guess.get(h), matches_keys=sum(_fold(v) in existing for v in values)))
    return UploadPreview(rows=len(rows), existing=len(existing), columns=out, targets=targets)


# --- relationship types ------------------------------------------------------------


def _relationship_columns(db: Session, rel: RelationshipType) -> tuple[list[Column], list[AttributeDef]]:
    attributes = _attributes(db, relationship_type_id=rel.id)
    names = {t.id: t.name for t in db.execute(select(EntityType).where(EntityType.id.in_([rel.from_type_id, rel.to_type_id]))).scalars()}
    return [
        Column("from", "text", required=True, structural=True, note=f"key of a {names.get(rel.from_type_id)}"),
        Column("to", "text", required=True, structural=True, note=f"key of a {names.get(rel.to_type_id)}"),
        Column("valid_from", "date"),
        Column("valid_to", "date"),
        *_attribute_columns(attributes),
    ], attributes


def _mirror(db: Session, rel: RelationshipType) -> None:
    owner = db.execute(select(AttributeDef.name).where(AttributeDef.references_id == rel.id)).scalar_one_or_none()
    if owner is not None:
        raise HTTPException(409, f"{rel.name!r} is the reference attribute {owner!r}; upload it with the entities")


def _keyed(db: Session, type_id: int) -> dict[str, list[int]]:
    """Key -> the ids of the entities of that type's family bearing it."""
    found: dict[str, list[int]] = {}
    for eid, key in db.execute(text("SELECT id, key FROM entity WHERE entity_type_id = ANY (entity_type_family(:t))"),
                               {"t": type_id}).all():
        found.setdefault(key, []).append(eid)
    return found


def _resolve(index: dict[str, list[int]], key: str, what: str) -> tuple[int | None, str | None]:
    ids = index.get(key, [])
    if not ids:
        return None, f"no {what} has the key {key!r}"
    if len(ids) > 1:
        return None, f"{len(ids)} entities of the {what} family have the key {key!r}; it is ambiguous"
    return ids[0], None


@router.get("/relationship-types/{relationship_type_id}/template")
def relationship_template(relationship_type_id: int, format: str = Query("csv", pattern="^(csv|xlsx)$"), rows: bool = False,
                          db: Session = Depends(get_db), _: UserAccount = Depends(get_current_user)) -> Response:
    rel = _get(db, RelationshipType, relationship_type_id, "relationship type")
    columns, attributes = _relationship_columns(db, rel)
    body: list[list[Any]] = []
    if rows:
        for a, b, valid_from, valid_to, attrs in db.execute(text(
                "SELECT ef.key, et.key, r.valid_from, r.valid_to, r.attrs FROM relationship r "
                "JOIN entity ef ON ef.id = r.from_entity_id JOIN entity et ON et.id = r.to_entity_id "
                "WHERE r.relationship_type_id = :t ORDER BY ef.key, et.key"), {"t": rel.id}).all():
            body.append([a, b, valid_from.isoformat() if valid_from else None,
                         valid_to.isoformat() if valid_to else None,
                         *[(attrs or {}).get(x.name) for x in attributes]])
    return _file(rel.name, columns, body, format)


def relationship_writer(db: Session, rel: RelationshipType, header: list[str]):
    """The columns of a relationship type and the function that writes one parsed row of it:
    shared by a file upload and a database import, as `entity_writer`."""
    columns, attributes = _relationship_columns(db, rel)
    names = {a.name for a in attributes}
    froms, tos = _keyed(db, rel.from_type_id), _keyed(db, rel.to_type_id)
    seen: set[tuple[int, int]] = set()

    def write(values: dict[str, Any]):
        a, fault = _resolve(froms, values["from"], "from end")
        if fault:
            return "from", fault
        b, fault = _resolve(tos, values["to"], "to end")
        if fault:
            return "to", fault
        if (a, b) in seen:
            return "to", f"{values['from']} -> {values['to']} appears earlier in this file"
        seen.add((a, b))
        found = db.execute(select(Relationship).where(Relationship.relationship_type_id == rel.id,
                                                      Relationship.from_entity_id == a,
                                                      Relationship.to_entity_id == b)).scalar_one_or_none()
        attrs = dict(found.attrs or {}) if found else {}
        for name in names & set(header):
            if values.get(name) is None:
                attrs.pop(name, None)
            else:
                attrs[name] = values[name]
        if found is None:
            found = Relationship(relationship_type_id=rel.id, from_entity_id=a, to_entity_id=b, attrs=attrs)
            db.add(found)
        else:
            found.attrs = attrs
        for when in ("valid_from", "valid_to"):
            if when in header:
                setattr(found, when, date.fromisoformat(values[when]) if values.get(when) else None)
        return None

    return columns, write


@router.post("/relationship-types/{relationship_type_id}/upload")
def relationship_upload(relationship_type_id: int, request: Request, file: UploadFile = File(...),
                        clean_only: bool = False, dry_run: bool = False, db: Session = Depends(get_db),
                        user: UserAccount = Depends(requires("domain.edit"))) -> UploadReport:
    rel = _get(db, RelationshipType, relationship_type_id, "relationship type")
    _mirror(db, rel)
    header, rows = _read(file)
    columns, write = relationship_writer(db, rel, header)
    return run_rows(db, header, rows, columns, write, clean_only, dry_run,
                user=user, request=request, audit_object=("relationship_type", relationship_type_id))


# --- parameters ----------------------------------------------------------------------


def _parameter_columns(db: Session, parameter: ParameterDef) -> tuple[list[Column], list[str]]:
    names = {t.id: t.name for t in db.execute(select(EntityType).where(
        EntityType.id.in_([*parameter.index_type_ids, *([parameter.value_type_id] if parameter.value_type_id else [])]))).scalars()}
    index_names = [names.get(t, f"type_{t}") for t in parameter.index_type_ids]
    # A type that indexes twice (distance[site, site]) gets its position.
    heads = [f"{n}_{i + 1}" if index_names.count(n) > 1 else n for i, n in enumerate(index_names)]
    value = (Column("value", "text", note=f"key of a {names.get(parameter.value_type_id)}; empty for none")
             if parameter.value_type_id else
             Column("value", "number", note=f"empty for the default, {parameter.default_value}"
                    + (f"; unit {parameter.unit}" if parameter.unit else "")))
    return [*[Column(h, "text", required=True, structural=True, note=f"key of a {n}") for h, n in zip(heads, index_names)],
            value], heads


@router.get("/parameters/{parameter_id}/template")
def parameter_template(parameter_id: int, format: str = Query("csv", pattern="^(csv|xlsx)$"),
                       db: Session = Depends(get_db), _: UserAccount = Depends(get_current_user)) -> Response:
    """Every cell of the index, each with its stored value or empty."""
    parameter = _get(db, ParameterDef, parameter_id, "parameter")
    columns, _heads = _parameter_columns(db, parameter)
    axes = [list(db.execute(select(Entity.id, Entity.key).where(Entity.entity_type_id == t)
                            .order_by(Entity.sort_order, Entity.key)).all()) for t in parameter.index_type_ids]
    total = 1
    for axis in axes:
        total *= len(axis)
    if total > MAX_TEMPLATE_CELLS:
        raise HTTPException(413, f"{total:,} cells is more than one template lists ({MAX_TEMPLATE_CELLS:,})")
    stored = {tuple(r.entity_ids): r for r in db.execute(
        select(ParameterValue.entity_ids, ParameterValue.value, ParameterValue.value_entity_id)
        .where(ParameterValue.parameter_def_id == parameter.id)).all()}
    keys = dict(db.execute(select(Entity.id, Entity.key).where(
        Entity.id.in_([r.value_entity_id for r in stored.values() if r.value_entity_id]))).all()) if stored else {}
    body = []
    for combo in product(*axes):
        cell = stored.get(tuple(e.id for e in combo))
        value = None
        if cell is not None:
            value = keys.get(cell.value_entity_id) if cell.value_entity_id else (
                int(cell.value) if cell.value % 1 == 0 else float(cell.value))
        body.append([*[e.key for e in combo], value])
    return _file(parameter.name, columns, body, format)


def parameter_writer(db: Session, parameter: ParameterDef):
    """The columns of a parameter (one per index, then `value`) and the function that
    writes one parsed cell: shared by a file upload and a database import."""
    columns, heads = _parameter_columns(db, parameter)
    axes = [_keyed(db, t) for t in parameter.index_type_ids]
    of = _keyed(db, parameter.value_type_id) if parameter.value_type_id else None
    seen: set[tuple[int, ...]] = set()

    def write(values: dict[str, Any]):
        ids = []
        for head, axis in zip(heads, axes):
            eid, fault = _resolve(axis, values[head], head)
            if fault:
                return head, fault
            ids.append(eid)
        if tuple(ids) in seen:
            return heads[-1], "this cell appears earlier in this file"
        seen.add(tuple(ids))
        where = (ParameterValue.parameter_def_id == parameter.id) & (ParameterValue.entity_ids == ids)
        raw = values.get("value")
        if raw is None:
            db.execute(delete(ParameterValue).where(where))  # the default, or no entity
            return None
        if of is not None:
            eid, fault = _resolve(of, str(raw), "value")
            if fault:
                return "value", fault
            cell = {"value": None, "value_entity_id": eid}
        else:
            cell = {"value": Decimal(str(raw)), "value_entity_id": None}
        upsert = insert(ParameterValue).values(parameter_def_id=parameter.id, entity_ids=ids, **cell)
        db.execute(upsert.on_conflict_do_update(
            index_elements=[ParameterValue.parameter_def_id, ParameterValue.entity_ids],
            set_={"value": upsert.excluded.value, "value_entity_id": upsert.excluded.value_entity_id}))
        if of is None and cell["value"] == parameter.default_value:
            db.execute(delete(ParameterValue).where(where))  # sparse: the default is not stored
        return None

    return columns, write


@router.post("/parameters/{parameter_id}/upload")
def parameter_upload(parameter_id: int, request: Request, file: UploadFile = File(...), clean_only: bool = False,
                     dry_run: bool = False, mapping: str | None = Form(None), db: Session = Depends(get_db),
                     user: UserAccount = Depends(requires("domain.edit"))) -> UploadReport:
    """`mapping`: JSON, file column -> the index column or `value` it is read as, or "" to leave
    it out -- a sheet whose columns are `Crop`, `Soil type` and `score` (benchmark, October 2026)."""
    parameter = _get(db, ParameterDef, parameter_id, "parameter")
    header, rows = _read(file)
    header, rows = apply_mapping(header, rows, _mapping(mapping))
    columns, write = parameter_writer(db, parameter)
    return run_rows(db, header, rows, columns, write, clean_only, dry_run,
                user=user, request=request, audit_object=("parameter", parameter_id))


@router.post("/parameters/{parameter_id}/upload/preview")
def parameter_upload_preview(parameter_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                             _: UserAccount = Depends(get_current_user)) -> UploadPreview:
    """A values file's columns and what each would be read as: an index by its name or by holding
    that kind's keys, the value by its name or as the one column left of numbers (benchmark,
    October 2026: a sheet of `Crop`, `Soil type`, `score` had to be renamed before it went in)."""
    parameter = _get(db, ParameterDef, parameter_id, "parameter")
    header, rows = _read(file)
    columns, heads = _parameter_columns(db, parameter)
    names = {t.id: t.name for t in db.execute(select(EntityType).where(EntityType.id.in_(parameter.index_type_ids))).scalars()}
    targets = [Target(name=c.name, kind=c.kind, required=c.structural, links_to=names.get(t) if c.structural else None)
               for c, t in zip(columns, [*parameter.index_type_ids, None])]
    keys = [{_fold(k) for k in _keyed(db, t)} for t in parameter.index_type_ids]
    norm = {h: re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_") for h in header}

    def column(h: str) -> list[str]:
        i = header.index(h)
        return [str(r[i]).strip() for r in rows if i < len(r) and r[i] is not None and str(r[i]).strip() != ""]

    guess: dict[str, str | None] = {h: None for h in header}
    for head, kind in zip(heads, [names.get(t) for t in parameter.index_type_ids]):
        by_name = [h for h in header if guess[h] is None and norm[h] in (head, kind, f"{kind}_id", f"{kind}_key", f"{kind}_code")]
        if by_name:
            guess[by_name[0]] = head
    for head, known in zip(heads, keys):
        if head in guess.values():
            continue
        best = max(((sum(_fold(v) in known for v in column(h)), h) for h in header if guess[h] is None), default=(0, None))
        if best[0] > 0 and best[0] * 2 >= len(column(best[1])):
            guess[best[1]] = head
    if "value" not in guess.values():
        named = [h for h in header if guess[h] is None and norm[h] in ("value", "amount", "score", "number", parameter.name)]

        def numeric(h: str) -> bool:
            try:
                return bool(column(h)) and all(float(v) == float(v) for v in column(h))
            except ValueError:
                return False

        pick = named or [h for h in header if guess[h] is None and numeric(h)][:1]
        if pick:
            guess[pick[0]] = "value"
    out = [ColumnGuess(name=h, sample=list(dict.fromkeys(column(h)))[:3], filled=len(column(h)), unique=len(set(column(h))) == len(column(h)) == len(rows) and bool(rows),
                       suggestion=guess[h]) for h in header]
    return UploadPreview(rows=len(rows), columns=out, targets=targets)
