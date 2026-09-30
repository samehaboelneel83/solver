"""Start a problem from a spreadsheet (simplification plan, phase 3).

    POST /api/v1/domains/{id}/spreadsheet/propose   read a workbook, say what it would make
    POST /api/v1/domains/{id}/spreadsheet/import    make it, as the proposal (edited) says

A person with their data in Excel had to make every kind of record, every
field and every link by hand before a single row could be uploaded. Now the
workbook is read first and the structure is proposed from it:

- each sheet is a kind of record (a CSV is one sheet, named after the file);
- each column a field, its type read off the values -- whole numbers,
  numbers, yes/no, dates, one of a short list, or text;
- the first column whose values are all there and all different is the key
  (a column called id, key, code or name is preferred);
- a column whose every value is a key of another sheet is a link to it
  (a many-to-one link type, and one link per row), not a field.

The person corrects the proposal -- names, types, which column is the key --
and the import builds it through the same `plant_domain_seed` the ready
examples use, so it creates only what is missing: a kind or field that
already exists is reused, a record whose key exists is left as it is.
Every value is checked against its field's type before anything is written;
one bad cell refuses the whole import, with the sheet, row and column.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from app.api.bulk import MAX_BYTES, MAX_ROWS
from app.api.deps import requires
from app.api.validation import validate_name
from app.core.db import get_db
from app.models.iam import UserAccount

router = APIRouter(prefix="/api/v1", tags=["start"])

MAX_SHEETS = 20
ENUM_MAX_CHOICES = 8
SAMPLES = 3
TYPES = ("integer", "number", "boolean", "date", "enum", "text")
_TRUE = {"true", "yes", "y"}
_FALSE = {"false", "no", "n"}
_KEY_NAMES = ("id", "key", "code", "name")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")


# --- reading ------------------------------------------------------------------------


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _sheets(upload: UploadFile) -> list[tuple[str, list[str], list[list[Any]]]]:
    """(sheet name, header, rows) for every sheet with a header row."""
    data = upload.file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"the file is over {MAX_BYTES // (1024 * 1024)} MB; split it")
    filename = upload.filename or "data"
    tables: list[tuple[str, list[list[Any]]]] = []
    if filename.lower().endswith(".xlsx") or data[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 -- any unreadable workbook is the same answer
            raise HTTPException(422, "the file is not a readable .xlsx workbook") from exc
        for sheet in book.worksheets[:MAX_SHEETS]:
            tables.append((sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]))
    else:
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(422, "the file is neither UTF-8 CSV nor .xlsx") from exc
        stem = re.sub(r"\.[^.]*$", "", filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1])
        tables.append((stem or "data", [list(row) for row in csv.reader(io.StringIO(decoded))]))

    out = []
    for title, table in tables:
        table = [row for row in table if not all(_blank(v) for v in row)]
        if not table:
            continue
        if len(table) - 1 > MAX_ROWS:
            raise HTTPException(413, f"sheet {title!r} has over {MAX_ROWS:,} rows; split it")
        header = [str(h).strip() if not _blank(h) else "" for h in table[0]]
        width = len(header)
        while width and header[width - 1] == "":
            width -= 1
        header = header[:width]
        rows = [(row + [None] * width)[:width] for row in table[1:]]
        out.append((title, header, rows))
    if not out:
        raise HTTPException(422, "the file is empty: the first row of each sheet names its columns")
    return out


def to_name(text_: str) -> str:
    """"Hours per week" -> "hours_per_week": a name the platform takes."""
    name = re.sub(r"[^a-z0-9]+", "_", str(text_).strip().lower()).strip("_")
    if name and name[0].isdigit():
        name = f"n_{name}"
    return name or "field"


def _singular(name: str) -> str:
    if name.endswith("ies") and len(name) > 4:
        return name[:-3] + "y"
    if name.endswith("s") and not name.endswith("ss") and len(name) > 3:
        return name[:-1]
    return name


# --- values -------------------------------------------------------------------------


def _as_key(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def convert(value: Any, data_type: str, enum_values: list[str] | None = None) -> Any:
    """The cell as a value of `data_type`; ValueError when it is not one."""
    if isinstance(value, str):
        value = value.strip()
    if data_type == "boolean":
        if isinstance(value, bool):
            return value
        word = str(value).lower()
        if word in _TRUE or word == "1":
            return True
        if word in _FALSE or word == "0":
            return False
        raise ValueError("is not yes or no")
    if data_type == "integer":
        if isinstance(value, bool):
            raise ValueError("is not a whole number")
        if isinstance(value, int):
            return value
        if isinstance(value, float) and value.is_integer():
            return int(value)
        if isinstance(value, str) and re.fullmatch(r"-?\d+", value):
            return int(value)
        raise ValueError("is not a whole number")
    if data_type == "number":
        if isinstance(value, bool):
            raise ValueError("is not a number")
        if isinstance(value, (int, float)):
            return value
        try:
            number = float(str(value))
        except ValueError:
            raise ValueError("is not a number") from None
        return int(number) if number.is_integer() and re.fullmatch(r"-?\d+", str(value)) else number
    if data_type == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, str) and _DATE.fullmatch(value):
            try:
                return date.fromisoformat(value).isoformat()
            except ValueError:
                pass
        raise ValueError("is not a date (YYYY-MM-DD)")
    word = _as_key(value)
    if data_type == "enum" and word not in (enum_values or []):
        raise ValueError(f"is not one of {', '.join(enum_values or [])}")
    return word


def _fits(values: list[Any], data_type: str) -> bool:
    try:
        for value in values:
            convert(value, data_type)
    except ValueError:
        return False
    return True


def infer(values: list[Any]) -> tuple[str, list[str] | None]:
    """The type the non-blank values of a column read as."""
    if not values:
        return "text", None
    if all(isinstance(v, bool) or str(v).strip().lower() in _TRUE | _FALSE for v in values):
        return "boolean", None
    for data_type in ("integer", "number", "date"):
        if _fits(values, data_type):
            return data_type, None
    distinct = sorted({_as_key(v) for v in values})
    if 2 <= len(distinct) <= ENUM_MAX_CHOICES and len(values) >= 2 * len(distinct):
        return "enum", distinct
    return "text", None


# --- the proposal -------------------------------------------------------------------


class FieldPlan(BaseModel):
    column: str
    name: str
    data_type: str
    enum_values: list[str] | None = None
    samples: list[Any] = Field(default_factory=list)
    skip: bool = False


class LinkPlan(BaseModel):
    column: str
    name: str
    to: str  # the name of the kind whose keys the column holds
    skip: bool = False


class KindPlan(BaseModel):
    sheet: str
    name: str
    key: str | None  # the column; None numbers the rows
    rows: int = 0
    exists: bool = False
    fields: list[FieldPlan] = Field(default_factory=list)
    links: list[LinkPlan] = Field(default_factory=list)
    skip: bool = False


class Proposal(BaseModel):
    kinds: list[KindPlan]


def _existing_types(db: Session, domain_id: int) -> set[str]:
    if db.execute(text("SELECT 1 FROM domain WHERE id = :d"), {"d": domain_id}).first() is None:
        raise HTTPException(404, "domain not found")
    return set(db.execute(text("SELECT name FROM entity_type WHERE domain_id = :d"), {"d": domain_id}).scalars())


def _column(rows: list[list[Any]], i: int) -> list[Any]:
    return [row[i] for row in rows if not _blank(row[i])]


def propose(sheets: list[tuple[str, list[str], list[list[Any]]]], existing: set[str]) -> Proposal:
    kinds: list[KindPlan] = []
    keys: dict[str, set[str]] = {}
    taken: set[str] = set()
    for title, header, rows in sheets:
        name = _singular(to_name(title))
        while name in taken:
            name = f"{name}_2"
        taken.add(name)
        candidates = [
            i for i, h in enumerate(header)
            if h and rows and len(_column(rows, i)) == len(rows)
            and len({_as_key(v) for v in _column(rows, i)}) == len(rows)
        ]
        preferred = [i for i in candidates if to_name(header[i]) in _KEY_NAMES or to_name(header[i]).endswith("_id")]
        key = (preferred or candidates or [None])[0]
        kinds.append(KindPlan(sheet=title, name=name, key=header[key] if key is not None else None,
                              rows=len(rows), exists=name in existing))
        keys[name] = {_as_key(row[key]) for row in rows} if key is not None else set()

    for plan, (_, header, rows) in zip(kinds, sheets, strict=True):
        used: set[str] = set()
        for i, column in enumerate(header):
            if not column or column == plan.key:
                continue
            values = _column(rows, i)
            field_name = to_name(column)
            while field_name in used:
                field_name = f"{field_name}_2"
            used.add(field_name)
            target = next((other.name for other in kinds
                           if other is not plan and keys[other.name] and values
                           and {_as_key(v) for v in values} <= keys[other.name]), None)
            if target is not None:
                plan.links.append(LinkPlan(column=column, name=field_name, to=target))
                continue
            data_type, choices = infer(values)
            plan.fields.append(FieldPlan(column=column, name=field_name, data_type=data_type, enum_values=choices,
                                         samples=[_as_key(v) for v in values[:SAMPLES]]))
    return Proposal(kinds=kinds)


# --- building it --------------------------------------------------------------------


def _check(proposal: Proposal, sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> list[str]:
    """What is wrong with the proposal itself, before any value is read."""
    faults: list[str] = []
    kept = [k for k in proposal.kinds if not k.skip]
    names = [k.name for k in kept]
    for kind in kept:
        where = f"sheet {kind.sheet!r}"
        if kind.sheet not in sheets:
            faults.append(f"{where} is not in the file")
            continue
        header = sheets[kind.sheet][0]
        for label, name in [("the kind", kind.name), *[(f"field {f.column!r}", f.name) for f in kind.fields if not f.skip],
                            *[(f"link {link.column!r}", link.name) for link in kind.links if not link.skip]]:
            try:
                validate_name(name)
            except ValueError:
                faults.append(f"{where}: {label} is named {name!r}; a name is lower case letters, digits and _, starting with a letter")
        if names.count(kind.name) > 1:
            faults.append(f"{where}: another sheet also makes {kind.name!r}")
        if kind.key is not None and kind.key not in header:
            faults.append(f"{where}: the key column {kind.key!r} is not in the sheet")
        own = [f.name for f in kind.fields if not f.skip] + [link.name for link in kind.links if not link.skip]
        for dup in sorted({n for n in own if own.count(n) > 1}):
            faults.append(f"{where}: two columns are named {dup!r}")
        for item in [*kind.fields, *kind.links]:
            if not item.skip and item.column not in header:
                faults.append(f"{where}: column {item.column!r} is not in the sheet")
        for field_ in kind.fields:
            if not field_.skip and field_.data_type not in TYPES:
                faults.append(f"{where}: field {field_.name!r} has type {field_.data_type!r}, not one of {', '.join(TYPES)}")
            if not field_.skip and field_.data_type == "enum" and not field_.enum_values:
                faults.append(f"{where}: field {field_.name!r} is one of a list, but the list is empty")
        for link in kind.links:
            if not link.skip and link.to not in names:
                faults.append(f"{where}: link {link.name!r} goes to {link.to!r}, which this import does not make")
    return faults


def build_seed(proposal: Proposal, sheets: dict[str, tuple[list[str], list[list[Any]]]]) -> tuple[dict[str, Any], list[str]]:
    """The proposal and the rows as a template seed, and every cell that does not fit."""
    faults: list[str] = []
    kept = [k for k in proposal.kinds if not k.skip]
    seed: dict[str, Any] = {"entity_types": [], "relationship_types": [], "entities": [], "relationships": []}
    row_keys: dict[str, list[str]] = {}
    for kind in kept:
        header, rows = sheets[kind.sheet]
        at = {column: i for i, column in enumerate(header)}
        fields = [f for f in kind.fields if not f.skip]
        seed["entity_types"].append({"name": kind.name, "attributes": [
            {"name": f.name, "data_type": f.data_type, **({"enum_values": f.enum_values} if f.data_type == "enum" else {})}
            for f in fields]})
        keys: list[str] = []
        seen: set[str] = set()
        for n, row in enumerate(rows, start=2):
            key = f"{kind.name}_{n - 1}" if kind.key is None else _as_key(row[at[kind.key]]) if not _blank(row[at[kind.key]]) else ""
            if not key:
                faults.append(f"{kind.sheet!r} row {n}: the key {kind.key!r} is empty")
            elif key in seen:
                faults.append(f"{kind.sheet!r} row {n}: the key {key!r} is used twice")
            seen.add(key)
            keys.append(key)
            attrs: dict[str, Any] = {}
            for f in fields:
                cell = row[at[f.column]]
                if _blank(cell):
                    continue
                try:
                    attrs[f.name] = convert(cell, f.data_type, f.enum_values)
                except ValueError as exc:
                    faults.append(f"{kind.sheet!r} row {n}, column {f.column!r}: {_as_key(cell)!r} {exc}")
            seed["entities"].append({"type": kind.name, "key": key, "sort_order": n - 1, "attrs": attrs})
        row_keys[kind.name] = keys

    targets = {k.name: set(row_keys[k.name]) for k in kept}
    for kind in kept:
        header, rows = sheets[kind.sheet]
        at = {column: i for i, column in enumerate(header)}
        for link in (item for item in kind.links if not item.skip):
            seed["relationship_types"].append({"name": link.name, "from": kind.name, "to": link.to, "cardinality": "many_to_one"})
            for n, (row, key) in enumerate(zip(rows, row_keys[kind.name], strict=True), start=2):
                cell = row[at[link.column]]
                if _blank(cell):
                    continue
                other = _as_key(cell)
                if other not in targets[link.to]:
                    faults.append(f"{kind.sheet!r} row {n}, column {link.column!r}: {other!r} is not a {link.to}")
                    continue
                seed["relationships"].append({"type": link.name, "from": [kind.name, key], "to": [link.to, other]})
    return seed, faults


def _counts(db: Session, domain_id: int) -> dict[str, int]:
    return dict(db.execute(text(
        "SELECT 'kinds', count(*) FROM entity_type WHERE domain_id = :d"
        " UNION ALL SELECT 'fields', count(*) FROM attribute_def a JOIN entity_type t ON t.id = a.entity_type_id WHERE t.domain_id = :d"
        " UNION ALL SELECT 'records', count(*) FROM entity e JOIN entity_type t ON t.id = e.entity_type_id WHERE t.domain_id = :d"
        " UNION ALL SELECT 'link_types', count(*) FROM relationship_type WHERE domain_id = :d"
        " UNION ALL SELECT 'links', count(*) FROM relationship r JOIN relationship_type t ON t.id = r.relationship_type_id"
        "   WHERE t.domain_id = :d"), {"d": domain_id}).all())


@router.post("/domains/{domain_id}/spreadsheet/propose")
def propose_spreadsheet(domain_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                        user: UserAccount = Depends(requires("domain.edit"))) -> Proposal:
    existing = _existing_types(db, domain_id)
    return propose(_sheets(file), existing)


@router.post("/domains/{domain_id}/spreadsheet/import")
def import_spreadsheet(domain_id: int, file: UploadFile = File(...), proposal: str = Form(...),
                       db: Session = Depends(get_db), user: UserAccount = Depends(requires("domain.edit"))) -> dict[str, Any]:
    from app.seed import plant_domain_seed

    _existing_types(db, domain_id)
    try:
        plan = Proposal.model_validate(json.loads(proposal))
    except (ValueError, ValidationError) as exc:
        raise HTTPException(422, f"the proposal is not readable: {exc}") from exc
    sheets = {title: (header, rows) for title, header, rows in _sheets(file)}
    faults = _check(plan, sheets)
    if not faults:
        seed, faults = build_seed(plan, sheets)
    if faults:
        raise HTTPException(422, {"message": "Nothing was imported: fix these first.", "faults": faults[:50],
                                  "more": max(0, len(faults) - 50)})
    before = _counts(db, domain_id)
    try:
        plant_domain_seed(db, domain_id, seed)
        db.commit()
    except DBAPIError as exc:
        # A kind or field that already exists and disagrees with the sheet -- a field
        # of another type, a required one the sheet leaves out: the database says which.
        db.rollback()
        says = str(getattr(exc, "orig", exc)).splitlines()[0]
        raise HTTPException(422, {"message": "Nothing was imported: fix these first.", "faults": [says], "more": 0}) from exc
    after = _counts(db, domain_id)
    return {"made": {name: after[name] - before[name] for name in after}}
