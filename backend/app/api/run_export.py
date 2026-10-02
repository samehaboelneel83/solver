"""A run's answer as files (improvement plan 3.2).

    GET /api/v1/runs/{id}/export?format=xlsx|csv|geojson|html|pdf

What a field supervisor, a dispatcher or a spreadsheet downstream needs, from
the run's own frozen record (never today's data):

- **xlsx**: a Summary sheet (status, goal or goals in order, solver, time),
  one sheet per decision -- every chosen cell or non-zero amount, one column
  per index with the record's label beside its key; a two-index yes/no
  decision also as a grid -- and a Rules sheet: each rule, held or not, its
  slack, what bending it cost, where it fell short;
- **csv**: every decision in one long table: decision, keys, value;
- **geojson**: the answer map (`app.api.answer_map`), for GIS tools;
- **html**: a printable report -- summary, goals, the answer map drawn, the rules and each decision --
  that the browser prints (`print=true` opens the print dialog);
- **pdf**: that same report as a PDF file, laid out by WeasyPrint on A4 (user trial: "a real download").
  Nothing is fetched while laying it out: the base map is already drawn into the page as an image.
"""
from __future__ import annotations

import csv
import io
import json
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.api.answer_map import answer_map
from app.api.deps import get_current_user
from app.core.db import get_db
from app.models.iam import UserAccount

router = APIRouter(prefix="/api/v1", tags=["runs"])

SHEET_ROWS = 200_000


def _record(db: Session, run_id: int) -> dict[str, Any]:
    row = db.execute(
        text("SELECT r.id, r.status, r.objective, r.solver, r.params, r.error, r.finished_at, mv.ir, d.data,"
             "       sol.assignments, sol.amounts, s.name AS scenario, p.name AS problem, p.domain_id"
             "  FROM run r JOIN scenario s ON s.id = r.scenario_id JOIN model_version mv ON mv.id = r.model_version_id"
             "  JOIN problem p ON p.id = s.problem_id JOIN dataset d ON d.id = r.dataset_id"
             "  LEFT JOIN solution sol ON sol.run_id = r.id WHERE r.id = :r"),
        {"r": run_id},
    ).mappings().one_or_none()
    if row is None:
        raise HTTPException(404, "run not found")
    results = [dict(r) for r in db.execute(
        text("SELECT constraint_id, label, hard, satisfied, total_violation, penalty_paid, slack, violations"
             "  FROM constraint_result WHERE run_id = :r ORDER BY constraint_id"), {"r": run_id}).mappings()]
    return {**dict(row), "results": results}


def _labels(data: dict[str, Any]) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {s: dict(v) for s, v in (data.get("labels") or {}).items() if isinstance(v, dict)}
    for set_name, rows in (data.get("sets") or {}).items():
        for row in rows if isinstance(rows, list) else []:
            if isinstance(row, dict):
                label = row.get("label") or row.get("name")
                if isinstance(label, str) and label:
                    out.setdefault(set_name, {}).setdefault(str(row.get("id")), label)
    return out


def decision_rows(rec: dict[str, Any]) -> dict[str, tuple[list[str], list[list[Any]]]]:
    """Per decision: the index sets and its rows (keys..., value) -- chosen yes/no cells and non-zero amounts."""
    variables = (rec["ir"] or {}).get("variables") or {}
    out: dict[str, tuple[list[str], list[list[Any]]]] = {}
    for var, spec in variables.items():
        index = list(spec.get("index") or [])
        # A solution lists every non-zero cell under assignments, and the amount of a decision that is
        # not yes/no again under amounts: one row per cell, with its amount when it has one.
        cells: dict[tuple[str, ...], Any] = {tuple(map(str, cell)): 1 for cell in (rec["assignments"] or {}).get(var, [])}
        for entry in (rec["amounts"] or {}).get(var, []):
            key = tuple(map(str, entry.get("index") or []))
            value = float(entry.get("value") or 0)
            if abs(value) > 1e-9:
                cells[key] = value
            else:
                cells.pop(key, None)
        rows = [[*key, value] for key, value in cells.items()]
        out[var] = (index, sorted(rows, key=lambda r: [str(x) for x in r[:-1]]))
    return out


def _filename(rec: dict[str, Any], ext: str) -> str:
    stem = "".join(c if c.isalnum() or c in "-_" else "_" for c in f"{rec['problem']}-run-{rec['id']}")[:80]
    return f"{stem}.{ext}"


def to_xlsx(rec: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    params = rec["params"] or {}
    summary.append(["Problem", rec["problem"]])
    summary.append(["Scenario", rec["scenario"]])
    summary.append(["Run", rec["id"]])
    summary.append(["Status", rec["status"]])
    if params.get("objective_mode") == "lex" and params.get("objective_terms"):
        for n, term in enumerate(params["objective_terms"], start=1):
            summary.append([f"Goal {n}: {term['id']}", term["value"]])
    else:
        summary.append(["Goal", float(rec["objective"]) if rec["objective"] is not None else None])
    summary.append(["Solver", rec["solver"]])
    summary.append(["Finished", rec["finished_at"].isoformat() if rec["finished_at"] else None])
    if rec["error"]:
        summary.append(["Error", rec["error"]])
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    made_of = params.get("objective_breakdown")
    if made_of and made_of.get("terms"):
        # What the goal is made of, by term and by record (benchmark, October 2026).
        goal = wb.create_sheet("Goal")
        goal.append(["Goal term", "Weight", "Value", "Share", "Kind", "Record", "Record's part"])
        for cell in goal[1]:
            cell.font = Font(bold=True)
        for t in made_of["terms"]:
            goal.append([t["id"], t["weight"], t["value"], t["share"]])
            for r in t.get("records") or []:
                goal.append([t["id"], None, None, None, r["kind"], r["key"], r["value"]])
            if t.get("rest"):
                goal.append([t["id"], None, None, None, None, f"{t['rest']['records']} others", t["rest"]["value"]])
        if made_of.get("soft_rules"):
            goal.append(["soft rules broken", None, made_of["soft_rules"]])
    labels = _labels(rec["data"] or {})
    taken: set[str] = {"Summary", "Rules", "Goal"}
    for var, (index, rows) in decision_rows(rec).items():
        title = var[:31]
        while title in taken:
            title = (title[:28] + "_" + str(len(taken)))[:31]
        taken.add(title)
        ws = wb.create_sheet(title)
        header = []
        for s in index:
            header += [s, f"{s} name"]
        ws.append([*header, "value"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows[:SHEET_ROWS]:
            keys = row[:-1]
            line: list[Any] = []
            for s, k in zip(index, keys):
                line += [k, labels.get(s, {}).get(k, "")]
            ws.append([*line, row[-1]])
        # A yes/no decision over two sets also as a grid: rows of one, columns of the other.
        spec = ((rec["ir"] or {}).get("variables") or {}).get(var) or {}
        if len(index) == 2 and (spec.get("domain") or "binary") == "binary" and rows:
            grid = wb.create_sheet((title[:26] + " grid")[:31])
            taken.add(grid.title)
            cols = sorted({r[1] for r in rows})
            grid.append([f"{index[0]} \\ {index[1]}", *cols])
            on = {(r[0], r[1]) for r in rows}
            for a in sorted({r[0] for r in rows}):
                grid.append([a, *("✓" if (a, b) in on else "" for b in cols)])
    rules = wb.create_sheet("Rules")
    rules.append(["rule", "must hold", "held", "slack", "short by (total)", "cost of bending", "where it fell short"])
    for c in rules[1]:
        c.font = Font(bold=True)
    for r in rec["results"]:
        where = "; ".join(f"{' · '.join(map(str, v.get('index') or v.get('instance') or []))} by "
                          f"{float(v.get('by') if v.get('by') is not None else v.get('amount') or 0):g}"
                          for v in (r["violations"] or [])[:50])
        rules.append([r["constraint_id"], r["hard"], r["satisfied"],
                      float(r["slack"]) if r["slack"] is not None else None,
                      float(r["total_violation"] or 0), float(r["penalty_paid"] or 0), where])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(rec: dict[str, Any]) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    width = max((len(index) for index, _ in decision_rows(rec).values()), default=0)
    writer.writerow(["decision", *[f"key{i + 1}" for i in range(width)], "value"])
    for var, (index, rows) in decision_rows(rec).items():
        for row in rows:
            keys = row[:-1]
            writer.writerow([var, *keys, *([""] * (width - len(keys))), row[-1]])
    return out.getvalue()


def _goal_name(term: dict[str, Any] | None, gid: str) -> str:
    """As the run page says it: the author's note, else for an editor's placeholder id (o_1) what it adds up."""
    import re

    if term and str(term.get("note") or "").strip():
        return str(term["note"]).strip()
    if not re.fullmatch(r"o_?\d+", gid):
        return re.sub(r"^[a-z]_(?=[a-z])", "", gid).replace("_", " ")

    def read(node: Any, key: str) -> list[str]:
        if isinstance(node, list):
            return [x for n in node for x in read(n, key)]
        if not isinstance(node, dict):
            return []
        own = []
        if key == "data":
            own = [node["par"]] if isinstance(node.get("par"), str) else (
                [node["attr"]["name"]] if isinstance(node.get("attr"), dict) and node["attr"].get("name") else [])
        elif isinstance(node.get("var"), str):
            own = [node["var"]]
        return own + [x for v in node.values() for x in read(v, key)]

    expression = (term or {}).get("expression")
    data = list(dict.fromkeys(read(expression, "data")))
    if data:
        return " and ".join(d.replace("_", " ") for d in data)
    decided = list(dict.fromkeys(read(expression, "var")))
    return "total " + " and ".join(d.replace("_", " ") for d in decided) if decided else gid


STATUS_COLOUR = {"chosen": "#2563eb", "not_chosen": "#94a3b8", "short": "#dc2626", "place": "#64748b"}


#: The base maps the app offers without any setting (the same as the map's built-in choices).
BUILTIN_BASEMAPS = {
    "builtin-satellite": ("https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
                          "Imagery © Esri, Maxar, Earthstar Geographics"),
    "builtin-streets": ("https://tile.openstreetmap.org/{z}/{x}/{y}.png", "© OpenStreetMap contributors"),
}
TILE = 256
MAX_TILES = 64


def basemap_of(db: Session | None, domain_id: int | None, basemap: str | None) -> tuple[str, str] | None:
    """The tile address and credit line for a base map the person chose on the run's map: a built-in one, or
    one from the organization's tile index (the `spatial.tiles_index` setting). Never an address from the
    request itself: the server only fetches what the app already offers."""
    if not basemap or basemap == "none":
        return None
    if basemap in BUILTIN_BASEMAPS:
        if db is not None:
            from app.settings_resolve import resolve

            # An installation that keeps its sites inside (spatial.internet_basemaps = false) fetches no
            # internet imagery for a report either.
            if resolve(db, domain_id=domain_id)["spatial.internet_basemaps"].value is False:
                return None
        return BUILTIN_BASEMAPS[basemap]
    if db is None:
        return None
    try:
        from urllib.parse import urljoin

        from app.settings_resolve import resolve
        from app.spatial.terrain import candidates, fetch_bytes

        index_url = str(resolve(db, domain_id=domain_id)["spatial.tiles_index"].value or "").strip()
        if not index_url.startswith(("http://", "https://")):
            return None
        body = None
        for url in candidates(index_url):
            try:
                body = fetch_bytes(url, timeout=5)
                if body:
                    break
            except Exception:  # noqa: BLE001 -- an unreachable index means no base map, not a failed report
                continue
        for t in json.loads(body or b"[]"):
            if isinstance(t, dict) and str(t.get("id") or t.get("name") or "") == basemap and t.get("tiles"):
                return urljoin(index_url, str(t["tiles"][0])), str(t.get("attribution") or "")
    except Exception:  # noqa: BLE001
        return None
    return None


def _mercator(lon: float, lat: float, z: float) -> tuple[float, float]:
    import math

    size = TILE * 2 ** z
    s = math.sin(math.radians(max(-85.05112878, min(85.05112878, lat))))
    return (lon + 180) / 360 * size, (0.5 - math.log((1 + s) / (1 - s)) / (4 * math.pi)) * size


def _tile_picture(template: str, z: int, left: float, top: float, width: float, height: float,
                  out_w: int, out_h: int) -> str | None:
    """The tiles at zoom `z` covering world pixels [left, left+width) x [top, top+height), stitched, cut and
    scaled to out_w x out_h, as a JPEG data address; None when the tiles cannot be had."""
    import base64
    import math
    import urllib.request
    from concurrent.futures import ThreadPoolExecutor

    from PIL import Image

    from app.spatial.terrain import candidates

    x0, y0 = int(math.floor(left / TILE)), int(math.floor(top / TILE))
    x1, y1 = int(math.floor((left + width - 1e-6) / TILE)), int(math.floor((top + height - 1e-6) / TILE))
    tiles = [(x, y) for x in range(x0, x1 + 1) for y in range(y0, y1 + 1)]
    n = 2 ** z
    if not tiles or len(tiles) > MAX_TILES:
        return None

    def get(xy: tuple[int, int]):
        x, y = xy
        if not 0 <= y < n:
            return xy, None
        url = template.replace("{z}", str(z)).replace("{x}", str(x % n)).replace("{y}", str(y)).replace("{s}", "a")
        for address in candidates(url):
            try:
                request = urllib.request.Request(address, headers={"User-Agent": "ProblemSolver-report/1.0"})
                with urllib.request.urlopen(request, timeout=8) as response:  # noqa: S310 -- an address the app offers
                    return xy, Image.open(io.BytesIO(response.read())).convert("RGB")
            except Exception:  # noqa: BLE001 -- a missing tile leaves a blank square
                continue
        return xy, None

    with ThreadPoolExecutor(max_workers=8) as pool:
        got = dict(pool.map(get, tiles))
    if not any(got.values()):
        return None
    sheet = Image.new("RGB", ((x1 - x0 + 1) * TILE, (y1 - y0 + 1) * TILE), (241, 245, 249))
    for (x, y), image in got.items():
        if image is not None:
            sheet.paste(image.resize((TILE, TILE)), ((x - x0) * TILE, (y - y0) * TILE))
    cx, cy = left - x0 * TILE, top - y0 * TILE
    picture = sheet.crop((int(cx), int(cy), int(cx + width), int(cy + height))).resize((out_w, out_h))
    buffer = io.BytesIO()
    picture.save(buffer, "JPEG", quality=82)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()


def _svg_map(features: list[dict[str, Any]], width: int = 720, height: int = 440,
             basemap: tuple[str, str] | None = None) -> str:
    """The answer map as an SVG drawing in Web Mercator (as the app's map and every tile server draw it), fitted
    and centred, over the chosen base map when its tiles can be fetched."""
    import math
    from html import escape

    def coords(g: dict[str, Any]) -> list[list[float]]:
        t, c = g.get("type"), g.get("coordinates")
        if t == "Point":
            return [c]
        if t in ("LineString", "MultiPoint"):
            return list(c)
        if t in ("Polygon", "MultiLineString"):
            return [p for ring in c for p in ring]
        if t == "MultiPolygon":
            return [p for poly in c for ring in poly for p in ring]
        return []

    pts = [p for f in features for p in coords(f.get("geometry") or {})]
    if not pts:
        return ""
    world = [_mercator(p[0], p[1], 0) for p in pts]
    xs, ys = [w[0] for w in world], [w[1] for w in world]
    span = max(max(xs) - min(xs), 1e-9), max(max(ys) - min(ys), 1e-9)
    # The zoom that fits the places with a margin; never closer than street level.
    zf = min(math.log2(min((width - 60) / span[0], (height - 60) / span[1])), 18.0)
    k = 2 ** zf
    cx, cy = (min(xs) + max(xs)) / 2 * k, (min(ys) + max(ys)) / 2 * k
    left, top = cx - width / 2, cy - height / 2

    def at(p: list[float]) -> tuple[float, float]:
        x, y = _mercator(p[0], p[1], zf)
        return x - left, y - top

    def xy(p: list[float]) -> str:
        x, y = at(p)
        return f"{x:.1f},{y:.1f}"

    under, credit = "", ""
    if basemap:
        z = max(0, min(19, math.ceil(zf)))
        s = 2 ** (z - zf)  # world pixels at z per pixel of the drawing
        picture = _tile_picture(basemap[0], z, left * s, top * s, width * s, height * s, width, height)
        if picture:
            under = f'<image href="{picture}" x="0" y="0" width="{width}" height="{height}" preserveAspectRatio="none"/>'
            credit = basemap[1]
    order = {"place": 0, "not_chosen": 1, "chosen": 2, "short": 3}
    out, names = [], []
    on_map = bool(under)
    for f in sorted(features, key=lambda f: (f.get("geometry") or {}).get("type") not in ("LineString", "MultiLineString")
                    and order.get((f.get("properties") or {}).get("status"), 1) + 1 or 0):
        g, props = f.get("geometry") or {}, f.get("properties") or {}
        status = props.get("status")
        served = str(props.get("layer") or "").endswith("_served")
        colour = "#f97316" if served and status == "chosen" else STATUS_COLOUR.get(status, "#64748b")
        title = f"<title>{escape(str(props.get('title') or ''))}</title>"
        if g.get("type") == "Point":
            x, y = at(g["coordinates"])
            big = status in ("chosen", "short") and not served
            ring = "#ffffff" if on_map else "#0f172a"
            out.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{7 if big else 3.5}" fill="{colour}" fill-opacity="0.95" '
                       f'stroke="{ring}" stroke-width="{1.5 if big else 0.6}">{title}</circle>')
            if big and status == "chosen" and props.get("label"):
                names.append((x, y, str(props["label"])))
        elif g.get("type") in ("LineString", "MultiLineString"):
            # A road or a canal, chosen or not: a thicker line for a chosen one.
            width = 1.2 if served else 3.0 if status == "chosen" else 1.8
            for line in ([g["coordinates"]] if g["type"] == "LineString" else g["coordinates"]):
                out.append(f'<polyline points="{" ".join(xy(p) for p in line)}" fill="none" stroke="{colour}" '
                           f'stroke-width="{width}" stroke-opacity="{0.75 if served else 0.85}">{title}</polyline>')
        elif g.get("type") in ("Polygon", "MultiPolygon"):
            polys = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
            for poly in polys:
                out.append(f'<polygon points="{" ".join(xy(p) for p in poly[0])}" fill="{colour}" fill-opacity="0.25" '
                           f'stroke="{colour}" stroke-width="0.8">{title}</polygon>')
    # The chosen places by name, readable on any background.
    for x, y, label in names[:40]:
        # The white halo, then the name over it: two texts, not paint-order, which a PDF renderer may ignore.
        at = f'x="{x + 9:.1f}" y="{y + 4:.1f}" font-size="12" font-weight="700"'
        out.append(f'<text {at} fill="#ffffff" stroke="#ffffff" stroke-width="3">{escape(label)}</text>'
                   f'<text {at} fill="#0f172a">{escape(label)}</text>')
    if credit:
        at = f'x="{width - 6}" y="{height - 6}" font-size="9" text-anchor="end"'
        out.append(f'<text {at} fill="#ffffff" stroke="#ffffff" stroke-width="2.5">{escape(credit)}</text>'
                   f'<text {at} fill="#0f172a">{escape(credit)}</text>')
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" width="100%" '
            f'style="border:1px solid #cbd5e1;border-radius:6px;background:#f8fafc">{under}{"".join(out)}</svg>')


def _negated(term: dict[str, Any] | None) -> bool:
    """A goal counted backwards: a negative weight, or its expression multiplied by a negative constant."""
    if not term:
        return False
    if isinstance(term.get("weight"), (int, float)) and term["weight"] < 0:
        return True
    factors = (term.get("expression") or {}).get("mul") if isinstance(term.get("expression"), dict) else None
    return isinstance(factors, list) and any(isinstance(f, dict) and isinstance(f.get("const"), (int, float)) and f["const"] < 0
                                             for f in factors)


def to_html(rec: dict[str, Any], *, print_now: bool = False, basemap: tuple[str, str] | None = None) -> str:
    """A report a person prints or saves as PDF: what a field supervisor takes away."""
    from html import escape

    params = rec["params"] or {}
    labels = _labels(rec["data"] or {})
    e = lambda v: escape("" if v is None else str(v))  # noqa: E731
    parts = [f"<h1>{e(rec['problem'])}</h1>",
             f"<p class=sub>Run {e(rec['id'])} · scenario {e(rec['scenario'])} · {e(rec['status'])} · solver {e(rec['solver'])}"
             f"{' · finished ' + e(rec['finished_at'].strftime('%Y-%m-%d %H:%M')) if rec['finished_at'] else ''}</p>"]
    if params.get("objective_mode") == "lex" and params.get("objective_terms"):
        terms = {t.get("id"): t for t in ((rec["ir"] or {}).get("objective") or {}).get("terms") or [] if isinstance(t, dict)}
        # A cost among maximised goals is written "maximise minus the cost": print it as the cost.
        goals = ", then ".join(
            f"{e(_goal_name(terms.get(t['id']), t['id']))} <b>{(-1 if _negated(terms.get(t['id'])) else 1) * float(t['value']):,.6g}</b>"
            + (" (kept as low as it can go)" if _negated(terms.get(t["id"])) else "")
            for t in params["objective_terms"] if t.get("id") not in ("stay_close", "preferences"))
        parts.append(f"<p>Goals, in order: {goals}.</p>")
    elif rec["objective"] is not None:
        parts.append(f"<p>Goal: <b>{float(rec['objective']):g}</b></p>")
    made_of = params.get("objective_breakdown") or {}
    shown = [t for t in made_of.get("terms") or [] if t.get("id") not in ("stay_close", "preferences")]
    if len(shown) > 1 or (shown and len(shown[0].get("records") or []) > 1):
        rows = "".join(
            f"<tr><td>{e(t['id'].replace('_', ' '))}</td><td>{t['value']:,.6g}</td><td>{t['share'] * 100:.1f}%</td>"
            f"<td>{e(', '.join(str((labels.get(r['kind']) or {}).get(r['key'], r['key'])) + ' ' + format(r['value'], ',.6g') for r in (t.get('records') or [])[:5]))}</td></tr>"
            for t in shown)
        parts.append("<h2>What the goal is made of</h2><table><tr><th>Goal term</th><th>Value</th><th>Share</th><th>Most of it from</th></tr>"
                     f"{rows}</table>")
    if rec["error"]:
        parts.append(f"<p class=bad>{e(rec['error'])}</p>")
    if rec["assignments"] is not None or rec["amounts"] is not None:
        mapped = answer_map(rec["ir"] or {}, rec["data"] or {}, rec["assignments"], rec["amounts"], rec["results"], labels)
        drawing = _svg_map(mapped["features"], basemap=basemap)
        if drawing:
            parts.append("<h2>On the map</h2>" + drawing +
                         "<p class=key><span style='color:#2563eb'>●</span> chosen &nbsp; <span style='color:#94a3b8'>●</span> "
                         "not chosen &nbsp; <span style='color:#dc2626'>●</span> a rule fell short here"
                         + (" &nbsp; <span style='color:#f97316'>—</span> served from" if any(
                             str(f["properties"].get("layer", "")).endswith("_served") for f in mapped["features"]) else "")
                         + "</p>")
    if rec["results"]:
        rows = []
        for r in rec["results"]:
            short = "; ".join(f"{' · '.join(map(str, v.get('index') or v.get('instance') or []))} by "
                              f"{float(v.get('by') if v.get('by') is not None else v.get('amount') or 0):g}"
                              for v in (r["violations"] or [])[:12])
            slack = "" if r["slack"] is None else f"{float(r['slack']):g}"
            rows.append(f"<tr><td>{e(r['label'] or r['constraint_id'])}</td><td>{'must hold' if r['hard'] else 'preference'}</td>"
                        f"<td class={'ok' if r['satisfied'] else 'bad'}>{'held' if r['satisfied'] else 'not met'}</td>"
                        f"<td>{slack}</td><td>{e(short)}</td></tr>")
        parts.append("<h2>Rules</h2><table><tr><th>Rule</th><th>Kind</th><th>Result</th><th>Slack</th><th>Where it fell short</th></tr>"
                     + "".join(rows) + "</table>")
    for var, (index, rows) in decision_rows(rec).items():
        if not rows:
            continue
        head = "".join(f"<th>{e(s)}</th>" for s in index) + "<th>value</th>"
        body = "".join("<tr>" + "".join(f"<td>{e(labels.get(s, {}).get(k) or k)}</td>" for s, k in zip(index, row[:-1]))
                       + f"<td>{row[-1]:g}</td></tr>" for row in rows[:1000])
        more = f"<p class=sub>{len(rows) - 1000} more rows in the Excel export.</p>" if len(rows) > 1000 else ""
        parts.append(f"<h2>{e(var)} <span class=sub>({len(rows)})</span></h2><table><tr>{head}</tr>{body}</table>{more}")
    style = ("body{font:13px/1.45 system-ui,sans-serif;color:#0f172a;max-width:900px;margin:24px auto;padding:0 16px}"
             "h1{font-size:22px;margin:0}h2{font-size:16px;margin:24px 0 8px}.sub{color:#64748b}.bad{color:#b91c1c}.ok{color:#15803d}"
             ".key{font-size:12px;color:#475569}table{border-collapse:collapse;width:100%;font-size:12px}"
             "th,td{border:1px solid #e2e8f0;padding:3px 6px;text-align:left;vertical-align:top}th{background:#f1f5f9}"
             "@media print{body{margin:0}h2{break-after:avoid}tr{break-inside:avoid}}"
             "@page{size:A4;margin:14mm 12mm;@bottom-right{content:counter(page) ' / ' counter(pages);font-size:10px;color:#64748b}}"
             "svg{max-width:100%;height:auto}")
    script = "<script>window.addEventListener('load',()=>setTimeout(()=>window.print(),300))</script>" if print_now else ""
    return (f"<!doctype html><html><head><meta charset=utf-8><title>{e(rec['problem'])} — run {e(rec['id'])}</title>"
            f"<style>{style}</style></head><body>{''.join(parts)}{script}</body></html>")


def _no_fetch(url: str, *args: Any, **kwargs: Any) -> dict[str, Any]:
    """WeasyPrint asks for every resource a page names; the report names only `data:` ones, and anything
    else is refused rather than fetched from wherever the page says."""
    if url.startswith("data:"):
        from weasyprint.urls import default_url_fetcher

        return default_url_fetcher(url, *args, **kwargs)
    raise ValueError(f"the report fetches nothing: {url[:60]}")


def to_pdf(rec: dict[str, Any], *, basemap: tuple[str, str] | None = None) -> bytes:
    """The printable report as a PDF file (WeasyPrint, A4, numbered pages)."""
    try:
        from weasyprint import HTML
    except (ImportError, OSError) as exc:  # the library, or the Pango it draws text with, is not installed
        raise HTTPException(503, "PDF export is not installed on this server; use the printable report and "
                                 "save it as PDF from the browser") from exc
    return HTML(string=to_html(rec, basemap=basemap), url_fetcher=_no_fetch).write_pdf()


@router.get("/runs/{run_id}/export")
def export_run(run_id: int, format: str = Query("xlsx", pattern="^(xlsx|csv|geojson|html|pdf)$"),
               print: bool = Query(False),  # noqa: A002 -- the query word a person types
               basemap: str | None = Query(None, max_length=255),
               db: Session = Depends(get_db), user: UserAccount = Depends(get_current_user)) -> Response:
    rec = _record(db, run_id)
    if format == "html":
        # The base map the person had under the run's map, by its id (never an address from the request).
        under = basemap_of(db, rec.get("domain_id"), basemap)
        return Response(to_html(rec, print_now=print, basemap=under), media_type="text/html; charset=utf-8")
    if format == "pdf":
        under = basemap_of(db, rec.get("domain_id"), basemap)
        return Response(to_pdf(rec, basemap=under), media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{_filename(rec, "pdf")}"'})
    if rec["assignments"] is None and rec["amounts"] is None and format != "xlsx":
        raise HTTPException(409, f"run {run_id} has no answer to export (it is {rec['status']})")
    if format == "xlsx":
        return Response(to_xlsx(rec), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{_filename(rec, "xlsx")}"'})
    if format == "csv":
        return Response(to_csv(rec), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{_filename(rec, "csv")}"'})
    mapped = answer_map(rec["ir"] or {}, rec["data"] or {}, rec["assignments"], rec["amounts"], rec["results"],
                        _labels(rec["data"] or {}))
    body = {"type": "FeatureCollection", "name": f"{rec['problem']} run {rec['id']}", "layers": mapped["layers"],
            "features": mapped["features"]}
    return Response(json.dumps(body), media_type="application/geo+json",
                    headers={"Content-Disposition": f'attachment; filename="{_filename(rec, "geojson")}"'})
