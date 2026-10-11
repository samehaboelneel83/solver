"""Files a person attaches to the assistant: read once into plain tables.

A CSV, TSV, Excel workbook or JSON list becomes `{name, sheets: [{name,
columns, rows, total_rows, truncated}]}` -- what the model can read with the
`read_file` tool, and what a plan's `entities_from_file` /
`parameter_values_from_file` load from without copying rows through the chat.

Like the conversation, the parsed file is the client's: it comes back in the
upload's answer and is sent with each turn. Nothing is stored here.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time
from typing import Any

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
#: Rows read from a file the Assistant generated with run_python and loads straight from its working folder
#: (never sent through the chat): a layout's candidates run to tens of thousands (the camp-bed retest).
MAX_GENERATED_ROWS = __import__("app.core.limits", fromlist=["x"]).GENERATED_ROWS
_ROWS = __import__("contextvars").ContextVar("agent_file_rows", default=MAX_ROWS)
MAX_SHEETS = 10
MAX_COLUMNS = 60


class FileRefused(ValueError):
    pass


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    text = str(value).strip()
    if text == "":
        return None
    try:
        number = float(text.replace(",", "")) if text.replace(",", "").replace(".", "", 1).lstrip("-").isdigit() else None
    except ValueError:
        number = None
    if number is not None:
        return int(number) if number.is_integer() else number
    return text


def _table(name: str, raw_rows: list[list[Any]]) -> dict[str, Any] | None:
    rows = [r for r in raw_rows if any(c not in (None, "") for c in r)]
    if not rows:
        return None
    header = [str(c).strip() if c not in (None, "") else f"column_{i + 1}" for i, c in enumerate(rows[0])][:MAX_COLUMNS]
    seen: dict[str, int] = {}
    columns = []
    for col in header:
        seen[col] = seen.get(col, 0) + 1
        columns.append(col if seen[col] == 1 else f"{col}_{seen[col]}")
    body = [[_cell(c) for c in (r + [None] * len(columns))[: len(columns)]] for r in rows[1:]]
    limit = _ROWS.get()
    return {"name": name, "columns": columns, "rows": body[:limit],
            "total_rows": len(body), "truncated": len(body) > limit}


def parse(filename: str, data: bytes, max_rows: int | None = None, max_bytes: int | None = None) -> dict[str, Any]:
    """A data file as tables. `max_rows` / `max_bytes`: wider limits for a generated file read on the server."""
    token = _ROWS.set(max_rows) if max_rows else None
    try:
        return _parse(filename, data, max_bytes or MAX_BYTES)
    finally:
        if token is not None:
            _ROWS.reset(token)


def _parse(filename: str, data: bytes, max_bytes: int) -> dict[str, Any]:
    if len(data) > max_bytes:
        raise FileRefused(f"the file is larger than {max_bytes // (1024 * 1024)} MB")
    lower = filename.lower()
    sheets: list[dict[str, Any]] = []
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 -- any unreadable workbook is the same refusal
            raise FileRefused(f"not a readable Excel workbook: {exc}") from exc
        for ws in book.worksheets[:MAX_SHEETS]:
            rows = [list(r) for _, r in zip(range(_ROWS.get() + 2), ws.iter_rows(values_only=True))]
            table = _table(ws.title, rows)
            if table:
                sheets.append(table)
    elif lower.endswith(".json"):
        try:
            doc = json.loads(data.decode("utf-8-sig"))
        except ValueError as exc:
            raise FileRefused(f"not valid JSON: {exc}") from exc
        items = doc if isinstance(doc, list) else next((v for v in doc.values() if isinstance(v, list)), None) \
            if isinstance(doc, dict) else None
        if not items or not all(isinstance(i, dict) for i in items):
            raise FileRefused("a JSON file must be a list of objects (one per row)")
        columns = list(dict.fromkeys(k for item in items for k in item))[:MAX_COLUMNS]
        table = _table(filename.rsplit(".", 1)[0], [columns] + [[i.get(c) for c in columns] for i in items])
        if table:
            sheets.append(table)
    elif lower.endswith((".csv", ".tsv", ".txt")):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1256", errors="replace")  # Arabic Windows exports
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if lower.endswith(".tsv") else csv.excel
        rows = [r for _, r in zip(range(_ROWS.get() + 2), csv.reader(io.StringIO(text), dialect))]
        table = _table(filename.rsplit(".", 1)[0], rows)
        if table:
            sheets.append(table)
    else:
        raise FileRefused("attach data (.csv, .tsv, .txt, .xlsx, .json) or a map file (.dxf, .geojson, .kml, .kmz, "
                          ".gpx, a shapefile as a .zip of its .shp, .shx, .dbf and .prj, .gpkg); a .dwg must be saved "
                          "as DXF from your CAD program first")
    if not sheets:
        raise FileRefused("the file has no rows")
    return {"name": filename, "sheets": sheets}


# -- tables pasted into a message ---------------------------------------------------

_TIGHT = re.compile(r",(?! )")


def _row_like(line: str) -> int:
    """How many cells a line holds if it reads as a CSV row (commas with no space after them), else 0."""
    line = line.strip()
    if not line or ", " in line or not _TIGHT.search(line):
        return 0
    return line.count(",") + 1


def _slug(heading: str, taken: set[str]) -> str:
    named = re.search(r"([\w-]+\.csv)\b", heading)
    if named:
        base = named.group(1)[:-4].lower()
    else:
        words = re.findall(r"[A-Za-z0-9]+", re.split(r"[(:]", heading, 1)[0].lower())[:4]
        base = "_".join(words) or "table"
    name, n = f"{base}.csv", 2
    while name in taken:
        name, n = f"{base}_{n}.csv", n + 1
    return name


def pasted_tables(text: str) -> list[tuple[str, str]]:
    """The tables a person pasted into a message (a heading, a header row, rows with as many cells), as
    (file name, CSV text). The phase 1 evaluation, 11 October 2026: a 200-row table pasted into the request was
    copied into the plan by hand -- 34,000-character replies, cut at the length limit, in every run."""
    lines = (text or "").splitlines()
    found: list[tuple[str, str]] = []
    taken: set[str] = set()
    i = 0
    while i < len(lines):
        width = _row_like(lines[i])
        if width < 2 or any(re.fullmatch(r"-?[\d.]+", c.strip()) for c in lines[i].split(",")):
            i += 1
            continue
        j = i + 1
        while j < len(lines) and _row_like(lines[j]) == width:
            j += 1
        if j - i >= 2:
            heading = next((lines[k].strip() for k in range(i - 1, max(-1, i - 3), -1) if lines[k].strip()), "")
            name = _slug(heading, taken)
            taken.add(name)
            found.append((name, "\n".join(l.strip() for l in lines[i:j]) + "\n"))
        i = max(j, i + 1)
    return found


# -- for the model ------------------------------------------------------------------


SMALL_SHEET_ROWS = 25
# Or this few cells: a narrow sheet of 30 places (6 columns) was shown 3 rows, and the Assistant asked whether
# "Abu Nomros" was in it (row 15; fibre test, October 2026).
SMALL_SHEET_CELLS = 480
# A bigger sheet's first column (its names) is listed up to this many values.
NAMES_LISTED = 200


def outline(files: list[dict[str, Any]]) -> str:
    """What the system prompt says about the attached files: names, columns, sizes, a taste of the rows."""
    if not files:
        return ""
    lines = ["ATTACHED FILES (read more with read_file; count, filter and add up with query_file; load rows with entities_from_file / parameter_values_from_file):"]
    for f in files:
        spatial = f.get("spatial")
        if spatial:
            if spatial.get("local_metres"):
                where = ("LOCAL METRES (a drawing with no coordinate system: do NOT ask the user for one, never use "
                         "EPSG:4326): x_m/y_m (centre), min_x_m/min_y_m, width_m/height_m are metres from the drawing's "
                         "lower-left corner; area_m2 and length_m are exact; \"geometry\" is the shape for records; "
                         "\"shape_m\" is the same shape in those metres as WKT (for run_python: shapely.wkt.loads). "
                         "Only if it must line up with other map data, ask where it is and use place_file. To import it "
                         "as map data: placement " + json.dumps({k: v for k, v in (spatial.get("placement") or {}).items()
                                                                 if k != "name"}))
            elif spatial.get("placed"):
                where = (f'placed in EPSG:{(spatial.get("placement") or {}).get("code")} '
                         f'({(spatial.get("placement") or {}).get("name")}): lon/lat in degrees, area_m2 and length_m in metres, '
                         f'"geometry" is the shape a geometry attribute takes')
            else:
                options = "; ".join(f'EPSG:{c["code"]} {c["name"]} ({c["reason"]}, lands at {c["lands_at"]})'
                                    for c in spatial.get("candidates") or [])
                where = ("NOT PLACED YET: x/y are in the file's own units. Before using its positions, ask the user "
                         f"which coordinate system it is in (likely: {options or 'none found; ask for the EPSG code'}), "
                         "then call place_file")
            lines.append(f'- MAP FILE "{f.get("name")}" ({spatial.get("format")}, {spatial.get("layers")} layers, '
                         f'one sheet per layer): {where}. upload_id {spatial.get("upload_id")}.')
            for note in spatial.get("notes") or []:
                lines.append(f"    note: {note}")
        sheets = f.get("sheets") or []
        for s in sheets:
            from app.agent.sandbox import csv_name

            lines.append(f'- file "{f.get("name")}", sheet "{s.get("name")}": {s.get("total_rows")} rows; '
                         f'columns {json.dumps(s.get("columns"), ensure_ascii=False)}; run_python file '
                         f'"{csv_name(str(f.get("name", "file")), str(s.get("name", "sheet")), len(sheets) == 1)}"')
            totals = profile(s)
            if totals:
                lines.append(f"    totals (exact; use these, never add up yourself): {totals}")
            rows = s.get("rows") or []
            # A small sheet is shown whole (the blend test: 6 ingredients, 3 shown, and the person was asked to
            # name the other 3); a big one, its first rows and how to read the rest.
            width = max(len(s.get("columns") or []), 1)
            small = len(rows) <= SMALL_SHEET_ROWS or len(rows) * width <= SMALL_SHEET_CELLS
            shown = rows if small else rows[:3]
            for row in shown:
                lines.append("    " + json.dumps([_shown(v) for v in row], ensure_ascii=False, default=str))
            total = int(s.get("total_rows") or len(rows))
            if not small and total <= NAMES_LISTED and len(rows) >= total:
                names = [r[0] for r in rows if r and isinstance(r[0], str)]
                if len(names) == total and len(set(names)) == total:
                    lines.append(f"    every {json.dumps((s.get('columns') or ['first'])[0], ensure_ascii=False)}: "
                                 + json.dumps(names, ensure_ascii=False))
            lines.append("    (every row is shown above)" if len(shown) >= total else
                         f"    ...{total - len(shown)} more rows: read_file / query_file, never ask the user for them")
    return "\n".join(lines)


def _number(v: Any) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return None
    return None


def _fmt(x: float) -> str:
    return str(int(x)) if float(x).is_integer() else f"{x:.6g}"


def profile(s: dict[str, Any]) -> str:
    """Totals the model must not work out in its head (the platform test of October 2026: it summed five
    capacities to 2,450 instead of 2,850 and counted 11 lanes over a limit instead of 13): per numeric
    column count, sum, min and max; per text column the number of distinct values and any repeated."""
    columns, rows = s.get("columns") or [], s.get("rows") or []
    if not rows:
        return ""
    parts = []
    for i, col in enumerate(columns):
        if col in ("geometry", "kind"):
            continue
        cells = [r[i] for r in rows if i < len(r) and r[i] not in (None, "")]
        nums = [n for n in (_number(c) for c in cells) if n is not None]
        if nums and len(nums) == len(cells):
            total = "" if col in ("lon", "lat", "longitude", "latitude", "x", "y", "x_m", "y_m", "min_x_m", "min_y_m") \
                else f" sum {_fmt(sum(nums))}"
            parts.append(f"{col}:{total} min {_fmt(min(nums))} max {_fmt(max(nums))}")
        elif cells:
            distinct = {str(c).strip().casefold() for c in cells}
            repeated = len(cells) - len(distinct)
            parts.append(f"{col}: {len(distinct)} distinct" + (f" ({repeated} repeated)" if repeated and len(distinct) * 2 > len(cells) else "")
                         + (f", {len(rows) - len(cells)} empty" if len(cells) < len(rows) else ""))
    return "; ".join(parts)


_OPS = ("=", "!=", "<", "<=", ">", ">=", "in", "notIn", "contains", "empty", "notEmpty")


def _keep(row: list[Any], at: dict[str, int], where: list[dict[str, Any]]) -> bool:
    for f in where:
        col, op, want = f.get("column"), f.get("op", "="), f.get("value")
        if col not in at:
            raise FileRefused(f'no column "{col}" to filter on; columns: {list(at)}')
        if op not in _OPS:
            raise FileRefused(f'filter op "{op}" is not one of {list(_OPS)}')
        v = row[at[col]] if at[col] < len(row) else None
        if op in ("empty", "notEmpty"):
            ok = (v in (None, "")) == (op == "empty")
        elif op in ("in", "notIn"):
            wanted = {str(w).strip().casefold() for w in (want if isinstance(want, list) else [want])}
            ok = (str(v).strip().casefold() in wanted) == (op == "in")
        elif op == "contains":
            ok = str(want).casefold() in str(v or "").casefold()
        else:
            a, b = _number(v), _number(want)
            if a is None or b is None:  # text compares as people read it
                a, b = str(v or "").strip().casefold(), str(want or "").strip().casefold()
                if op not in ("=", "!="):
                    raise FileRefused(f'"{col}" {op} {want!r} compares text; use = or != or a number')
            ok = {"=": a == b, "!=": a != b, "<": a < b, "<=": a <= b, ">": a > b, ">=": a >= b}[op]
        if not ok:
            return False
    return True


def query(files: list[dict[str, Any]], file: str, sheet: str | None, where: list[dict[str, Any]] | None = None,
          group_by: list[str] | None = None, aggregate: dict[str, str] | None = None,
          columns: list[str] | None = None, limit: int = 50) -> str:
    """Count, filter, add up and group an attached sheet exactly, so the model never does arithmetic itself.
    Without `aggregate` it returns the matching rows (only `columns`, if given) and their count; with it,
    one line per group: {"capacity": "sum", "lane": "count"} (sum, min, max, mean, count, distinct)."""
    s = find_sheet(files, file, sheet)
    cols = s.get("columns") or []
    at = {c: i for i, c in enumerate(cols)}
    for c in list(group_by or []) + list(aggregate or {}) + list(columns or []):
        if c not in at and c != "*":
            raise FileRefused(f'sheet "{s.get("name")}" has no column "{c}"; columns: {cols}')
    rows = [r for r in s.get("rows") or [] if _keep(r, at, where or [])]
    head = f"{len(rows)} of {len(s.get('rows') or [])} rows match" + (f" {json.dumps(where)}" if where else "")
    out = io.StringIO()
    writer = csv.writer(out)
    if not aggregate:
        shown = columns or cols
        writer.writerow(shown)
        for r in rows[: max(1, min(int(limit or 50), 200))]:
            writer.writerow([_shown(r[at[c]]) if at[c] < len(r) else None for c in shown])
        more = len(rows) - min(len(rows), max(1, min(int(limit or 50), 200)))
        return head + "\n" + out.getvalue() + (f"...[{more} more matching rows]" if more > 0 else "")
    groups: dict[tuple, list[list[Any]]] = {}
    for r in rows:
        groups.setdefault(tuple(r[at[g]] for g in group_by or []), []).append(r)
    if not groups and not group_by:
        groups[()] = []
    writer.writerow(list(group_by or []) + [f"{fn}({c})" for c, fn in aggregate.items()])
    for key, members in groups.items():
        line = list(key)
        for c, fn in aggregate.items():
            if fn == "count":
                line.append(len(members) if c == "*" else sum(1 for m in members if m[at[c]] not in (None, "")))
                continue
            vals = [m[at[c]] for m in members if m[at[c]] not in (None, "")]
            if fn == "distinct":
                line.append(len({str(v).strip().casefold() for v in vals}))
                continue
            nums = [n for n in (_number(v) for v in vals) if n is not None]
            if fn not in ("sum", "min", "max", "mean"):
                raise FileRefused(f'aggregate "{fn}" is not one of sum, min, max, mean, count, distinct')
            if not nums:
                line.append(None)
                continue
            x = {"sum": sum(nums), "min": min(nums), "max": max(nums), "mean": sum(nums) / len(nums)}[fn]
            line.append(_fmt(round(x, 6)))
        writer.writerow(line)
    return head + "\n" + out.getvalue()


def find_sheet(files: list[dict[str, Any]], file: str, sheet: str | None) -> dict[str, Any]:
    match = next((f for f in files if f.get("name") == file), None)
    if match is None:
        raise FileRefused(f'no attached file named "{file}"; attached: {[f.get("name") for f in files]}')
    sheets = match.get("sheets") or []
    if sheet in (None, ""):
        if len(sheets) == 1:
            return sheets[0]
        raise FileRefused(f'"{file}" has several sheets; name one of {[s.get("name") for s in sheets]}')
    found = next((s for s in sheets if s.get("name") == sheet), None)
    if found is None:
        raise FileRefused(f'"{file}" has no sheet "{sheet}"; sheets: {[s.get("name") for s in sheets]}')
    return found


def read(files: list[dict[str, Any]], file: str, sheet: str | None, offset: int, limit: int) -> str:
    s = find_sheet(files, file, sheet)
    rows = (s.get("rows") or [])[max(0, offset): max(0, offset) + max(1, min(limit, 200))]
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(s.get("columns") or [])
    writer.writerows([[_shown(v) for v in r] for r in rows])
    rest = len(s.get("rows") or []) - offset - len(rows)
    return out.getvalue() + (f"...[{rest} more rows; read_file offset={offset + len(rows)}]" if rest > 0 else "")


def _column(s: dict[str, Any], col: str, what: str = "a column") -> int:
    columns = s.get("columns") or []
    if col in (None, ""):
        # The roster field test: 'sheet "nurses" has no column "None"' -- a missing field, not a missing column.
        raise FileRefused(f'{what} is missing: name one of the columns of sheet "{s.get("name")}": {columns}')
    if col not in columns:
        raise FileRefused(f'sheet "{s.get("name")}" has no column "{col}"; columns: {columns}')
    return columns.index(col)


def _columns(s: dict[str, Any], col: Any, what: str = "a column") -> list[int]:
    """One column, or several (["job", "step"]) whose values joined with "-" make a key: a step of a job, a
    shift of a day (the job-shop field test, October 2026: operations had no single key column)."""
    if isinstance(col, (list, tuple)) and col:
        return [_column(s, c, what) for c in col]
    return [_column(s, col, what)]


def _joined(row: list[Any], at: list[int]) -> str | None:
    parts = [row[i] for i in at]
    if any(p in (None, "") for p in parts):
        return None
    if len(parts) == 1:
        return str(parts[0])  # one column: the key as it always was

    def part(v: Any) -> str:
        return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    return "-".join(part(p) for p in parts)


_YES = {"yes", "y", "true", "t", "on"}
_NO = {"no", "n", "false", "f", "off"}


# A number written with its unit or currency: "$1,200", "1 200 kg", "EGP 45.5", "12.5 %".
_WITH_UNIT = re.compile(r"^\s*([$€£¥]|[A-Za-z]{3}\s)?\s*(-?\d[\d, ]*(?:\.\d+)?|-?\.\d+)\s*(%|[A-Za-z/²³°]{1,12})?\s*$")


def _numeric(value: Any) -> tuple[Any, str | None]:
    """A cell where a number belongs, as a number, and how it was read (None: as it was). A yes/no cell is 1/0
    (phase 1 evaluation, 10 October 2026: an `open` column of yes/no was refused "a number is required" three
    times, and nothing was built); a number with its unit or currency is the number; a percentage is a fraction."""
    if isinstance(value, bool):
        return int(value), "yes/no"
    if not isinstance(value, str):
        return value, None
    text = value.strip()
    if text.casefold() in _YES:
        return 1, "yes/no"
    if text.casefold() in _NO:
        return 0, "yes/no"
    m = _WITH_UNIT.match(text)
    if not m:
        return value, None
    digits = m.group(2).replace(",", "").replace(" ", "")
    try:
        number = float(digits)
    except ValueError:
        return value, None
    if m.group(3) == "%":
        number, how = number / 100, "percent"
    else:
        how = "unit" if m.group(1) or m.group(3) else None
    number = round(number, 12)
    return (int(number) if number.is_integer() else number), how


def _as_number(value: Any) -> Any:
    return _numeric(value)[0]


class _Load:
    """What one *_from_file entry did with its sheet, in the file's terms (neutrality plan, step 4: the plan card
    says what was stored against what the file holds -- phase 1, supply chain L, stored every route as open)."""

    READ_AS = {"yes/no": "yes/no read as 1/0", "percent": "percentages read as fractions (35% = 0.35)",
               "unit": "numbers read without their unit or currency"}

    def __init__(self, what: str, s: dict[str, Any], spec: dict[str, Any]):
        self.what, self.s, self.spec = what, s, spec
        self.total = int(s.get("total_rows") or len(s.get("rows") or []))
        self.kept = 0
        self.no_key = 0
        self.no_value = 0
        self.repeated = 0
        self.stored = 0
        self.read_as: dict[str, int] = {}
        self.not_number = 0
        self.columns_note = ""

    def number(self, value: Any) -> Any:
        out, how = _numeric(value)
        if how:
            self.read_as[how] = self.read_as.get(how, 0) + 1
        elif isinstance(out, str):
            self.not_number += 1
        return out

    def line(self) -> str:
        where = self.spec.get("where")
        parts = [f'{self.what} from "{self.spec.get("file")}"'
                 + (f' sheet "{self.s.get("name")}"' if self.spec.get("sheet") else "")
                 + f": {self.total} rows"]
        if where:
            parts.append(f"{self.kept} kept by the filter {json.dumps(where, ensure_ascii=False)}")
        if self.s.get("truncated"):
            parts.append(f"only the first {len(self.s.get('rows') or [])} read")
        if self.columns_note:
            parts.append(self.columns_note)
        if self.no_key:
            parts.append(f"{self.no_key} skipped: a key cell empty")
        if self.no_value:
            parts.append(f"{self.no_value} skipped: the value empty")
        if self.repeated:
            parts.append(f"{self.repeated} repeated keys counted once")
        for how, n in self.read_as.items():
            parts.append(f"{n} {self.READ_AS[how]}")
        if self.not_number:
            parts.append(f"{self.not_number} values are not numbers")
        parts.append(f"{self.stored} stored")
        return "; ".join(parts)


def _rows(s: dict[str, Any], spec: dict[str, Any]) -> list[list[Any]]:
    """The sheet's rows a *_from_file entry reads: all of them, or those its `where` keeps (phase 1 evaluation,
    10 October 2026: routes loaded with `"where": [{"column": "open", "value": "yes"}]` stored every route as
    open -- the filter was read by queries only, and a load ignored it without a word)."""
    where = spec.get("where")
    if not where:
        return list(s.get("rows") or [])
    if isinstance(where, dict):
        where = [where]
    if not isinstance(where, list) or not all(isinstance(f, dict) for f in where):
        raise FileRefused('"where" is a list of filters: [{"column": ..., "op": "=", "value": ...}]')
    at = {c: i for i, c in enumerate(s.get("columns") or [])}
    return [r for r in s.get("rows") or [] if _keep(r, at, where)]


def _wide(s: dict[str, Any], spec: dict[str, Any], pairs: list[tuple[str, Any]], wide: list[int],
          entities: list[dict[str, Any]], load: _Load) -> list[dict[str, Any]]:
    """A wide table's cells: one per row and per column whose header is a record of the "*" type (a cost per
    plant and customer, one column per customer)."""
    name = spec.get("parameter")
    if len(wide) > 1:
        raise FileRefused(f'parameter_values_from_file for "{name}": only one index can be read from the column '
                          'headers ("*")')
    if spec.get("value") is not None:
        raise FileRefused(f'parameter_values_from_file for "{name}" reads a wide table ("*"): its cells are the '
                          'values, so leave "value" out')
    col_type = pairs[wide[0]][0]
    row_ends = [(t, _columns(s, c)) for i, (t, c) in enumerate(pairs) if i != wide[0]]
    used = {i for _, at in row_ends for i in at}
    columns = s.get("columns") or []
    if spec.get("columns"):
        at = {c: _column(s, c) for c in spec["columns"]}
    else:
        known = {str(e.get("key")) for e in entities if isinstance(e, dict) and e.get("type") == col_type}
        others = [c for i, c in enumerate(columns) if i not in used]
        named = [c for c in others if str(c) in known]
        if known and not named:
            raise FileRefused(f'parameter_values_from_file for "{name}": no column header of sheet "{s.get("name")}" '
                              f'is a key of "{col_type}" (headers {others}); name the columns with "columns"')
        at = {c: columns.index(c) for c in (named if known else others)}
        left = [c for c in others if c not in at]
        if left:
            load.columns_note = f"columns {left} are not {col_type} keys and were not read"
    load.columns_note = (load.columns_note + "; " if load.columns_note else "") + f"{len(at)} columns of {col_type}"
    rows = _rows(s, spec)
    load.kept = len(rows)
    cells = []
    for row in rows:
        keys = [_joined(row, a) for _, a in row_ends]
        if None in keys:
            load.no_key += 1
            continue
        for header, i in at.items():
            value = load.number(row[i]) if i < len(row) else None
            if value is None:
                load.no_value += 1
                continue
            named_keys = [[t, k] for (t, _), k in zip(row_ends, keys)]
            named_keys.insert(wide[0], [col_type, str(header)])
            cells.append({"parameter": name, "entities": named_keys, "value": value})
            load.stored += 1
    return cells


def expand(seed: dict[str, Any], files: list[dict[str, Any]], report: list[str] | None = None) -> dict[str, Any]:
    """`entities_from_file` and `parameter_values_from_file` as plain `entities` and `parameter_values`; each may
    keep only the rows its `where` matches. `report`, if given, gets one line per entry: what it stored against
    what the file holds."""
    loads: list[_Load] = []
    out = {k: v for k, v in seed.items()
           if k not in ("entities_from_file", "parameter_values_from_file", "relationships_from_file")}
    if isinstance(seed.get("parameters"), list):  # a default may be corrected below: never the caller's copy
        out["parameters"] = [dict(p) if isinstance(p, dict) else p for p in seed["parameters"]]
    bindings = list(seed.get("source_bindings") or [])
    kept_files: dict[str, dict[str, Any]] = {}

    def bound(kind: str, spec: dict[str, Any], target: Any) -> None:
        """A *_from_file entry reading a workspace data source (use_source): remembered with the build, so a
        refresh can read the source again and update what this entry made (migration 0114)."""
        found = next((f for f in files if f.get("name") == spec.get("file")), None)
        if found is None:
            return
        src = found.get("source")
        mapping = {k: v for k, v in spec.items() if k != "file"}
        if isinstance(src, dict) and src.get("connection_id"):
            bindings.append({"connection_id": src["connection_id"], "job_id": src.get("job_id"),
                             "sha256": src.get("sha256"), "kind": kind, "target": str(target), "mapping": mapping})
            return
        # A file the person attached (or one kept in the workspace, re-attached): kept with the build as a
        # workspace file, version by version (migration 0115), and bound like a database source.
        from app.integrations.refresh import keepable

        name = (src or {}).get("file") or found["name"]
        if isinstance(src, dict) and src.get("file") or keepable(found):
            bindings.append({"file_name": name, "kind": kind, "target": str(target), "mapping": mapping})
            if not (isinstance(src, dict) and src.get("file")) and name not in kept_files:
                kept_files[name] = found

    entities = list(seed.get("entities") or [])
    # Fields declared as numbers: a yes/no cell read into one is 1/0.
    numeric = {(t.get("name"), a.get("name")) for t in seed.get("entity_types") or [] if isinstance(t, dict)
               for a in t.get("attributes") or [] if isinstance(a, dict)
               and str(a.get("data_type") or "").lower() in ("number", "integer", "int", "float", "decimal")}
    for spec in seed.get("entities_from_file") or []:
        s = find_sheet(files, spec.get("file"), spec.get("sheet"))
        bound("entities", spec, spec.get("type"))
        key_at = _columns(s, spec.get("key"), f'entities_from_file for "{spec.get("type")}": "key" (the column '
                                                        "with each record's key, or a list of columns)")
        label_at = _column(s, spec["label"]) if spec.get("label") else None
        attrs_at = {attr: _column(s, col) for attr, col in (spec.get("attrs") or {}).items()}
        made: dict[str, dict[str, Any]] = {}
        load = _Load(f'records "{spec.get("type")}"', s, spec)
        loads.append(load)
        rows = _rows(s, spec)
        load.kept = len(rows)
        for row in rows:
            key = _joined(row, key_at)
            if key is None:
                load.no_key += 1
                continue
            entity = {"type": spec.get("type"), "key": key,
                      "attrs": {a: load.number(row[i]) if (spec.get("type"), a) in numeric else _shape_or_value(row[i])
                                for a, i in attrs_at.items() if row[i] is not None}}
            # A key repeated in the file (days from a sheet with a row per day and shift -- the retest):
            # one record, when the rows agree on its fields; when they disagree, the file is asked about.
            first = made.get(entity["key"])
            if first is not None:
                if first["attrs"] != entity["attrs"]:
                    raise FileRefused(f'"{spec.get("file")}" gives {spec.get("type")} "{entity["key"]}" twice with '
                                      f'different fields ({first["attrs"]} and {entity["attrs"]}); which is right?')
                load.repeated += 1
                continue
            made[entity["key"]] = entity
            load.stored += 1
            if label_at is not None and row[label_at] is not None:
                entity["label"] = str(row[label_at])
            entities.append(entity)
    if entities:
        out["entities"] = entities
    links = list(seed.get("relationships") or [])
    for spec in seed.get("relationships_from_file") or []:
        # One link per row: {"type": "works_in", "from": ["employee", "employee column"], "to": ["unit", "unit column"]}
        s = find_sheet(files, spec.get("file"), spec.get("sheet"))
        bound("relationships", spec, spec.get("type"))
        ends = []
        for end in ("from", "to"):
            pair = spec.get(end)
            if not (isinstance(pair, (list, tuple)) and len(pair) == 2):
                raise FileRefused(f'relationships_from_file "{end}" must be [record type, column]')
            ends.append((str(pair[0]), _columns(s, pair[1])))
        load = _Load(f'links "{spec.get("type")}"', s, spec)
        loads.append(load)
        rows = _rows(s, spec)
        load.kept = len(rows)
        for row in rows:
            keys = [_joined(row, at) for _, at in ends]
            if None in keys:
                load.no_key += 1
                continue
            links.append({"type": spec.get("type"), "from": [ends[0][0], keys[0]], "to": [ends[1][0], keys[1]]})
            load.stored += 1
    if links:
        out["relationships"] = links
    from app.seed import field_distances, order_links, order_links_problems

    problems = order_links_problems(out)
    if problems:
        raise FileRefused("; ".join(problems))
    out = order_links(out)
    links = out.get("relationships") or []
    cells = list(seed.get("parameter_values") or [])
    for spec in seed.get("parameter_values_from_file") or []:
        s = find_sheet(files, spec.get("file"), spec.get("sheet"))
        bound("parameter_values", spec, spec.get("parameter"))
        load = _Load(f'parameter "{spec.get("parameter")}"', s, spec)
        loads.append(load)
        pairs = [(str(t), c) for t, c in (spec.get("entities") or [])]
        wide = [i for i, (_, c) in enumerate(pairs) if c == "*"]
        if wide:
            cells.extend(_wide(s, spec, pairs, wide, out.get("entities") or [], load))
            continue
        ends = [(t, _columns(s, c)) for t, c in pairs]
        if len(ends) >= 2 and len({tuple(at) for _, at in ends}) < len(ends):
            # Two indices read from one column give each record only with itself (the evaluation's routing
            # test: distance[place, place] from ["place","place"] twice -- 9 cells of 81, the rest 999,999).
            raise FileRefused(
                f'parameter_values_from_file for "{spec.get("parameter")}" reads two of its indices from the same '
                f'column, so only each record paired with itself would get a value. Name a different column for '
                f'each index (a file of pairs), or, for distances between records with coordinate columns, use '
                f'"distances_from_fields": [{{"name": "{spec.get("parameter")}", "of": <type>, "x": <column>, '
                f'"y": <column>}}] in the seed instead')
        # "value": a column, or a number for every row: a file that only LISTS pairs (days off asked for,
        # bans, skills held) is a 0/1 parameter with "value": 1 (the nurse roster field test, October 2026).
        constant = spec.get("value") if isinstance(spec.get("value"), (int, float)) \
            and not isinstance(spec.get("value"), bool) else None
        if constant not in (None, 0):
            # A file that LISTS pairs, loaded as `constant` on a default of the same: the list would change
            # nothing (the evaluation, 11 October 2026: allowed pairs as 1 on a default of 1, three times). When
            # the list names EVERY record of its first index, what it does not name is the opposite: 0. When it
            # names only some, it restricts those alone (the others are free), and the default must stay: the
            # same day, flipping it made every unlisted record unservable and the plan infeasible.
            declared = next((p for p in out.get("parameters") or []
                             if isinstance(p, dict) and p.get("name") == spec.get("parameter")), None)
            if declared is not None and _number(declared.get("default_value")) == float(constant):
                ends0 = [(str(t), _columns(s, c)) for t, c in (spec.get("entities") or [])[:1] if c != "*"]
                named = {_joined(r, ends0[0][1]) for r in _rows(s, spec)} - {None} if ends0 else set()
                every = {str(e.get("key")) for e in out.get("entities") or []
                         if isinstance(e, dict) and e.get("type") == (ends0[0][0] if ends0 else None)}
                if every and named >= every:
                    declared["default_value"] = 0
                    load.columns_note = (f"the default was {_fmt(float(constant))}, the same as every value the list "
                                         "gives, and the list names every record: every pair it does not name is now 0")
                else:
                    load.columns_note = (f"every value is {_fmt(float(constant))}, the same as the default, and the list "
                                         f"names {len(named)} of {len(every) or '?'} records: as loaded it changes "
                                         "nothing. If it restricts only the records it names, add a 0/1 parameter "
                                         "per record (1 for those named) and write the rule for them: amount <= most "
                                         "x (1 - listed[r] + pair[r, o]); if it is complete, give the default 0")
        value_at = None if constant is not None else _column(
            s, spec.get("value"), f'parameter_values_from_file for "{spec.get("parameter")}": "value" (a column, or a '
                                  "number for every row)")
        rows = _rows(s, spec)
        load.kept = len(rows)
        for row in rows:
            keys = [_joined(row, at) for _, at in ends]
            if None in keys:
                load.no_key += 1
                continue
            value = constant if value_at is None else load.number(row[value_at])
            if value is None:
                load.no_value += 1
                continue
            cells.append({"parameter": spec.get("parameter"),
                          "entities": [[t, k] for (t, _), k in zip(ends, keys)], "value": value})
            load.stored += 1
    if cells:
        out["parameter_values"] = cells
    if report is not None:
        report.extend(load.line() for load in loads)
    if bindings:
        out["source_bindings"] = bindings
    if kept_files:
        out["source_files"] = list(kept_files.values())
    from app.seed import fitting_patterns, pattern_problems

    problems = pattern_problems(out)
    if problems:
        raise FileRefused("; ".join(problems))
    return fitting_patterns(field_distances(out))


# -- map files: the platform's own CAD / GIS readers (app/gis), as tables of features ----------------

SPATIAL_SUFFIXES = (".dxf", ".geojson", ".kml", ".kmz", ".gpx", ".zip", ".shp", ".gpkg")
GEOGRAPHIC_FORMATS = ("GeoJSON", "KML", "KMZ", "GPX")


def is_spatial(filename: str, data: bytes) -> bool:
    """A map file: by its ending, or a .json that is GeoJSON, or a .csv with a shape or coordinate columns."""
    lower = filename.lower()
    if lower.endswith(SPATIAL_SUFFIXES):
        return True
    if lower.endswith(".json"):
        head = data[:4000].decode("utf-8-sig", errors="ignore")
        return '"FeatureCollection"' in head or '"Feature"' in head
    if lower.endswith(".csv"):
        from app.gis.formats import WKT_COLUMNS, XY_COLUMNS

        first = data[:4000].decode("utf-8-sig", errors="ignore").splitlines()[:1]
        header = {c.strip().strip('"').lower() for c in (first[0].replace(";", ",").split(",") if first else [])}
        return bool(header & set(WKT_COLUMNS)) or any(a in header and b in header for a, b in XY_COLUMNS)
    return False


#: Where a drawing in local metres is shown on the map until it is placed for real (GIS_LOCAL_ANCHOR="lon,lat").
LOCAL_ANCHOR = tuple(float(v) for v in __import__("os").environ.get("GIS_LOCAL_ANCHOR", "31.2357,30.0444").split(","))
#: Wider than this, in metres, a drawing with no coordinate system is probably not in metres.
LOCAL_SANE_M = 5000.0


def _local_metres(placement: Any):
    """A placed shape (lon/lat) back to local metres from the placement's anchor, as WKT -- what run_python's
    geometry code needs (shapely.wkt.loads): the camp-bed test's numbers were degrees read as metres."""
    from pyproj import Transformer
    from shapely.geometry import shape as to_shape
    from shapely.ops import transform
    from shapely import wkt

    lon0, lat0 = placement.lonlat
    back = Transformer.from_crs("EPSG:4326", f"+proj=aeqd +lat_0={lat0} +lon_0={lon0} +x_0=0 +y_0=0 +units=m "
                                "+ellps=WGS84", always_xy=True)

    def run(g: dict[str, Any]) -> str | None:
        try:
            return wkt.dumps(transform(back.transform, to_shape(g)), rounding_precision=3)
        except Exception:  # noqa: BLE001 -- a shape that cannot be read has no local form
            return None

    return run


def _shape_or_value(value: Any) -> Any:
    """A shape written into a CSV as GeoJSON text (a run_python output) is a shape again, so a geometry field
    takes it (the camp-bed retest: "attribute geometry must be geometry")."""
    if isinstance(value, str) and value.startswith("{") and '"coordinates"' in value and '"type"' in value:
        try:
            shape = json.loads(value)
        except ValueError:
            return value
        if isinstance(shape, dict) and "type" in shape and "coordinates" in shape:
            return shape
    return value


def _bbox(coords: Any) -> tuple[float, float, float, float] | None:
    xs: list[float] = []
    ys: list[float] = []

    def walk(c: Any) -> None:
        if isinstance(c, (list, tuple)) and len(c) >= 2 and all(isinstance(v, (int, float)) for v in c[:2]):
            xs.append(float(c[0]))
            ys.append(float(c[1]))
        elif isinstance(c, (list, tuple)):
            for item in c:
                walk(item)

    walk(coords)
    return (min(xs), min(ys), max(xs), max(ys)) if xs else None


def _centre(coords: Any) -> tuple[float, float] | None:
    flat: list[tuple[float, float]] = []

    def walk(c: Any) -> None:
        if isinstance(c, (list, tuple)) and len(c) >= 2 and all(isinstance(v, (int, float)) for v in c[:2]):
            flat.append((float(c[0]), float(c[1])))
        elif isinstance(c, (list, tuple)):
            for item in c:
                walk(item)

    walk(coords)
    if not flat:
        return None
    return (sum(p[0] for p in flat) / len(flat), sum(p[1] for p in flat) / len(flat))


def parse_spatial(filename: str, data: bytes, placement: dict[str, Any] | None = None,
                  upload_id: str | None = None) -> dict[str, Any]:
    """A map file as one table per layer. Placed (in longitude/latitude, with each feature's shape, area and
    length) when `placement` is given or the platform is sure of the file's coordinate system; otherwise in
    the file's own coordinates, with the likely systems to ask the person about."""
    from shapely.geometry import shape as to_shape

    from app.gis import cad, crs, store
    from app.gis.convert import to_geojson
    from app.gis.to_records import measure, storable

    try:
        drawing = store.read_bytes(data, filename)
    except cad.CadError as exc:
        raise FileRefused(str(exc)) from exc
    if not drawing.features:
        raise FileRefused("the file has nothing that can be shown on a map")
    offered = crs.candidates(drawing.extent, drawing.unit_metres, geodata=drawing.geodata)
    local = False
    notes = list(drawing.notes or [])[:5]
    if placement is None:
        if drawing.version in GEOGRAPHIC_FORMATS:
            placement = {"kind": "epsg", "code": 4326}
        elif offered and offered[0].get("sure"):
            placement = offered[0]["placement"]
        elif drawing.extent is not None:
            # A drawing with no known coordinate system (a CAD site plan, a shapefile without .prj): LOCAL METRES,
            # never degrees. The camp-bed test: the Assistant offered EPSG:4326 for a drawing in metres, the
            # user took it, and an 88 m camp became "86 km" -- degrees read as metres. Positions are measured
            # from the drawing's lower-left corner; the map position is nominal until place_file is used.
            local = True
            placement = {"kind": "local", "anchor": [drawing.extent[0], drawing.extent[1]],
                         "lonlat": list(LOCAL_ANCHOR), "rotation": 0, "scale": 1}
            unit = drawing.unit_metres or 1.0
            if not drawing.unit_metres:
                notes.append("the drawing does not say its units: metres are assumed")
            span = max(drawing.extent[2] - drawing.extent[0], drawing.extent[3] - drawing.extent[1]) * unit
            if span > LOCAL_SANE_M:
                notes.append(f"SCALE CHECK: in metres the drawing is {span / 1000:,.1f} km across. For a site plan that "
                             "is too big: its units are probably millimetres or centimetres. Ask the user before "
                             "using any length or area.")
            elif 0 < span < 1:
                notes.append(f"SCALE CHECK: in metres the drawing is {span:.3f} m across: its units are probably "
                             "kilometres. Ask the user before using any length or area.")
    placed = None
    if placement is not None:
        try:
            placed = crs.Placement.parse(placement, drawing.unit_metres)
        except crs.PlacementError as exc:
            raise FileRefused(str(exc)) from exc

    sheets: list[dict[str, Any]] = []
    for layer in list(drawing.layers)[:MAX_SHEETS * 5]:
        feats = [f for f in drawing.features if f.layer == layer][:MAX_ROWS]
        if not feats:
            continue
        keys = list(dict.fromkeys(k for f in feats for k in (f.props or {})))[:MAX_COLUMNS]
        rows: list[list[Any]] = []
        if placed is not None and local:
            # Local metres for the model; lon/lat (nominal) and the shape kept for the platform's own measures.
            from app.gis.convert import place as place_shapes

            shapes, _ = place_shapes(feats, placed)  # (index into feats, shape): broken polygons are left out
            unit = drawing.unit_metres or 1.0
            x0, y0 = drawing.extent[0], drawing.extent[1]  # type: ignore[index]
            columns = ["feature", "kind", "x_m", "y_m", "min_x_m", "min_y_m", "width_m", "height_m", "area_m2",
                       "length_m", "geometry", "shape_m", *keys]
            to_metres = _local_metres(placed)
            for n, (i, shape) in enumerate(shapes, start=1):
                raw = feats[i]
                f = {"geometry": shape, "properties": {"kind": raw.kind, **(raw.props or {})}}
                box = _bbox(raw.coords)
                c = _centre(raw.coords)
                try:
                    m = measure(shape)
                except Exception:  # noqa: BLE001
                    m = {}
                r = lambda v: None if v is None else round(v, 3)  # noqa: E731
                rows.append([f"{layer}_{n}", f["properties"].get("kind"),
                             r((c[0] - x0) * unit) if c else None, r((c[1] - y0) * unit) if c else None,
                             r((box[0] - x0) * unit) if box else None, r((box[1] - y0) * unit) if box else None,
                             r((box[2] - box[0]) * unit) if box else None, r((box[3] - box[1]) * unit) if box else None,
                             m.get("area_m2"), m.get("length_m"), storable(f["geometry"]), to_metres(shape),
                             *[_cell((f["properties"] or {}).get(k)) for k in keys]])
        elif placed is not None:
            geo, _ = to_geojson(feats, placed)
            columns = ["feature", "kind", "lon", "lat", "area_m2", "length_m", "geometry", *keys]
            for n, f in enumerate(geo, start=1):
                g = f["geometry"]
                try:
                    point = to_shape(g).representative_point()
                    lon, lat = round(point.x, 7), round(point.y, 7)
                except Exception:  # noqa: BLE001 -- a shape shapely cannot read keeps no position
                    lon = lat = None
                try:
                    m = measure(g)
                except Exception:  # noqa: BLE001
                    m = {}
                props = f["properties"]
                rows.append([f"{layer}_{n}", props.get("kind"), lon, lat, m.get("area_m2"), m.get("length_m"),
                             storable(g), *[_cell(props.get(k)) for k in keys]])
        else:
            columns = ["feature", "kind", "x", "y", *keys]
            for n, f in enumerate(feats, start=1):
                c = _centre(f.coords)
                x, y = (round(c[0], 3), round(c[1], 3)) if c else (None, None)
                rows.append([f"{layer}_{n}", f.kind, x, y, *[_cell((f.props or {}).get(k)) for k in keys]])
        total = sum(1 for f in drawing.features if f.layer == layer)
        sheets.append({"name": layer, "columns": columns, "rows": rows, "total_rows": total,
                       "truncated": total > len(rows)})
    if not sheets:
        raise FileRefused("the file has nothing that can be shown on a map")
    chosen = None
    if placed is not None:
        info = crs.crs_info(placed.code) if placed.kind == "epsg" else None
        chosen = {**placement, "name": info["name"] if info else
                  ("local metres (drawing coordinates)" if local else "local placement")}
    return {"name": filename, "sheets": sheets, "spatial": {
        "format": drawing.version, "upload_id": upload_id, "placed": placed is not None, "placement": chosen,
        "local_metres": local, "units": drawing.units_name, "unit_metres": drawing.unit_metres,
        "extent": list(drawing.extent) if drawing.extent else None, "layers": len(sheets),
        "candidates": [{"code": c["placement"].get("code"), "name": c["name"], "reason": c["reason"],
                        "lands_at": c["centre"], "sure": c.get("sure", False)} for c in offered[:5]],
        "notes": notes}}


_WKT = ("POINT", "LINESTRING", "POLYGON", "MULTIPOINT", "MULTILINESTRING", "MULTIPOLYGON", "GEOMETRYCOLLECTION")


def _shown(value: Any) -> Any:
    """A shape is long; in what the model reads it is its kind and size."""
    if isinstance(value, dict) and "type" in value and "coordinates" in value:
        return f"<{value['type']}>"
    if isinstance(value, str) and len(value) > 60 and value.startswith(_WKT):
        return f"<WKT {value.split(' ', 1)[0].title()} in metres>"
    return value
