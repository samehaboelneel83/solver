"""The camp as a spreadsheet: what a planner edits between the drawing and the engine.

    drawing (.dxf) --dxf_to_workbook--> camp.xlsx --read_workbook--> CampProblem --solve--> layout

A drawing says where things are; it cannot say how many people a door takes,
how big a bed is, or which beds may shrink. The workbook holds both: the
geometry the converter read (editable -- coordinates are "x,y; x,y; ...") and
the rules, prefilled with sensible defaults and marked for review.

Sheets
  camp            setting | value | meaning          (name, grid, corridor width, origin, ...)
  doors           id, x1, y1, x2, y2, depth, capacity, zone_depth, zone_max_depth, zone_step, zone_margin, area_per_bed
  obstacles       id, kind, coordinates, area_m2
  no_beds         id, coordinates, area_m2
  zones           id, coordinates, area_m2
  bed_types       id, length, width, min_length, max_length, min_width, max_width, sizes, rotation,
                  side_gap, min_count, max_count, zone, priority
  boundary        coordinates (one row), area_m2
"""
from __future__ import annotations

from pathlib import Path

from shapely.geometry import Polygon

from .problem import (BedType, CampProblem, CorridorSpec, Door, DoorZone, Obstacle, ObjectiveSpec, PlacementZone,
                      Prohibited)

SETTINGS = [
    ("name", "Camp", "the camp's name"),
    ("grid", 0.5, "grid pitch, m: smaller is more exact and slower"),
    ("corridor_min_width", 1.0, "narrowest corridor, m: a whole number of grid cells"),
    ("corridor_access", "long_sides", "long_sides: a bed is entered from a long side; any_side: from any side"),
    ("objective_mode", "lexicographic", "lexicographic or weighted"),
    ("objective_order", "beds, distance, corridor, modifications", "lexicographic order"),
    ("trade_beds", "no", "may later objectives give beds back? yes/no"),
    ("origin_lon", 31.60, "longitude of the local origin (0, 0), for GIS"),
    ("origin_lat", 30.10, "latitude of the local origin (0, 0), for GIS"),
    ("offset_x", 0.0, "metres subtracted from the drawing's x (surveyed drawings)"),
    ("offset_y", 0.0, "metres subtracted from the drawing's y"),
]
DOOR_COLS = ["id", "x1", "y1", "x2", "y2", "depth", "capacity", "zone_depth", "zone_max_depth", "zone_step",
             "zone_margin", "area_per_bed"]
BED_COLS = ["id", "length", "width", "min_length", "max_length", "min_width", "max_width", "sizes", "rotation",
            "side_gap", "min_count", "max_count", "zone", "priority"]
DEFAULT_BEDS = [BedType("cot", length=2.0, width=0.9, rotation=True, side_gap=0.1)]


def _coords(ring) -> str:
    pts = list(ring)
    if len(pts) > 1 and tuple(pts[0]) == tuple(pts[-1]):
        pts = pts[:-1]
    return "; ".join(f"{x:g},{y:g}" for x, y in pts)


def _ring(text: str, where: str) -> list[tuple[float, float]]:
    try:
        pts = [tuple(float(v) for v in part.split(",")) for part in str(text).split(";") if part.strip()]
    except ValueError as exc:
        raise ValueError(f"{where}: coordinates must read 'x,y; x,y; ...' ({exc})") from None
    if len(pts) < 3 or any(len(p) != 2 for p in pts):
        raise ValueError(f"{where}: a shape needs at least three 'x,y' points")
    return [(p[0], p[1]) for p in pts]


def _yes(v) -> bool:
    return str(v).strip().lower() in ("yes", "y", "true", "1")


def _num(v, default=None):
    if v is None or (isinstance(v, str) and not v.strip()):
        return default
    return float(v)


def write_workbook(problem: CampProblem, path: str | Path, *, offset=(0.0, 0.0), notes: list[str] = ()) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    review = PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths or [16] * len(header), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "camp"
    spec = problem.objectives
    values = {
        "name": problem.name, "grid": problem.grid, "corridor_min_width": problem.corridor.min_width,
        "corridor_access": problem.corridor.access, "objective_mode": spec.mode,
        "objective_order": ", ".join(spec.order), "trade_beds": "yes" if spec.trade_beds else "no",
        "origin_lon": problem.origin_lonlat[0], "origin_lat": problem.origin_lonlat[1],
        "offset_x": offset[0], "offset_y": offset[1],
    }
    ws.append(["setting", "value", "meaning"])
    for c in ws[1]:
        c.font = bold
    for key, _, meaning in SETTINGS:
        ws.append([key, values[key], meaning])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 22, 40, 70
    for n in notes:
        ws.append(["note", n, "from the converter"])

    sheet("boundary", ["coordinates", "area_m2"], [[_coords(problem.boundary), round(Polygon(problem.boundary).area, 2)]],
          [120, 12])
    zone_of = {z.door: z for z in problem.zones}
    rows = []
    for d in problem.doors:
        z = zone_of.get(d.id, DoorZone(d.id))
        rows.append([d.id, d.a[0], d.a[1], d.b[0], d.b[1], d.depth, d.capacity, z.depth, z.max_depth, z.step,
                     z.margin, z.area_per_bed])
    ws_d = sheet("doors", DOOR_COLS, rows, [14] + [9] * 11)
    sheet("obstacles", ["id", "kind", "coordinates", "area_m2"],
          [[o.id, o.kind, _coords(o.ring), round(Polygon(o.ring).area, 2)] for o in problem.obstacles], [18, 12, 90, 10])
    sheet("no_beds", ["id", "coordinates", "area_m2"],
          [[p.id, _coords(p.ring), round(Polygon(p.ring).area, 2)] for p in problem.prohibited], [18, 90, 10])
    sheet("zones", ["id", "coordinates", "area_m2"],
          [[z.id, _coords(z.ring), round(Polygon(z.ring).area, 2)] for z in problem.placement_zones], [18, 90, 10])
    beds = []
    for b in problem.bed_types:
        beds.append([b.id, b.length, b.width, b.min_length, b.max_length, b.min_width, b.max_width,
                     "; ".join(f"{l:g}x{w:g}" for l, w in b.sizes), "yes" if b.rotation else "no", b.side_gap,
                     b.min_count, b.max_count, b.zone, b.priority])
    ws_b = sheet("bed_types", BED_COLS, beds, [12] + [10] * 13)
    if notes:
        # A converted drawing's rules are defaults: shade them for review.
        for ws_ in (ws_d, ws_b):
            for row in ws_.iter_rows(min_row=2):
                for c in row[5:] if ws_ is ws_d else row[1:]:
                    c.fill = review
    path = Path(path)
    wb.save(path)
    return path


def read_workbook(path: str | Path) -> CampProblem:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    need = {"camp", "boundary", "doors", "bed_types"}
    missing = need - set(wb.sheetnames)
    if missing:
        raise ValueError(f"the workbook has no {', '.join(sorted(missing))} sheet")

    def rows(name):
        if name not in wb.sheetnames:
            return []
        ws = wb[name]
        header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        out = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in r):
                continue
            out.append(dict(zip(header, r)))
        return out

    s = {r["setting"]: r["value"] for r in rows("camp") if r.get("setting") and r["setting"] != "note"}
    defaults = {k: v for k, v, _ in SETTINGS}
    get = lambda k: s.get(k, defaults[k]) if s.get(k) not in (None, "") else defaults[k]
    boundary = _ring(rows("boundary")[0]["coordinates"], "boundary")
    doors, zones = [], []
    for i, r in enumerate(rows("doors"), start=2):
        did = str(r["id"])
        doors.append(Door(did, (float(r["x1"]), float(r["y1"])), (float(r["x2"]), float(r["y2"])),
                          depth=_num(r.get("depth"), 0.3),
                          capacity=int(r["capacity"]) if r.get("capacity") not in (None, "") else None))
        zones.append(DoorZone(did, depth=_num(r.get("zone_depth"), 2.0), max_depth=_num(r.get("zone_max_depth"), 4.0),
                              step=_num(r.get("zone_step"), 1.0), margin=_num(r.get("zone_margin"), 0.5),
                              area_per_bed=_num(r.get("area_per_bed"), None)))
    obstacles = [Obstacle(str(r["id"]), _ring(r["coordinates"], f"obstacle {r['id']}"), str(r.get("kind") or "closed"))
                 for r in rows("obstacles")]
    prohibited = [Prohibited(str(r["id"]), _ring(r["coordinates"], f"no_beds {r['id']}")) for r in rows("no_beds")]
    pzones = [PlacementZone(str(r["id"]), _ring(r["coordinates"], f"zone {r['id']}")) for r in rows("zones")]
    beds = []
    for r in rows("bed_types"):
        sizes = tuple(tuple(float(v) for v in part.lower().split("x")) for part in str(r.get("sizes") or "").split(";")
                      if part.strip())
        beds.append(BedType(
            str(r["id"]), float(r["length"]), float(r["width"]),
            min_length=_num(r.get("min_length")), max_length=_num(r.get("max_length")),
            min_width=_num(r.get("min_width")), max_width=_num(r.get("max_width")),
            sizes=sizes, rotation=_yes(r.get("rotation", "yes")), side_gap=_num(r.get("side_gap"), 0.1),
            min_count=int(r.get("min_count") or 0),
            max_count=int(r["max_count"]) if r.get("max_count") not in (None, "") else None,
            zone=str(r["zone"]) if r.get("zone") not in (None, "") else None,
            priority=_num(r.get("priority"), 1.0)))
    order = tuple(x.strip() for x in str(get("objective_order")).split(",") if x.strip())
    return CampProblem(
        name=str(get("name")), boundary=boundary, doors=tuple(doors), zones=tuple(zones),
        obstacles=tuple(obstacles), prohibited=tuple(prohibited), placement_zones=tuple(pzones),
        bed_types=tuple(beds), grid=float(get("grid")),
        corridor=CorridorSpec(min_width=float(get("corridor_min_width")), access=str(get("corridor_access"))),
        objectives=ObjectiveSpec(mode=str(get("objective_mode")), order=order, trade_beds=_yes(get("trade_beds"))),
        origin_lonlat=(float(get("origin_lon")), float(get("origin_lat"))),
    )


def dxf_to_workbook(dxf_path: str, xlsx_path: str, *, name: str | None = None, units: str | None = None,
                    crs: str | None = None, bed_types=None, door_capacity: int | None = None) -> tuple[Path, list[str]]:
    """Read a drawing and write the workbook to review: geometry from the
    drawing, rules from the defaults (shaded)."""
    from .dxf import read_dxf

    d = read_dxf(dxf_path, units=units, crs=crs)
    ring = lambda poly: [(round(x, 4), round(y, 4)) for x, y in poly.exterior.coords[:-1]]
    problem = CampProblem(
        name=name or Path(dxf_path).stem.replace("_", " "),
        boundary=ring(d.boundary),
        doors=tuple(Door(x.id, x.a, x.b, capacity=door_capacity) for x in d.doors),
        zones=tuple(DoorZone(x.id, area_per_bed=0.25) for x in d.doors),
        obstacles=tuple(Obstacle(o.id, ring(o.polygon)) for o in d.obstacles),
        prohibited=tuple(Prohibited(p.id, ring(p.polygon)) for p in d.prohibited),
        placement_zones=tuple(PlacementZone(z.id, ring(z.polygon)) for z in d.zones),
        bed_types=tuple(bed_types or DEFAULT_BEDS),
        origin_lonlat=d.origin_lonlat or (31.60, 30.10),
    )
    notes = list(d.notes) + [f"read {len(d.doors)} doors, {len(d.obstacles)} obstacles, {len(d.prohibited)} no-bed areas, "
                             f"{len(d.zones)} zones; shaded cells are defaults to review"]
    return write_workbook(problem, xlsx_path, offset=d.offset, notes=notes), notes
